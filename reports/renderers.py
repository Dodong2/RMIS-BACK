import csv
import io
import re

CONTENT_TYPES = {
    "csv": "text/csv",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


def normalize_to_rows(data):
    """Every appendix export is either a list of dicts (Appendix E, one row
    per year) or a single dict (Appendix F/G, one snapshot) — flatten the
    dict case into field/value rows so all four renderers below can share
    one code path regardless of shape."""
    if isinstance(data, list):
        rows = data
    elif isinstance(data, dict):
        rows = [{"field": k, "value": v} for k, v in data.items()]
    else:
        rows = [{"value": data}]

    fieldnames = list(rows[0].keys()) if rows else []
    stringified = [{k: ("" if v is None else str(v)) for k, v in row.items()} for row in rows]
    return stringified, fieldnames


def render_csv(title, data):
    rows, fieldnames = normalize_to_rows(data)
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=fieldnames)
    writer.writeheader()
    writer.writerows(rows)
    return buf.getvalue().encode("utf-8")


def render_xlsx(title, data):
    from openpyxl import Workbook

    rows, fieldnames = normalize_to_rows(data)
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", "-", title)[:31] or "Report"  # Excel forbids []:*?/\ in sheet names
    ws.append(fieldnames)
    for row in rows:
        ws.append([row[f] for f in fieldnames])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render_pdf(title, data):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet

    rows, fieldnames = normalize_to_rows(data)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter)
    styles = getSampleStyleSheet()
    elements = [Paragraph(title, styles["Title"]), Spacer(1, 12)]

    if fieldnames:
        table_data = [fieldnames] + [[row[f] for f in fieldnames] for row in rows]
        table = Table(table_data, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#333333")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            )
        )
        elements.append(table)
    else:
        elements.append(Paragraph("No data available.", styles["Normal"]))

    doc.build(elements)
    return buf.getvalue()


def render_docx(title, data):
    from docx import Document

    rows, fieldnames = normalize_to_rows(data)
    doc = Document()
    doc.add_heading(title, level=1)

    if fieldnames:
        table = doc.add_table(rows=1, cols=len(fieldnames))
        table.style = "Light Grid Accent 1"
        for i, name in enumerate(fieldnames):
            table.rows[0].cells[i].text = name
        for row in rows:
            cells = table.add_row().cells
            for i, name in enumerate(fieldnames):
                cells[i].text = row[name]
    else:
        doc.add_paragraph("No data available.")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


RENDERERS = {"csv": render_csv, "xlsx": render_xlsx, "pdf": render_pdf, "docx": render_docx}


def render(format, title, data):
    if format not in RENDERERS:
        raise ValueError(f"Unsupported format: {format}. Choose one of {list(RENDERERS)}.")
    return RENDERERS[format](title, data)
