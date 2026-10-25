"""Official LSPU forms for Appendix E (LSPU-RDO-SF-017) and Appendix F (SF-16 terminal narrative outline).

Client follow-up 2026-10-02: the export must look like the form whichever file type is picked. So each form is
described once as a list of blocks (header lines, tables, sections, signatures) and render_form() draws the same
blocks as PDF, Word, Excel or CSV. The generic field/value renderers in renderers.py stay for the other reports.

Block kinds:
  {"kind": "header", "lines": [...], "title": str}          university header + form title
  {"kind": "line", "text": str, "italic": bool}             one line of text
  {"kind": "table", "header": [[...], ...], "rows": [[...]], "widths": [...], "spans": [(row, first, last)]}
      spans merge header cells of header row `row` from column `first` to `last`
      optional: "header_fill" (hex), "row_fills" {row: hex}, "cell_fills" [(row, col, hex)], "bold_rows" [row],
      "bold_first" (bold first column; default: tables without a header). Rows here are body rows.
  {"kind": "heading", "text": str, "level": 1|2}            outline heading (SF-16)
  {"kind": "section", "heading": str, "body": str, "level": 1|2}  heading + text; blank body = space to write in
  {"kind": "signatures", "items": [(label, name, role), ...]}
  {"kind": "footer", "parts": [left, center, right]}
  {"kind": "page_break"}
"""
import csv
import datetime
import io
import re
from xml.sax.saxutils import escape, unescape

from monitoring.models import MidtermReport, TerminalReport
from outputs.serializers import COMPUTED_6P

from .richtext import rich_text_lines, rich_text_marks

UNIVERSITY = ["Republic of the Philippines", "Laguna State Polytechnic University", "Province of Laguna"]
HEADER_FILL = "BDD7EE"  # the light blue of the printed SF-017 headers
# 6P wording of the LSPU forms (SF-017/SF-018), not the RMIS labels
FORM_6P_LABELS = {
    "publications": "Publications", "patents": "Patent", "products": "Products",
    "people_services": "People Services", "places_partnerships": "Places/Partnerships",
    "policies": "Policy Recommendations",
}


def _name(user):
    return (user.get_full_name() or user.email) if user else ""


def _pct(value):
    return "" if value is None else f"{value:g}%"


def _objectives(project):
    """Project objectives as a list; manual registration saves them as "1. ...\\n2. ..."."""
    lines = [re.sub(r"^\d+[.)]\s*", "", line).strip() for line in (project.objectives or "").splitlines()]
    return [line for line in lines if line]


def _dean(project):
    endorser = project.endorsers.filter(designation__icontains="dean").first()
    return endorser.name if endorser else ""


# ---- Appendix E: LSPU-RDO-SF-017 Midterm/Terminal Report ----------------------------------------------------------

def sf017_blocks(report):
    project = report.project
    objectives = report.objective_accomplishments or [
        {"objective": text, "q1": None, "q2": None, "q3": None, "q4": None} for text in _objectives(project)
    ]
    outputs = []
    for i, output in enumerate(project.expected_outputs.order_by("id"), start=1):
        related = COMPUTED_6P.get(output.category)
        actual = getattr(project, related).count() if related else output.manual_actual_count
        outputs.append([str(i), FORM_6P_LABELS.get(output.category, output.get_category_display()),
                        output.description, str(output.target_count), f"Actual: {actual}"])
    return [
        {"kind": "header", "lines": UNIVERSITY, "title": "MIDTERM/TERMINAL REPORT (FOR-LSPU-FUNDED PROJECT)"},
        {"kind": "line", "text": f"Project: {project.title} ({project.project_code}) · Project Year {report.project_year}"},
        {
            "kind": "table",
            "header": [["", "SPECIFIC OBJECTIVES", "% Accomplishment", "", "", ""], ["", "", "Q1", "Q2", "Q3", "Q4"]],
            "spans": [(0, 2, 5)],
            "rows": [[str(i), row["objective"], *(_pct(row.get(q)) for q in ("q1", "q2", "q3", "q4"))]
                     for i, row in enumerate(objectives, start=1)],
            "widths": [4, 52, 11, 11, 11, 11],
        },
        {"kind": "line", "text": "*please add additional lines as needed", "italic": True},
        {
            "kind": "table",
            "header": [["", "LSPU 6Ps Matrix", "", "", ""], ["", "Item", "Particulars", "Quantity", "Remarks"]],
            "spans": [(0, 1, 4)],
            "rows": outputs,
            "widths": [4, 22, 46, 14, 14],
        },
        {"kind": "line", "text": "**For Terminal Report, please attach narrative report and financial statement", "italic": True},
        {"kind": "signatures", "items": [
            ("Submitted by:", _name(project.lead), "Project Leader"),
            ("Noted by:", _dean(project), "Dean/Associate Dean"),
        ]},
        {"kind": "footer", "parts": ["LSPU-RDO-SF-017", "Rev. 0", "8 August 2018"]},
    ]


def appendix_e_form(project_id):
    """One SF-017 per submitted project year, or None when there is no midterm report yet."""
    reports = MidtermReport.objects.filter(project_id=project_id).select_related("project__lead").order_by("project_year")
    blocks = []
    for report in reports:
        if blocks:
            blocks.append({"kind": "page_break"})
        blocks.extend(sf017_blocks(report))
    return {"title": "Appendix E - Midterm Progress Report", "landscape": True, "blocks": blocks} if blocks else None


# ---- Appendix F: SF-16 Terminal Narrative Report -------------------------------------------------------------------

def appendix_f_form(project_id):
    """SF-16 outline. Sections RMIS already holds are filled from the registration (SF-018) and the terminal report;
    the rest are left blank for the leader to write (client follow-up 2026-10-02)."""
    report = TerminalReport.objects.filter(project_id=project_id).select_related("project__lead").first()
    if not report:
        return None
    project = report.project
    team = list(project.team_members.order_by("id"))
    duration = " – ".join(f"{d:%B %Y}" for d in (project.start_date, project.target_end_date) if d)
    objectives = "\n".join(f"{i}. {text}" for i, text in enumerate(_objectives(project), start=1))
    main_text = [
        ("Introduction", "", 1),
        ("Project Rationale", rich_text_lines(project.background), 1),
        ("Research Framework", "", 1),
        ("Objectives of the Project/Study", objectives, 1),
        ("Conceptual Model", "", 1),
        ("Significance of the Project/Study", rich_text_lines(project.socio_economic_significance), 1),
        ("Research Methodology", rich_text_lines(project.methodology), 1),
        ("Research/Experimental Design", "", 2),
        ("Sampling Design (if applicable)", "", 2),
        ("Instruments and Procedures", "", 2),
        ("Data Collection and Analysis", "", 2),
        ("Results and Discussion", "", 1),
        ("Summary of Findings", "", 1),
        ("Conclusion/s", "", 1),
        ("Recommendation/s", "", 1),
        ("Literature Cited", rich_text_lines(project.references), 1),
    ]
    contents = ["Acknowledgement", "Mentoring Beneficiaries", "Executive Summary"] + [h for h, _, _ in main_text]
    blocks = [
        {"kind": "header", "lines": UNIVERSITY, "title": "TERMINAL REPORT"},
        {"kind": "line", "text": project.title},
        {
            "kind": "table",
            "header": [],
            "rows": [
                ["Project Code", project.project_code],
                ["Project Leader", _name(project.lead)],
                ["Co-Project Leader/s", "; ".join(m.name for m in team if m.member_role == "co_leader")],
                ["Project Team", "; ".join(m.name for m in team if m.member_role != "co_leader")],
                ["Campus", project.campus],
                ["College Unit", project.college],
                ["Implementing Unit", project.implementing_unit],
                ["Duration", duration],
                ["Total Project Cost", f"PhP {project.total_cost:,.2f}" if project.total_cost is not None else ""],
                ["Date Submitted", f"{report.submitted_at:%B %d, %Y}"],
                ["Certified by RDO/RIUH", f"Yes, {_name(report.certified_by)}" if report.is_certified else "Not yet"],
            ],
            "widths": [30, 70],
        },
        {"kind": "page_break"},
        {"kind": "heading", "text": "PRELIMINARY PAGES", "level": 1},
        {"kind": "section", "heading": "Acknowledgement", "body": "", "level": 2},
        {"kind": "section", "heading": "Mentoring Beneficiaries", "body": "", "level": 2},
        {"kind": "section", "heading": "Table of Contents", "body": "\n".join(contents), "level": 2},
        {"kind": "section", "heading": "Executive Summary", "body": report.narrative, "level": 2},
        {"kind": "heading", "text": "MAIN TEXT", "level": 1},
        *({"kind": "section", "heading": h, "body": body or "", "level": level + 1} for h, body, level in main_text),
    ]
    return {"title": "Appendix F - Terminal Report", "landscape": False, "blocks": blocks}


# ---- LSPU-RDO-SF-018 Research Proposal Form ----------------------------------------------------------------------

PROPOSAL_FILL = "E7DFC6"  # the beige section bars of the frontend's ProposalPreview
GANTT_FILL = "FDE047"
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
SDG_LABELS = [
    "No Poverty", "Zero Hunger", "Good Health and Well-being", "Quality Education", "Gender Equality",
    "Clean Water and Sanitation", "Affordable and Clean Energy", "Decent Work and Economic Growth",
    "Industry, Innovation and Infrastructure", "Reduced Inequalities", "Sustainable Cities and Communities",
    "Responsible Consumption and Production", "Climate Action", "Life Below Water", "Life on Land",
    "Peace, Justice and Strong Institutions", "Partnerships for the Goals",
]
BUDGET_GROUPS = [
    ("ps", "PERSONAL SERVICES (PS)"),
    ("mooe", "MAINTENANCE AND OTHER OPERATING EXPENSES (MOOE)"),
    ("co", "EQUIPMENT OUTLAY (CO)"),
]


def _box(checked, label):
    return f"[{'/' if checked else ' '}] {label}"


def _boxes(choices, picked):
    return "    ".join(_box(code in picked, label) for code, label in choices)


def _form_date(value):
    return f"{datetime.date.fromisoformat(value):%B %d, %Y}" if value else ""


def _with_gender(person):
    return ", ".join(v for v in (person["name"], (person.get("gender") or "").capitalize()) if v)


def _money(n):
    return f"{n:,.2f}" if n else "-"


def _section(title, body, marked=None):
    # "split": the PDF gives each line its own row so a section longer than a page can break across pages
    # "marked": the same text with <b>/<i>/<u> (richtext.rich_text_marks) for the PDF and Word; CSV and Excel use the body
    block = {"kind": "table", "header": [[title]], "rows": [[body or "Not provided"]], "widths": [100], "header_fill": PROPOSAL_FILL,
             "split": True}
    if marked:
        block["marked"] = marked
    return block


def _rich_section(title, value):
    return _section(title, rich_text_lines(value), rich_text_marks(value))


def _work_plan_block(data):
    """Section XI as the preview's Gantt chart: year 1 = calendar year of the earliest date, one column per month."""
    dates = [d for d in [data["start_date"], data["target_end_date"]] + [x for w in data["work_plan"] for x in (w["start_date"], w["target_date"])] if d]
    first = min(int(d[:4]) for d in dates) if dates else datetime.date.today().year
    years = max(1, (max(int(d[:4]) for d in dates) if dates else first) - first + 1)
    months = years * 12

    def index(d):
        return (int(d[:4]) - first) * 12 + int(d[5:7]) - 1

    def month(d):
        return MONTHS[int(d[5:7]) - 1] if d else ""

    width = 3 + months
    header = [
        ["XI. WORK PLAN"] + [""] * (width - 1),
        ["", "", ""] + [f"{q // 4 + 1}Q{q % 4 + 1}" if m == 0 else "" for q in range(years * 4) for m in range(3)],
        ["Workplan", "Start", "End"] + [MONTHS[m % 12] if years == 1 else MONTHS[m % 12][0] for m in range(months)],
    ]
    spans = [(0, 0, width - 1)] + [(1, 3 + q * 3, 5 + q * 3) for q in range(years * 4)]
    rows, fills = [], []
    for i, w in enumerate(data["work_plan"]):
        start, end = w["start_date"] or w["target_date"], w["target_date"] or w["start_date"]
        rows.append([f"{i + 1}. {w['title']}", month(w["start_date"]), month(w["target_date"])] + [""] * months)
        if start and end:
            fills += [(i, 3 + m, GANTT_FILL) for m in range(index(start), index(end) + 1)]
    if not rows:
        rows = [["Set under Work Plan after registration, or through the Excel upload."] + [""] * (width - 1)]
    return {"kind": "table", "header": header, "rows": rows, "spans": spans, "widths": [26, 8, 8] + [3] * months,
            "header_fill": PROPOSAL_FILL, "cell_fills": fills, "bold_first": True, "narrow_from": 3}


def proposal_form(project_id):
    """The Research Proposal Form (LSPU-RDO-SF-018) of a registered project, laid out like the wizard's Research
    Proposal Form Preview (client request 2026-10-09). Built from importer.proposal_preview, the same data the
    preview shows."""
    from research_projects.importer import proposal_preview
    from research_projects.models import Project

    project = Project.objects.filter(pk=project_id).select_related("lead").first()
    if project is None:
        return None
    data = proposal_preview(project)
    total_cost = f"PhP {float(data['total_cost']):,.2f}" if data["total_cost"] else "—"
    duration = " – ".join(v for v in (_form_date(data["start_date"]), _form_date(data["target_end_date"])) if v)
    team = [[f"Team Member {i}", _with_gender(m)] for i, m in enumerate(data["team"], start=1)] or [["Team Member", "None listed"]]
    sector_other = f": {data['sector_other']}" if data["sector_other"] else ""
    continuing = f"Continuing (Year {data['continuing_year']})" if data["is_continuing"] and data["continuing_year"] else "Continuing (i.e., Year 2, Year 3, and so on…)"

    # Allocation columns (Q1..Q4, or Y1Q1..YnQ4 for a multi-year project; client 2026-10-14/10-17) only when the LIB has a split
    periods = data["allocation_periods"] if any(b["allocations"] for b in data["budget"]) else []

    def split(items):
        sums = {p: 0.0 for p in periods}
        for b in items:
            for a in b["allocations"]:
                if a["period"] in sums:
                    sums[a["period"]] += a["amount"]
        return [_money(sums[p]) for p in periods]

    blank = [""] * len(periods)
    budget_rows, row_fills, bold_rows = [], {}, []
    for key, label in BUDGET_GROUPS:
        items = [b for b in data["budget"] if b["category"] == key]
        if not items:
            continue
        row_fills[len(budget_rows)] = PROPOSAL_FILL
        bold_rows.append(len(budget_rows))
        budget_rows.append([label, "", "", "", *blank, ""])
        for b in items:
            budget_rows.append([f"    {b['description']}", b["unit"], f"{b['quantity']:g}" if b["quantity"] is not None else "",
                                _money(b["unit_cost"]) if b["unit_cost"] is not None else "", *split([b]), _money(b["total"])])
        row_fills[len(budget_rows)] = "E2E8F0"
        bold_rows.append(len(budget_rows))
        budget_rows.append(["SUBTOTAL", "", "", "", *split(items), _money(sum(b["total"] for b in items))])
    if budget_rows:
        row_fills[len(budget_rows)] = "FECACA"
        bold_rows.append(len(budget_rows))
        budget_rows.append(["GRAND TOTAL", "", "", "", *split(data["budget"]), _money(sum(b["total"] for b in data["budget"]))])
    else:
        budget_rows = [["No LIB entered.", "", "", "", *blank, ""]]
    budget_columns = 5 + len(periods)
    # Y1Q1..YnQ4 share the room left after the fixed columns, labelled like the work plan (1Q1, 1Q2, ...)
    fixed_widths = [30, 7, 6, 11] if len(periods) <= 4 else [18, 5, 4, 8]
    period_width = (90 - sum(fixed_widths)) / len(periods) if periods else 0
    period_labels = [p[1:] if p[0] == "Y" else p for p in periods]

    outputs = [[FORM_6P_LABELS.get(o["category"], o["category"]), o["description"], str(o["target_count"])] for o in data["outputs"]]
    beneficiaries = [[b["group"] + (f" – {b['description']}" if b["description"] else ""), str(b["total"])] for b in data["beneficiaries"]]
    endorsers = [row for e in data["endorsers"] for row in (
        ["Name", e["name"].upper()],
        ["Designation", (e["designation"] or "—") + (f" · Date Signed: {_form_date(e['signed_on'])}" if e["signed_on"] else "")],
    )]
    submitted = f" · {_form_date(data['proposal_submitted_on'])}" if data["proposal_submitted_on"] else ""

    blocks = [
        {"kind": "header", "lines": UNIVERSITY, "title": "RESEARCH PROPOSAL FORM (LSPU-FUNDED RESEARCH)"},
        {
            "kind": "table",
            "header": [["I. PROJECT/STUDY DETAILS", ""]],
            "spans": [(0, 0, 1)],
            "header_fill": PROPOSAL_FILL,
            "bold_first": True,
            "rows": [
                ["TITLE", data["title"] or "—"],
                ["PROJECT LEADER/GENDER", _with_gender({"name": data["lead_name"], "gender": data["lead_gender"]}) or "—"],
                ["CO-PROJECT LEADER/GENDER", "; ".join(map(_with_gender, data["co_leaders"])) or "—"],
                *team,
                ["DURATION", duration or "—"],
                ["START DATE", _form_date(data["start_date"]) or "—"],
                ["END DATE", _form_date(data["target_end_date"]) or "—"],
                ["TOTAL PROJECT/STUDY COST", total_cost],
                ["IMPLEMENTING UNIT", data["implementing_unit"] or "—"],
                ["CAMPUS", f"{data['campus']} Campus" if data["campus"] else "—"],
                ["CONTACT NO/S.", data["contact_number"] or "—"],
                ["E-MAIL ADDRESS", data["lead_email"] or "—"],
                ["COOPERATING AGENCY/IES", data["cooperating_agencies"]],
            ],
            "widths": [30, 70],
        },
        _section("SECTOR", _boxes(Project.SECTOR_CHOICES, data["sectors"]) + sector_other),
        _section("RESEARCH PROPOSAL CLASSIFICATION", "\n".join([
            f"{_box(not data['is_continuing'], 'New Proposal')}    {_box(data['is_continuing'], continuing)}",
            f"{_box(False, 'Program*')}    {_box(True, 'Project**')}    {_box(False, 'Study')}",
            # "Not Applicable" (client 2026-10-23) has no box on the form; it leaves both pairs unchecked
            _boxes([c for c in Project.RESEARCH_TYPE_CHOICES if c[0] != "not_applicable"], [data["research_type"]]),
            f"{_box(data['is_dry_research'] is False, 'Wet Research (i.e., with laboratory)')}    {_box(data['is_dry_research'] is True, 'Dry Research')}",
        ])),
        _section("STUDY COMPONENT TITLES", "\n".join(f"Study {i}: {t}" for i, t in enumerate(data["study_titles"] or ["", ""], start=1))),
        _section("RDE AGENDA", f"RDE THRUST: {data['rde_thrust'] or '—'}\nRDE PRIORITY AREA: {data['rde_priority'] or '—'}"),
        _section("RESEARCH PRIORITY AREA", "\n".join(_box(code == data["research_priority_area"], label) for code, label in Project.PRIORITY_AREA_CHOICES)),
        _section("RESEARCH TYPOLOGY", "\n".join(_box(code in data["research_typology"], label) for code, label in Project.TYPOLOGY_CHOICES)),
        _section("SUSTAINABLE DEVELOPMENT GOALS (SDGs)", "; ".join(f"SDG {n} — {SDG_LABELS[n - 1]}" for n in data["sdgs"]) or "None selected"),
        _rich_section("II. BACKGROUND OF THE STUDY", data["background"]),
        _section("III. OBJECTIVES OF THE STUDY", "\n".join(f"{i}. {o}" for i, o in enumerate(data["objectives"], start=1))),
        _rich_section("IV. PROJECT DESCRIPTIONS/METHODOLOGY", data["methodology"]),
        {
            "kind": "table",
            "header": [["V. QUANTIFIABLE EXPECTED OUTPUTS: 6Ps", "", ""], ["ITEM", "PARTICULARS", "QUANTITY"]],
            "spans": [(0, 0, 2)],
            "header_fill": PROPOSAL_FILL,
            "rows": outputs or [["Set under Research Outputs after registration, or through the Excel upload.", "", ""]],
            "widths": [25, 60, 15],
        },
        _rich_section("VI. SOCIO-ECONOMIC SIGNIFICANCE", data["socio_economic_significance"]),
        {
            "kind": "table",
            "header": [["VII. TARGET BENEFICIARIES", ""], ["Target Beneficiaries", "Total"]],
            "spans": [(0, 0, 1)],
            "header_fill": PROPOSAL_FILL,
            "rows": beneficiaries or [["None listed", ""]],
            "widths": [85, 15],
        },
        _rich_section("VIII. MONITORING/EVALUATION", data["monitoring_evaluation"]),
        _rich_section("IX. LIST OF REFERENCES", data["references"]),
        {
            "kind": "table",
            "header": [["X. BUDGET REQUIREMENTS"] + [""] * (budget_columns - 1),
                       ["DESCRIPTION", "UNIT", "QTY", "UNIT COST", *period_labels, "TOTAL"]],
            "spans": [(0, 0, budget_columns - 1)],
            "header_fill": PROPOSAL_FILL,
            "rows": budget_rows,
            "row_fills": row_fills,
            "bold_rows": bold_rows,
            "bold_first": False,
            "widths": [*fixed_widths, *[period_width] * len(periods), 10] if periods else [45, 10, 8, 14, 14],
        },
        {
            "kind": "table",
            "header": [["LIB APPROVAL", ""]],
            "spans": [(0, 0, 1)],
            "header_fill": PROPOSAL_FILL,
            "rows": [["Endorsed by (CRC Chair)", data.get("lib_endorsed_by") or "—"],
                     ["Certified by (Budget Officer)", data.get("lib_certified_by") or "—"]],
            "widths": [35, 65],
            "bold_first": False,
        },
        _work_plan_block(data),
        {"kind": "page_break"},
        {"kind": "header", "lines": ["Annex A"], "title": "ENDORSEMENT PAGE"},
        {
            "kind": "table",
            "header": [["SUBMITTED BY:", ""]],
            "spans": [(0, 0, 1)],
            "header_fill": PROPOSAL_FILL,
            "rows": [["Name", f"{(data['lead_name'] or '—').upper()} · Project Leader{submitted}"]],
            "widths": [25, 75],
            "bold_first": False,
        },
        {
            "kind": "table",
            "header": [["ENDORSED, NOTED, RECOMMENDED AND APPROVED BY (in order):", ""]],
            "spans": [(0, 0, 1)],
            "header_fill": PROPOSAL_FILL,
            "rows": endorsers or [["No endorsers entered.", ""]],
            "widths": [25, 75],
            "bold_first": False,
        },
    ]
    if data["college"]:
        blocks.append({"kind": "line", "text": f"College Unit: {data['college']}", "italic": True})
    return {"title": "Research Proposal Form", "landscape": False, "blocks": blocks}


# ---- Project staff Monthly Accomplishment Report -------------------------------------------------------------------

def month_bounds(month):
    """[start, end) datetimes of the month containing the date `month`."""
    from django.utils import timezone

    start = timezone.make_aware(datetime.datetime(month.year, month.month, 1))
    nxt = datetime.date(month.year + month.month // 12, month.month % 12 + 1, 1)
    return start, timezone.make_aware(datetime.datetime(nxt.year, nxt.month, 1))


def accomplishment_rows(user, month, projects=None):
    """Client follow-up 2026-10-02: instead of the staff recalling what they did, the report is pulled from their
    tasks: % completed (latest progress by month end), hours logged that month and what they noted. Covers tasks with
    an update that month or still open at month end, plus tasks since reassigned that they logged work on that month
    (client 2026-10-23). `projects` limits it to the requester's scope. Shared by the
    downloaded form and the Accomplishment Report page (client 2026-10-21), so both show the same numbers."""
    from django.db.models import Q, Sum

    from personnel.models import Task

    start, end = month_bounds(month)
    tasks = (
        Task.objects.filter(created_at__lt=end)
        .filter(Q(assignee=user) & (Q(updates__created_at__gte=start, updates__created_at__lt=end) | ~Q(status="done")
                                    | Q(completed_at__gte=start, completed_at__lt=end))
                | Q(updates__author=user, updates__kind="update", updates__created_at__gte=start, updates__created_at__lt=end))
        .select_related("project__lead").distinct().order_by("project__project_code", "due_date", "id")
    )
    if projects is not None:
        tasks = tasks.filter(project__in=projects)
    rows = []
    for task in tasks:
        month_updates = task.updates.filter(created_at__gte=start, created_at__lt=end).order_by("created_at", "id")
        rows.append({
            "task": task,
            "hours": month_updates.filter(author=user).aggregate(total=Sum("hours"))["total"] or 0,
            "progress": task.progress_as_of(end),
            "notes": "; ".join(u.note for u in month_updates if u.note),
            "accomplished": task.assignee_id == user.pk and task.status == "done" and task.completed_at is not None
            and start <= task.completed_at < end,
        })
    return rows


def accomplishment_form(user, month, projects=None):
    rows, total_hours, leaders = [], 0, []
    for i, row in enumerate(accomplishment_rows(user, month, projects), start=1):
        task, hours = row["task"], row["hours"]
        total_hours += hours
        rows.append([
            str(i), task.project.project_code, task.title, f"{row['progress']:g}%", f"{hours:g}",
            f"{task.due_date:%b %d, %Y}" if task.due_date else "", task.get_status_display(), row["notes"],
        ])
        if task.project.lead not in leaders:
            leaders.append(task.project.lead)
    if rows:
        rows.append(["", "", "TOTAL HOURS", "", f"{total_hours:g}", "", "", ""])
    position = " · ".join(filter(None, [getattr(user, "position", ""), user.role.name if user.role else ""]))
    return {
        "title": f"Monthly Accomplishment Report - {month:%B %Y}",
        "landscape": True,
        "blocks": [
            {"kind": "header", "lines": UNIVERSITY, "title": "MONTHLY ACCOMPLISHMENT REPORT"},
            {"kind": "line", "text": f"Name: {_name(user)}" + (f" ({position})" if position else "")},
            {"kind": "line", "text": f"Period Covered: {month:%B %Y}"},
            {
                "kind": "table",
                "header": [["", "Project", "Task / Activity", "% Completed", "Hours (this month)", "Deadline", "Status",
                            "Accomplishments / Remarks"]],
                "rows": rows or [["", "", "No task activity this month.", "", "", "", "", ""]],
                "widths": [4, 12, 26, 9, 9, 11, 9, 30],
            },
            {"kind": "line", "text": "Pulled from the RMIS task updates; % completed is the latest progress reported by the end "
                                     "of the month.", "italic": True},
            {"kind": "signatures", "items": [
                ("Submitted by:", _name(user), "Project Staff"),
                ("Noted by:", "; ".join(_name(lead) for lead in leaders), "Project Leader"),
            ]},
        ],
    }


# ---- R&D Quarterly Progress Report (client 2026-10-14, prototype QuarterlyReports.tsx export) -----------------------

def _quarterly_blocks(report):
    from monitoring import quarterly

    project = report.project
    fin = report.financial_snapshot or {}
    start, end = quarterly.quarter_range(report.year, report.quarter)
    reviews = [e for e in report.events.select_related("actor").order_by("at", "id")
               if e.action == "status_transition" and e.remarks and e.to_status != "submitted"]
    return [
        {"kind": "header", "lines": UNIVERSITY, "title": "R&D QUARTERLY PROGRESS REPORT"},
        {"kind": "line", "text": f"Project: {project.title} ({project.project_code})"},
        {"kind": "line", "text": f"Period Covered: Q{report.quarter} {report.year} ({start:%B %d, %Y} to {end:%B %d, %Y})"
                                 f" · Version {report.version}"},
        {
            "kind": "table",
            "header": [],
            "rows": [
                ["Project Leader", _name(project.lead), "Funding Source", quarterly.funding_agency(project)],
                ["Campus", project.campus, "College", project.college],
                ["Approved Duration", " to ".join(f"{d:%B %d, %Y}" if d else "-" for d in (project.start_date, project.target_end_date)),
                 "Original Budget", f"PHP {fin.get('original_budget', 0):,.2f}"],
                ["Quarter Disbursement", f"PHP {fin.get('quarter_disbursements', 0):,.2f}",
                 "Budget Utilization", f"{fin.get('utilization_pct', 0)}%"],
            ],
            "widths": [18, 32, 18, 32],
        },
        {"kind": "heading", "text": "Quarterly Accomplishments", "level": 2},
        {
            "kind": "table",
            "header": [["Approved Objective", "Planned Activity / Milestone", "Target", "Actual this Quarter", "Cumulative",
                        "Target Date", "Actual Date", "Completion %", "Variance", "Remarks / Delay Explanation"]],
            "rows": [[r.get("objective", ""), r.get("activity", ""), r.get("target", ""), r.get("actual", ""),
                      r.get("cumulative", ""), r.get("target_date", ""), r.get("actual_date", ""),
                      f"{r.get('completion_pct', 0)}%", r.get("variance", ""), r.get("remarks", "")]
                     for r in report.accomplishments] or [["No accomplishment rows.", "", "", "", "", "", "", "", "", ""]],
            "widths": [13, 12, 11, 13, 10, 8, 8, 9, 6, 10],
        },
        {"kind": "heading", "text": "Issues and Corrective Actions", "level": 2},
        {
            "kind": "table",
            "header": [["Issue", "Corrective Action", "Recommendation", "Responsible", "Target Resolution Date"]],
            "rows": [[i.get("issue", ""), i.get("corrective_action", ""), i.get("recommendation", ""),
                      i.get("responsible", ""), i.get("target_resolution_date", "")]
                     for i in report.issues] or [["None reported.", "", "", "", ""]],
            "widths": [24, 24, 22, 16, 14],
        },
        {"kind": "section", "heading": "Evidence and RMIS Document Links",
         "body": "\n".join(d.file_name for d in report.evidence.all()) or "None linked.", "level": 2},
        {"kind": "section", "heading": "Overall Recommendations", "body": report.recommendations or "None.", "level": 2},
        {"kind": "section", "heading": "Review History", "level": 2,
         "body": "\n".join(f"{_name(e.actor)} ({e.at:%b %d, %Y}): {e.remarks}" for e in reviews) or "No reviewer comments."},
        {"kind": "signatures", "items": [
            ("Prepared by:", _name(project.lead), "Project Leader"),
            ("Approved by:", _name(report.approved_by), f"{report.approved_at:%B %d, %Y}" if report.approved_at else ""),
        ]},
    ]


def quarterly_reports_form(reports):
    """Approved R&D Quarterly Progress Reports, one per page, or None when none were given."""
    blocks = []
    for report in reports:
        if blocks:
            blocks.append({"kind": "page_break"})
        blocks.extend(_quarterly_blocks(report))
    return {"title": "R&D Quarterly Progress Reports", "landscape": True, "blocks": blocks} if blocks else None


# ---- Renderers -----------------------------------------------------------------------------------------------------

def _table_width(block):
    return len(block["rows"][0]) if block["rows"] else len(block["header"][-1]) if block["header"] else 2


def _cell_style(block, r, c):
    """(fill hex or None, bold) of body cell r, c."""
    fill = next((h for fr, fc, h in block.get("cell_fills", ()) if fr == r and fc == c), None) or block.get("row_fills", {}).get(r)
    bold = r in block.get("bold_rows", ()) or (c == 0 and block.get("bold_first", not block["header"]))
    return fill, bold


def _flat_rows(form):
    """The form as plain rows, top to bottom; CSV and Excel both write these."""
    out = []
    for block in form["blocks"]:
        kind = block["kind"]
        if kind == "header":
            out += [[line] for line in block["lines"]] + [[], [block["title"]], []]
        elif kind == "line":
            out.append([block["text"]])
        elif kind == "table":
            out += [list(r) for r in block["header"]] + [list(r) for r in block["rows"]] + [[]]
        elif kind == "heading":
            out += [[], [block["text"]]]
        elif kind == "section":
            out.append([("    " if block["level"] > 2 else "") + block["heading"], block["body"]])
        elif kind == "signatures":
            out += [[]] + [[label, name, role] for label, name, role in block["items"]]
        elif kind == "footer":
            out += [[], block["parts"]]
        elif kind == "page_break":
            out += [[], []]
    return out


def render_form_csv(form):
    buf = io.StringIO()
    csv.writer(buf).writerows(_flat_rows(form))
    return buf.getvalue().encode("utf-8")


def render_form_xlsx(form):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", "-", form["title"])[:31]
    thin = Side(style="thin", color="000000")
    grid = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill = PatternFill("solid", fgColor=HEADER_FILL)
    wrap = Alignment(wrap_text=True, vertical="top")
    widths = {}
    row = 1
    span = max([_table_width(b) for b in form["blocks"] if b["kind"] == "table"] or [2])
    if form["landscape"]:
        ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0  # as many pages as needed, only the width is fitted
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    def full(r, value, center=False, **style):
        """A line across the whole form width (header, title, notes)."""
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=span)
        return put(r, 1, value, alignment=Alignment(horizontal="center" if center else "left", wrap_text=True), **style)

    def put(r, c, value, **style):
        cell = ws.cell(row=r, column=c, value=value)
        cell.alignment = style.pop("alignment", wrap)
        for k, v in style.items():
            setattr(cell, k, v)
        return cell

    for block in form["blocks"]:
        kind = block["kind"]
        if kind == "header":
            for line in block["lines"]:
                full(row, line, center=True, font=Font(bold=line == UNIVERSITY[1]))
                row += 1
            full(row + 1, block["title"], center=True, font=Font(bold=True, size=12))
            row += 3
        elif kind == "line":
            italic = block.get("italic", False)
            full(row, block["text"], font=Font(italic=italic, bold=not italic))
            ws.row_dimensions[row].height = 15 * (1 + len(block["text"]) // 90)
            row += 1
        elif kind == "table":
            width = _table_width(block)
            for i, w in enumerate(block.get("widths", []) if width > 1 else []):
                widths[i + 1] = max(widths.get(i + 1, 0), w)
            for h, header in enumerate(block["header"]):
                if width == 1:
                    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
                for c in range(width):
                    put(row, c + 1, header[c] or None, font=Font(bold=True), border=grid,
                        fill=PatternFill("solid", fgColor=block["header_fill"]) if "header_fill" in block else fill,
                        alignment=Alignment(horizontal="center", vertical="center", wrap_text=True))
                for span_row, first, last in block.get("spans", []):
                    if span_row == h:
                        ws.merge_cells(start_row=row, start_column=first + 1, end_row=row, end_column=last + 1)
                row += 1
            for r, values in enumerate(block["rows"]):
                if width == 1:  # a form section: one box across the whole form
                    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
                    ws.row_dimensions[row].height = 15 * sum(1 + len(line) // 120 for line in (values[0] or "").split("\n"))
                for c, value in enumerate(values):
                    cell_fill, bold = _cell_style(block, r, c)
                    put(row, c + 1, value, border=grid, font=Font(bold=bold),
                        **({"fill": PatternFill("solid", fgColor=cell_fill)} if cell_fill else {}))
                row += 1
            row += 1
        elif kind == "heading":
            row += 1
            full(row, block["text"], font=Font(bold=True, size=12))
            row += 1
        elif kind == "section":
            put(row, 1, ("    " if block["level"] > 2 else "") + block["heading"], font=Font(bold=True))
            put(row, 2, block["body"] or None)
            row += 1
        elif kind == "signatures":
            row += 1
            half = max(span // 2, 1)
            for c, (label, name, role) in enumerate(block["items"]):
                first = 1 + c * half
                last = span if c == len(block["items"]) - 1 else first + half - 1
                for r, value, font in ((row, label, Font()), (row + 2, name or "________________________", Font(bold=True)),
                                       (row + 3, role, Font())):
                    ws.merge_cells(start_row=r, start_column=first, end_row=r, end_column=last)
                    put(r, first, value, font=font)
            row += 5
        elif kind == "footer":
            full(row, "        ".join(block["parts"]), font=Font(size=8))
            row += 1
        elif kind == "page_break":
            row += 2
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _shade(cell, color=HEADER_FILL):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


MARK_TAG = re.compile(r"(</?[biu]>)")


def _add_marked_runs(paragraph, marked):
    """Lines with <b>/<i>/<u> (richtext.rich_text_marks) as Word runs, a line break between lines."""
    on = {"b": 0, "i": 0, "u": 0}
    for n, line in enumerate(marked.split("\n")):
        if n:
            paragraph.add_run().add_break()
        for part in MARK_TAG.split(line):
            if MARK_TAG.fullmatch(part):
                on[part.strip("</>")] += -1 if part[1] == "/" else 1
            elif part:
                run = paragraph.add_run(unescape(part))
                run.bold, run.italic, run.underline = on["b"] > 0 or None, on["i"] > 0 or None, on["u"] > 0 or None


def render_form_docx(form):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(11)
    section = doc.sections[0]
    if form["landscape"]:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width

    def para(text="", bold=False, italic=False, size=None, center=False):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold, run.italic = bold, italic
        if size:
            run.font.size = Pt(size)
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return p

    for block in form["blocks"]:
        kind = block["kind"]
        if kind == "header":
            for line in block["lines"]:
                para(line, bold=line == UNIVERSITY[1], center=True)
            para(block["title"], bold=True, size=13, center=True)
        elif kind == "line":
            para(block["text"], italic=block.get("italic", False), bold=not block.get("italic", False))
        elif kind == "table":
            width = _table_width(block)
            table = doc.add_table(rows=0, cols=width)
            table.style = "Table Grid"
            table.autofit = False
            usable = section.page_width - section.left_margin - section.right_margin
            weights = block.get("widths") or [1] * width
            col_widths = [int(usable * w / sum(weights)) for w in weights]
            for h, header in enumerate(block["header"]):
                cells = table.add_row().cells
                for c in range(width):
                    cells[c].text = header[c]
                    _shade(cells[c], block.get("header_fill", HEADER_FILL))
                    for run in cells[c].paragraphs[0].runs:
                        run.bold = True
                    cells[c].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                for span_row, first, last in block.get("spans", []):
                    if span_row == h:
                        merged = cells[first].merge(cells[last])
                        merged.text = header[first]
                        for run in merged.paragraphs[0].runs:
                            run.bold = True
                        merged.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r, values in enumerate(block["rows"]):
                cells = table.add_row().cells
                for c, value in enumerate(values):
                    if block.get("marked"):
                        _add_marked_runs(cells[c].paragraphs[0], block["marked"])
                    else:
                        cells[c].text = value
                    cell_fill, bold = _cell_style(block, r, c)
                    if cell_fill:
                        _shade(cells[c], cell_fill)
                    if bold:
                        for run in cells[c].paragraphs[0].runs:
                            run.bold = True
            for c, column in enumerate(table.columns):  # LibreOffice reads the grid, Word reads each cell
                column.width = col_widths[c]
            for r in table.rows:
                for c, cell in enumerate(r.cells):
                    cell.width = col_widths[c]
            para()
        elif kind == "heading":
            para(block["text"], bold=True, size=13)
        elif kind == "section":
            heading = doc.add_heading(block["heading"], level=block["level"])
            for run in heading.runs:  # the form's headings are black, not Word's default blue
                run.font.color.rgb = RGBColor(0, 0, 0)
            if block["body"]:
                for chunk in block["body"].split("\n"):
                    para(chunk)
            else:
                para("\n\n")  # space to write in
        elif kind == "signatures":
            table = doc.add_table(rows=4, cols=len(block["items"]))
            for c, (label, name, role) in enumerate(block["items"]):
                table.cell(0, c).text = label
                table.cell(2, c).text = name or "________________________"
                if table.cell(2, c).paragraphs[0].runs:
                    table.cell(2, c).paragraphs[0].runs[0].bold = True
                table.cell(3, c).text = role
        elif kind == "footer":
            footer = section.footer.paragraphs[0]
            footer.text = "\t".join(block["parts"])
        elif kind == "page_break":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_form_pdf(form):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    pagesize = landscape(A4) if form["landscape"] else A4
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=pagesize, leftMargin=40, rightMargin=40, topMargin=36, bottomMargin=40)
    base = getSampleStyleSheet()["Normal"]
    normal = ParagraphStyle("form", parent=base, fontName="Times-Roman", fontSize=10, leading=13)
    small = ParagraphStyle("cell", parent=normal, fontSize=9, leading=11)
    center = ParagraphStyle("center", parent=normal, alignment=TA_CENTER)
    h1 = ParagraphStyle("h1", parent=normal, fontName="Times-Bold", fontSize=12, spaceBefore=10, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=normal, fontName="Times-Bold", fontSize=11, spaceBefore=8, spaceAfter=3)
    h3 = ParagraphStyle("h3", parent=h2, fontName="Times-BoldItalic", leftIndent=18)
    footer_parts = []

    def p(text, style=normal):
        return Paragraph(escape(text or "").replace("\n", "<br/>"), style)

    elements = []
    for block in form["blocks"]:
        kind = block["kind"]
        if kind == "header":
            for line in block["lines"]:
                elements.append(Paragraph(f"<b>{escape(line)}</b>" if line == UNIVERSITY[1] else escape(line), center))
            elements += [Spacer(1, 10), Paragraph(f"<b>{escape(block['title'])}</b>", center), Spacer(1, 8)]
        elif kind == "line":
            text = escape(block["text"])
            elements.append(Paragraph(f"<i>{text}</i>" if block.get("italic") else f"<b>{text}</b>", normal))
            elements.append(Spacer(1, 4))
        elif kind == "table":
            width = _table_width(block)
            header = [[Paragraph(f"<b>{escape(v)}</b>", ParagraphStyle("th", parent=small, alignment=TA_CENTER))
                       for v in row] for row in block["header"]]
            def body_cell(r, c, v):
                if block.get("marked"):
                    return Paragraph(v, small)  # already escaped, only <b>/<i>/<u> left
                return Paragraph(f"<b>{escape(v or '')}</b>", small) if _cell_style(block, r, c)[1] else p(v, small)

            rows = block["rows"]
            if block.get("split"):
                # ReportLab can't break one table cell across pages (LayoutError, FRN-2026-005 2026-10-22)
                rows = [[line] for line in block.get("marked", rows[0][0]).split("\n")]
            body = [[body_cell(r, c, v) for c, v in enumerate(row)] for r, row in enumerate(rows)]
            widths = block.get("widths") or [100 / width] * width
            usable = pagesize[0] - 80
            table = Table(header + body or [[""] * width], colWidths=[usable * w / sum(widths) for w in widths],
                          repeatRows=len(header))
            if block.get("split"):
                # Lines of one section: outer box and header rule only, so it still reads as one cell
                style = [("BOX", (0, 0), (-1, -1), 0.5, colors.black), ("LINEBELOW", (0, len(header) - 1), (-1, len(header) - 1), 0.5, colors.black),
                         ("VALIGN", (0, 0), (-1, -1), "TOP")]
            else:
                style = [("GRID", (0, 0), (-1, -1), 0.5, colors.black), ("VALIGN", (0, 0), (-1, -1), "TOP")]
            if header:
                style.append(("BACKGROUND", (0, 0), (-1, len(header) - 1), colors.HexColor(f"#{block.get('header_fill', HEADER_FILL)}")))
            for r, values in enumerate(rows):
                for c in range(len(values)):
                    cell_fill = _cell_style(block, r, c)[0]
                    if cell_fill:
                        style.append(("BACKGROUND", (c, len(header) + r), (c, len(header) + r), colors.HexColor(f"#{cell_fill}")))
            if "narrow_from" in block:
                # Month columns of a multi-year Gantt are ~10pt wide, less than the default 6+6pt padding (FRN-2026-005)
                style += [("LEFTPADDING", (block["narrow_from"], 0), (-1, -1), 1), ("RIGHTPADDING", (block["narrow_from"], 0), (-1, -1), 1)]
            for span_row, first, last in block.get("spans", []):
                style.append(("SPAN", (first, span_row), (last, span_row)))
            table.setStyle(TableStyle(style))
            elements += [table, Spacer(1, 6)]
        elif kind == "heading":
            elements.append(Paragraph(escape(block["text"]), h1))
        elif kind == "section":
            elements.append(Paragraph(escape(block["heading"]), h3 if block["level"] > 2 else h2))
            elements.append(p(block["body"]) if block["body"] else Spacer(1, 48))
        elif kind == "signatures":
            items = block["items"]
            rows = [[label for label, _, _ in items], ["" for _ in items],
                    [name or "________________________" for _, name, _ in items], [role for _, _, role in items]]
            table = Table([[p(v) for v in row] for row in rows], colWidths=[(pagesize[0] - 80) / len(items)] * len(items))
            table.setStyle(TableStyle([("TOPPADDING", (0, 1), (-1, 1), 18)]))
            elements += [Spacer(1, 16), table]
        elif kind == "footer":
            footer_parts = block["parts"]
        elif kind == "page_break":
            elements.append(PageBreak())

    def draw_footer(canvas, _doc):
        if not footer_parts:
            return
        canvas.setFont("Helvetica", 8)
        left, middle, right = footer_parts
        canvas.drawString(40, 22, left)
        canvas.drawCentredString(pagesize[0] / 2, 22, middle)
        canvas.drawRightString(pagesize[0] - 40, 22, right)

    doc.build(elements, onFirstPage=draw_footer, onLaterPages=draw_footer)
    return buf.getvalue()


FORM_RENDERERS = {"csv": render_form_csv, "xlsx": render_form_xlsx, "pdf": render_form_pdf, "docx": render_form_docx}
