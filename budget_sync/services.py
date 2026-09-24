import re
from decimal import Decimal, InvalidOperation

import openpyxl
from django.db.models import Sum

from budget_lib.models import LineItem
from research_projects.models import Project

from .models import BudgetOfficeRecord

PROJECT_SHEET = re.compile(r"^P\d+$")
TOLERANCE = Decimal("0.01")


def _first_number(cells):
    for value in cells:
        if isinstance(value, (int, float)):
            return Decimal(str(round(value, 2)))
        if isinstance(value, str):
            try:
                return Decimal(value.replace(",", "").strip())
            except InvalidOperation:
                continue
    return None


def parse_sheet(ws):
    """Pull header fields and totals out of one Budget Office LIB sheet. The layout is the
    Budget Office's fixed template: 'PROJECT TITLE:', 'Project Study Leader:', 'Implementing Unit:',
    'Sub-total for MOOE', 'GRAND TOTAL' (no separate CO subtotal, so CO = grand total - MOOE)."""
    data = {"title": "", "leader_name": "", "implementing_unit": "", "mooe_total": None, "grand_total": None}
    for row in ws.iter_rows(values_only=True):
        cells = [v for v in row if v is not None and str(v).strip() != ""]
        if not cells or not isinstance(cells[0], str):
            continue
        label = re.sub(r"\s+", " ", cells[0]).strip().upper()
        rest = cells[1:]
        if label.startswith("PROJECT TITLE") and rest:
            data["title"] = str(rest[0]).strip()
        elif label.startswith("PROJECT STUDY LEADER") and rest:
            data["leader_name"] = str(rest[0]).strip()
        elif label.startswith("IMPLEMENTING UNIT") and rest:
            data["implementing_unit"] = str(rest[0]).strip()
        elif label.startswith("SUB-TOTAL FOR MOOE") and data["mooe_total"] is None:
            data["mooe_total"] = _first_number(rest)
        elif label == "GRAND TOTAL" and data["grand_total"] is None:
            data["grand_total"] = _first_number(rest)
    mooe = data["mooe_total"] or Decimal(0)
    grand = data["grand_total"] if data["grand_total"] is not None else mooe
    data["mooe_total"], data["grand_total"], data["co_total"] = mooe, grand, grand - mooe
    return data


def import_workbook(file_obj, source):
    """Create one BudgetOfficeRecord per filled-in project sheet (blank template sheets are skipped);
    auto-link by exact, case-insensitive title when exactly one RMIS project has it."""
    wb = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
    records = []
    for name in wb.sheetnames:
        if not PROJECT_SHEET.match(name):
            continue
        data = parse_sheet(wb[name])
        if not data["title"]:
            continue
        matches = list(Project.objects.filter(title__iexact=data["title"])[:2])
        project = matches[0] if len(matches) == 1 else None  # ambiguous titles are left for manual linking
        records.append(BudgetOfficeRecord(
            source=source, sheet_name=name, project=project, match_method="auto" if project else "", **data,
        ))
    BudgetOfficeRecord.objects.bulk_create(records)
    return len(records)


def _rmis_totals(project):
    budget = project.budgets.filter(is_current=True).first()
    if budget is None:
        return None
    by_category = dict(
        LineItem.objects.filter(budget=budget).values_list("category").annotate(total=Sum("amount"))
    )
    mooe = (by_category.get("mooe") or 0) + (by_category.get("ps") or 0)  # Budget Office files PS-type items under MOOE
    co = by_category.get("co") or 0
    return {"budget": budget.id, "mooe_total": mooe, "co_total": co, "grand_total": mooe + co}


def reconcile(source):
    """Compare each imported sheet against the linked project's current RMIS LIB."""
    rows = []
    for record in source.records.select_related("project").order_by("id"):
        row = {
            "record": record.id, "sheet_name": record.sheet_name, "title": record.title,
            "project": record.project_id, "match_method": record.match_method,
            "budget_office": {"mooe_total": record.mooe_total, "co_total": record.co_total, "grand_total": record.grand_total},
            "rmis": None, "difference": None,
        }
        if record.project is None:
            row["status"] = "unlinked"
        else:
            rmis = _rmis_totals(record.project)
            if rmis is None:
                row["status"] = "no_rmis_budget"
            else:
                diff = {k: rmis[k] - getattr(record, k) for k in ("mooe_total", "co_total", "grand_total")}
                row["rmis"], row["difference"] = rmis, diff
                row["status"] = "discrepancy" if any(abs(v) > TOLERANCE for v in diff.values()) else "matched"
        rows.append(row)
    summary = {s: sum(1 for r in rows if r["status"] == s) for s in ("matched", "discrepancy", "no_rmis_budget", "unlinked")}
    return {"import": source.id, "summary": summary, "records": rows}
