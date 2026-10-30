"""LIB review workflow (client 2026-10-21, replaces the 2026-10-18 endorse/return chain): the Project Leader submits,
the CRC Chair forwards it with Mark for checking (Task 19, 2026-10-30: one step, straight to Ready for Certification),
the Budget Officer certifies (Start Certification -> Certify). Either level can put it On Hold for the Project Leader, who edits and resubmits to the CRC
Chair. Every step is kept in the review history, each submission with a snapshot of its line items."""
import io
from datetime import date
from decimal import Decimal

import requests

from accounts.testing import RMISTestCase
from research_projects.models import Project

from .models import LineItem, LineItemBudget, LineItemEdit


def attach_approval_documents(project, user):
    """The Supporting Documents attachments (Approve Proposal, NTP, and the optional MOA); submitting the project's LIB needs
    the first two (client 2026-10-24, Task 13)."""
    from document_management.models import Document

    for document_type in ("notice_of_approval", "ntp", "moa"):
        Document.objects.create(project=project, document_type=document_type, version_number=1, storage_path=f"x/{document_type}",
                                file_name=f"{document_type}.pdf", file_size=1, uploaded_by=user)


def add_registration_rows(project):
    """The form's minimums that submitting the project also needs: 2 studies, a 6P, a beneficiary (client 2026-10-29)."""
    from outputs.models import ExpectedOutput
    from research_projects.models import Study, TargetBeneficiary

    Study.objects.bulk_create([Study(project=project, title="Study 1"), Study(project=project, title="Study 2")])
    ExpectedOutput.objects.create(project=project, category="publications", description="Journal article", target_count=1)
    TargetBeneficiary.objects.create(project=project, group="Faculty", description="CTE faculty", total=10)


class LibWorkflowTests(RMISTestCase):
    def setUp(self):
        super().setUp()
        self.leader = self.make_user("project_leader")
        self.crc = self.make_user("crc_chair")
        self.officer = self.make_user("finance_budget")
        self.project = Project.objects.create(
            title="P", project_code="P-1", funding_type="core_funded", lead=self.leader, campus="San Pablo",
            total_cost="1000", start_date=date(2026, 1, 1), target_end_date=date(2026, 12, 31),
        )
        self.budget = LineItemBudget.objects.create(project=self.project, version_number=1)
        self.item = LineItem.objects.create(budget=self.budget, category="mooe", description="Travel", amount=Decimal("1000"))
        attach_approval_documents(self.project, self.leader)
        add_registration_rows(self.project)

    def act(self, user, action, **data):
        return self.client_for(user).post(f"/api/budget/budgets/{self.budget.id}/{action}/", data, format="json")

    def status(self):
        self.budget.refresh_from_db()
        return self.budget.status

    def move_to(self, status):
        LineItemBudget.objects.filter(pk=self.budget.pk).update(status=status)

    def test_full_chain_from_draft_to_certified(self):
        steps = [(self.leader, "submit", "submitted"), (self.crc, "for-checking", "endorsed"),
                 (self.officer, "start-certification", "certifying"), (self.officer, "certify", "certified")]
        for user, action, expected in steps:
            response = self.act(user, action)
            self.assertEqual(response.status_code, 200, (action, response.data))
            self.assertEqual(self.status(), expected, action)
        self.assertEqual((self.budget.endorsed_by, self.budget.certified_by), (self.crc, self.officer))
        self.assertEqual([r["action"] for r in response.data["reviews"]],
                         ["submitted", "ready", "certification_started", "certified"])

    def test_steps_cannot_be_skipped(self):
        self.assertEqual(self.act(self.crc, "for-checking").status_code, 404)  # not submitted yet: a Draft is hidden from the CRC Chair (D2)
        self.act(self.leader, "submit")
        for action in ("start-certification", "certify"):
            self.assertEqual(self.act(self.officer, action).status_code, 404, action)  # Q10: hidden from Budget until Ready
        for action in ("start-review", "ready"):  # Task 19: the old CRC steps are gone
            self.assertEqual(self.act(self.crc, action).status_code, 404, action)
        self.assertEqual(self.status(), "submitted")
        self.act(self.crc, "for-checking")
        self.assertEqual(self.act(self.crc, "for-checking").status_code, 400)  # forwarded once; can't be repeated
        self.assertEqual(self.status(), "endorsed")
        self.assertEqual(self.act(self.officer, "certify").status_code, 400)  # Start Certification first (Q6)
        self.assertEqual(self.status(), "endorsed")

    def test_each_step_belongs_to_its_role(self):
        for user in (self.crc, self.officer, self.make_user("project_leader")):
            self.assertIn(self.act(user, "submit").status_code, (403, 404), user.role.code)
        self.act(self.leader, "submit")
        for user in (self.officer, self.leader, self.make_user("vprei"), self.make_user("riuh")):
            self.assertIn(self.act(user, "for-checking").status_code, (403, 404), user.role.code)  # Q10: hidden from Budget
        self.move_to("endorsed")
        for user in (self.crc, self.leader):
            self.assertEqual(self.act(user, "start-certification").status_code, 403, user.role.code)
        self.move_to("certifying")
        self.assertEqual(self.act(self.crc, "certify").status_code, 403)
        self.assertEqual(self.status(), "certifying")

    def test_mark_for_checking_forwards_older_libs_still_under_crc_review(self):
        for status in ("submitted", "crc_processing", "for_checking"):
            self.move_to(status)
            self.assertEqual(self.act(self.crc, "for-checking").status_code, 200, status)
            self.assertEqual(self.status(), "endorsed", status)
            self.budget.refresh_from_db()
            self.assertEqual(self.budget.endorsed_by, self.crc)

    def test_crc_chair_outside_campus_gets_404(self):
        self.act(self.leader, "submit")
        other = self.make_user("crc_chair", campus="Los Baños")
        self.assertEqual(self.act(other, "for-checking").status_code, 404)
        self.assertEqual(self.status(), "submitted")

    def test_submit_refuses_grand_total_not_equal_to_total_cost(self):
        LineItem.objects.filter(pk=self.item.pk).update(amount="900")
        response = self.act(self.leader, "submit")
        self.assertEqual(response.status_code, 400)
        self.assertIn("Total Project/Study Cost", str(response.data))
        self.assertEqual(self.status(), "draft")

    def test_submit_refuses_empty_lib(self):
        self.item.delete()
        self.assertEqual(self.act(self.leader, "submit").status_code, 400)

    def test_the_leader_edits_a_draft_but_not_a_submitted_lib(self):
        url = f"/api/budget/line-items/{self.item.id}/"
        leader = self.client_for(self.leader)
        self.assertEqual(leader.patch(url, {"justification": "Field visits"}, format="json").status_code, 200)
        new_row = {"budget": self.budget.id, "category": "mooe", "description": "Training Expenses", "amount": "10"}
        for status in ("submitted", "crc_processing", "for_checking", "endorsed", "certifying", "certified"):
            self.move_to(status)
            self.assertEqual(leader.patch(url, {"justification": "X"}, format="json").status_code, 400, status)
            self.assertEqual(leader.post("/api/budget/line-items/", new_row, format="json").status_code, 400, status)
            self.assertEqual(leader.delete(url).status_code, 400, status)

    def test_only_the_project_leader_encodes(self):
        # Client Q9 (2026-10-21): the Procurement Officer and Program Leaders no longer encode LIBs (Program Leaders merged
        # into the Project Leader 2026-10-27; vprei stands in for another non-encoding role).
        row = {"budget": self.budget.id, "category": "mooe", "description": "Training Expenses", "amount": "10"}
        url = f"/api/budget/line-items/{self.item.id}/"
        for code in ("procurement_officer_lib", "vprei", "crc_chair", "finance_budget", "system_admin"):
            client = self.client_for(self.make_user(code))
            self.assertEqual(client.post("/api/budget/line-items/", row, format="json").status_code, 403, code)
            self.assertEqual(client.patch(url, {"justification": "X"}, format="json").status_code, 403, code)
            self.assertEqual(client.delete(url).status_code, 403, code)
            self.assertEqual(client.post("/api/budget/budgets/", {"project": self.project.id}, format="json").status_code, 403, code)
        self.assertEqual(self.client_for(self.leader).post("/api/budget/line-items/", row, format="json").status_code, 201)

    def test_the_budget_officer_sees_a_lib_only_once_it_is_ready(self):
        # Client Q10 (2026-10-21): nothing before Ready for Certification reaches the Budget Officer.
        officer = self.client_for(self.officer)
        for status, seen in (("draft", False), ("submitted", False), ("crc_processing", False), ("for_checking", False),
                             ("returned", False), ("endorsed", True), ("certifying", True), ("certified", True)):
            self.move_to(status)
            listed = [b["id"] for b in officer.get("/api/budget/budgets/", {"project": self.project.id}).data]
            self.assertEqual(self.budget.id in listed, seen, status)
            self.assertEqual(officer.get(f"/api/budget/budgets/{self.budget.id}/").status_code, 200 if seen else 404, status)
            items = [i["id"] for i in officer.get("/api/budget/line-items/", {"budget": self.budget.id}).data]
            self.assertEqual(self.item.id in items, seen, status)

    def test_other_roles_still_see_every_status(self):
        self.move_to("for_checking")
        for user in (self.leader, self.crc, self.make_user("system_admin")):
            self.assertEqual(self.client_for(user).get(f"/api/budget/budgets/{self.budget.id}/").status_code, 200, user.role.code)

    def test_hold_needs_remarks(self):
        self.act(self.leader, "submit")
        self.assertEqual(self.act(self.crc, "hold").status_code, 400)
        self.assertEqual(self.status(), "submitted")

    def test_crc_chair_holds_from_its_review_states(self):
        for status in ("submitted", "crc_processing", "for_checking"):
            self.move_to(status)
            response = self.act(self.crc, "hold", remarks="Fix travel")
            self.assertEqual(response.status_code, 200, status)
            self.assertEqual((self.status(), response.data["hold_level"]), ("returned", "crc"), status)
        self.assertEqual(response.data["reviews"][-1]["action"], "held_crc")

    def test_budget_officer_holds_from_ready_and_in_progress(self):
        for status in ("endorsed", "certifying"):
            self.move_to(status)
            response = self.act(self.officer, "hold", remarks="Wrong unit cost")
            self.assertEqual(response.status_code, 200, status)
            self.assertEqual((self.status(), response.data["hold_level"]), ("returned", "budget"), status)

    def test_hold_belongs_to_the_level_holding_the_lib(self):
        self.move_to("submitted")
        self.assertEqual(self.act(self.officer, "hold", remarks="x").status_code, 404)  # Q10: not visible to Budget yet
        self.move_to("endorsed")
        self.assertEqual(self.act(self.crc, "hold", remarks="x").status_code, 403)
        for status, code in (("draft", 404), ("certified", 400), ("returned", 400)):  # a Draft is hidden from the CRC Chair (D2)
            self.move_to(status)
            self.assertEqual(self.act(self.crc, "hold", remarks="x").status_code, code, status)

    def test_the_leader_edits_an_on_hold_lib_and_resubmits_to_the_crc_chair(self):
        self.act(self.leader, "submit")
        self.move_to("endorsed")
        self.act(self.officer, "hold", remarks="Fix travel")
        edit = self.client_for(self.leader).patch(f"/api/budget/line-items/{self.item.id}/", {"justification": "Two trips"}, format="json")
        self.assertEqual(edit.status_code, 200, edit.data)

        response = self.act(self.leader, "resubmit", remarks="Fixed")

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(self.status(), "submitted")  # client Q7: back to Awaiting CRC Chair
        self.assertIsNone(response.data["hold_level"])
        self.assertEqual([(r["action"], r["remarks"]) for r in response.data["reviews"]],
                         [("submitted", ""), ("held_budget", "Fix travel"), ("resubmitted", "Fixed")])

    def test_only_the_project_leader_resubmits(self):
        self.move_to("returned")
        for user in (self.crc, self.officer, self.make_user("project_leader"), self.make_user("system_admin")):
            self.assertIn(self.act(user, "resubmit").status_code, (403, 404), user.role.code)
        self.assertEqual(self.status(), "returned")

    def test_resubmit_refuses_grand_total_not_equal_to_total_cost(self):
        self.move_to("returned")
        LineItem.objects.filter(pk=self.item.pk).update(amount="900")
        self.assertEqual(self.act(self.leader, "resubmit").status_code, 400)
        self.assertEqual(self.status(), "returned")

    def test_each_submission_keeps_a_snapshot_of_the_line_items(self):
        # Client Q11: what was submitted stays on record even after the leader edits an On Hold LIB.
        self.act(self.leader, "submit")
        self.act(self.crc, "hold", remarks="Fix")
        self.client_for(self.leader).patch(f"/api/budget/line-items/{self.item.id}/", {"justification": "Two trips"}, format="json")
        response = self.act(self.leader, "resubmit")

        submitted, held, resubmitted = response.data["reviews"]
        self.assertEqual(submitted["snapshot"][0]["justification"], "")
        self.assertEqual(resubmitted["snapshot"][0]["justification"], "Two trips")
        self.assertEqual((submitted["snapshot"][0]["amount"], held["snapshot"]), ("1000.00", []))

    def test_old_libs_keep_their_status_and_history_labels(self):
        self.move_to("returned")
        from .models import BudgetReview
        BudgetReview.objects.create(budget=self.budget, action="returned", remarks="Old return", acted_by=self.crc)

        data = self.client_for(self.leader).get(f"/api/budget/budgets/{self.budget.id}/").data

        self.assertEqual((data["status"], data["reviews"][0]["action"]), ("returned", "returned"))

    def test_a_line_item_cannot_move_out_of_a_locked_lib(self):
        # Audit 2026-10-21: PATCH {"budget": <a draft LIB>} only checked the new LIB, so an item could leave a
        # certified one.
        LineItemBudget.objects.filter(pk=self.budget.pk).update(status="certified")
        draft = LineItemBudget.objects.create(project=self.project, version_number=2)

        response = self.client_for(self.leader).patch(f"/api/budget/line-items/{self.item.id}/", {"budget": draft.id}, format="json")

        self.assertEqual(response.status_code, 400)
        self.item.refresh_from_db()
        self.assertEqual(self.item.budget_id, self.budget.id)

    def test_a_line_item_stays_in_its_own_lib_even_while_draft(self):
        other = LineItemBudget.objects.create(project=self.project, version_number=2)

        response = self.client_for(self.leader).patch(f"/api/budget/line-items/{self.item.id}/", {"budget": other.id}, format="json")

        self.assertIn("budget", response.data)
        self.item.refresh_from_db()
        self.assertEqual(self.item.budget_id, self.budget.id)


class LineItemSpecificationTests(RMISTestCase):
    """Task 12 (2026-10-29): an LIB row can name its specific item (SF-018 Section X, e.g. "PAPER, 80 gsm" under Office
    Supplies Expenses) beside the LIB Item it is picked from."""

    def setUp(self):
        super().setUp()
        self.leader = self.make_user("project_leader")
        self.project = Project.objects.create(
            title="P", project_code="P-1", funding_type="core_funded", lead=self.leader, campus="San Pablo",
            total_cost="1000", start_date=date(2026, 1, 1), target_end_date=date(2026, 12, 31),
        )
        self.budget = LineItemBudget.objects.create(project=self.project, version_number=1)
        self.item = LineItem.objects.create(budget=self.budget, category="mooe", description="Travel", amount=Decimal("1000"))
        attach_approval_documents(self.project, self.leader)
        add_registration_rows(self.project)

    act = LibWorkflowTests.act

    def test_a_row_keeps_its_specific_item_beside_the_lib_item(self):
        row = {"budget": self.budget.id, "category": "mooe", "description": "Office Supplies Expenses",
               "specification": "  PAPER, 80 gsm, A4  ", "amount": "10"}

        response = self.client_for(self.leader).post("/api/budget/line-items/", row, format="json")

        self.assertEqual(response.status_code, 201)
        self.assertEqual((response.data["description"], response.data["specification"]), ("Office Supplies Expenses", "PAPER, 80 gsm, A4"))

    def test_the_specific_item_is_optional_and_older_rows_have_none(self):
        row = {"budget": self.budget.id, "category": "mooe", "description": "Training Expenses", "amount": "10"}

        response = self.client_for(self.leader).post("/api/budget/line-items/", row, format="json")

        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data["specification"], "")
        self.assertEqual(LineItem.objects.get(pk=self.item.pk).specification, "")

    def test_the_lib_item_still_comes_from_lib_items(self):
        row = {"budget": self.budget.id, "category": "mooe", "description": "PAPER, 80 gsm", "specification": "PAPER, 80 gsm", "amount": "10"}

        response = self.client_for(self.leader).post("/api/budget/line-items/", row, format="json")

        self.assertEqual(response.status_code, 400)
        self.assertIn("description", response.data)

    def test_a_submission_snapshot_keeps_the_specific_item(self):
        LineItem.objects.filter(pk=self.item.pk).update(specification="Bus fare, Manila")

        submitted = self.act(self.leader, "submit").data["reviews"][0]

        self.assertEqual(submitted["snapshot"][0]["specification"], "Bus fare, Manila")


class SubmissionEmailTests(RMISTestCase):
    """Task 15 (2026-10-30): when the Project Leader submits (or resubmits) a proposal, every active CRC/CRD of the
    project's campus gets a Brevo e-mail with the leader, the title and the budget per project year."""

    def setUp(self):
        super().setUp()
        self.leader = self.make_user("project_leader", campus="Siniloan", first_name="Ana", last_name="Reyes")
        self.project = Project.objects.create(
            title="Bridging <Academia>", project_code="P-1", funding_type="core_funded", lead=self.leader, campus="Siniloan",
            total_cost="15000", start_date=date(2027, 1, 1), target_end_date=date(2028, 3, 31),
        )
        self.budget = LineItemBudget.objects.create(project=self.project, version_number=1)
        LineItem.objects.create(budget=self.budget, category="mooe", description="Training Expenses", amount=Decimal("10000"),
                                allocations=[{"period": "Y1Q1", "amount": "4000.00"}, {"period": "Y1Q4", "amount": "6000.00"}] +
                                            [{"period": p, "amount": "0.00"} for p in ("Y1Q2", "Y1Q3", "Y2Q1")])
        LineItem.objects.create(budget=self.budget, category="co", description="Office Equipment", amount=Decimal("5000"),
                                allocations=[{"period": "Y2Q1", "amount": "5000.00"}])
        attach_approval_documents(self.project, self.leader)
        add_registration_rows(self.project)
        self.crc = [self.make_user("crc_chair", campus="Siniloan"), self.make_user("crc_chair", campus="siniloan")]
        self.make_user("crc_chair", campus="Sta. Cruz")
        self.make_user("crc_chair", campus="Siniloan", is_active=False)
        self.make_user("crc_chair", campus="Siniloan", account_status="suspended")
        self.brevo_post.reset_mock()

    def step(self, name):
        with self.captureOnCommitCallbacks(execute=True):
            return self.client_for(self.leader).post(f"/api/budget/budgets/{self.budget.pk}/{name}/")

    def sent(self):
        return [call.kwargs["json"] for call in self.brevo_post.call_args_list]

    def test_submitting_e_mails_every_active_crc_crd_of_the_campus_with_the_budget_per_year(self):
        response = self.step("submit")

        self.assertEqual(response.status_code, 200, response.data)
        mails = self.sent()
        self.assertEqual(sorted(m["to"][0]["email"] for m in mails), sorted(u.email for u in self.crc))
        mail = mails[0]
        self.assertEqual(mail["subject"], "New Project Proposal Submitted — Bridging <Academia>")
        for text in ("Ana Reyes", self.leader.email, "Bridging &lt;Academia&gt;", "Year 1", "₱10,000.00", "Year 2", "₱5,000.00",
                     "₱15,000.00", "open your RMIS account to review"):
            self.assertIn(text, mail["htmlContent"])

    def test_no_e_mail_while_it_is_a_draft_or_when_a_submit_is_repeated(self):
        self.client_for(self.leader).patch(f"/api/budget/line-items/{LineItem.objects.first().pk}/", {"justification": "x"}, format="json")
        self.assertEqual(self.sent(), [])

        self.step("submit")
        repeated = self.step("submit")

        self.assertEqual(repeated.status_code, 400)
        self.assertEqual(len(self.sent()), 2)

    def test_a_failed_e_mail_keeps_the_submission(self):
        self.brevo_post.side_effect = requests.ConnectionError("Brevo down")

        response = self.step("submit")

        self.assertEqual(response.status_code, 200)
        self.budget.refresh_from_db()
        self.assertEqual(self.budget.status, "submitted")

    def test_a_resubmission_e_mails_the_campus_crc_crd_again(self):
        LineItemBudget.objects.filter(pk=self.budget.pk).update(status="returned")

        response = self.step("resubmit")

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(len(self.sent()), 2)
        self.assertEqual(self.sent()[0]["subject"], "Project Proposal Resubmitted — Bridging <Academia>")


class ForwardEmailTests(RMISTestCase):
    """Task 19 (2026-10-30): when the CRC/CRD marks a proposal for checking, every active Finance/Accounting/Budget of
    the project's campus gets a Brevo e-mail with the leader, the forwarding CRC/CRD and the budget per project year."""

    def setUp(self):
        super().setUp()
        self.leader = self.make_user("project_leader", campus="Siniloan", first_name="Ana", last_name="Reyes")
        self.crc = self.make_user("crc_chair", campus="Siniloan", first_name="Ben", last_name="Cruz")
        self.project = Project.objects.create(
            title="Bridging <Academia>", project_code="P-1", funding_type="core_funded", lead=self.leader, campus="Siniloan",
            total_cost="15000", start_date=date(2027, 1, 1), target_end_date=date(2028, 3, 31),
        )
        self.budget = LineItemBudget.objects.create(project=self.project, version_number=1, status="submitted")
        LineItem.objects.create(budget=self.budget, category="mooe", description="Training Expenses", amount=Decimal("10000"),
                                allocations=[{"period": "Y1Q1", "amount": "10000.00"}])
        LineItem.objects.create(budget=self.budget, category="co", description="Office Equipment", amount=Decimal("5000"),
                                allocations=[{"period": "Y2Q1", "amount": "5000.00"}])
        self.finance = [self.make_user("finance_budget", campus="Siniloan"), self.make_user("finance_budget", campus="SINILOAN")]
        self.make_user("finance_budget", campus="Sta. Cruz")
        self.make_user("finance_budget", campus="Siniloan", is_active=False)
        self.make_user("finance_budget", campus="Siniloan", account_status="suspended")
        self.brevo_post.reset_mock()

    def forward(self, user=None):
        with self.captureOnCommitCallbacks(execute=True):
            return self.client_for(user or self.crc).post(f"/api/budget/budgets/{self.budget.pk}/for-checking/")

    def sent(self):
        return [call.kwargs["json"] for call in self.brevo_post.call_args_list]

    def test_forwarding_e_mails_every_active_finance_of_the_campus(self):
        response = self.forward()

        self.assertEqual(response.status_code, 200, response.data)
        mails = self.sent()
        self.assertEqual(sorted(m["to"][0]["email"] for m in mails), sorted(u.email for u in self.finance))
        mail = mails[0]
        self.assertEqual(mail["subject"], "Project Proposal Ready for Budget Review — Bridging <Academia>")
        for text in ("reviewed by CRC/CRD", "Ana Reyes", self.leader.email, "Bridging &lt;Academia&gt;", "Ben Cruz",
                     "Year 1", "₱10,000.00", "Year 2", "₱5,000.00", "₱15,000.00", "budget review"):
            self.assertIn(text, mail["htmlContent"])

    def test_a_repeated_or_refused_forward_sends_nothing_more(self):
        self.forward()
        self.assertEqual(self.forward().status_code, 400)
        self.assertEqual(len(self.sent()), 2)
        self.brevo_post.reset_mock()
        LineItemBudget.objects.filter(pk=self.budget.pk).update(status="submitted")
        self.assertIn(self.forward(self.make_user("finance_budget", campus="Siniloan")).status_code, (403, 404))
        self.assertEqual(self.sent(), [])

    def test_a_failed_e_mail_keeps_the_forward(self):
        self.brevo_post.side_effect = requests.ConnectionError("Brevo down")

        self.assertEqual(self.forward().status_code, 200)
        self.budget.refresh_from_db()
        self.assertEqual(self.budget.status, "endorsed")

    def test_the_crc_crd_submit_e_mail_still_goes_to_the_crc_crd_only(self):
        from .notifications import campus_reviewers

        self.assertEqual(list(campus_reviewers(self.project)), [self.crc])


class LibDownloadTests(RMISTestCase):
    """Task 20 (2026-10-30): the LIB as saved, downloaded in Step 4's Download LIB Template layout. The uploaded
    workbook isn't kept, so manual entry, the LIB import and Upload via Excel all download from the saved items."""

    def setUp(self):
        super().setUp()
        self.leader = self.make_user("project_leader", campus="Siniloan")
        self.officer = self.make_user("finance_budget", campus="Siniloan")
        self.project = Project.objects.create(
            title="P", project_code="NTP-9", funding_type="core_funded", lead=self.leader, campus="Siniloan",
            total_cost="15500", start_date=date(2027, 1, 1), target_end_date=date(2028, 3, 31),
        )
        self.budget = LineItemBudget.objects.create(project=self.project, version_number=1, status="endorsed")
        LineItem.objects.create(budget=self.budget, category="mooe", description="Training Expenses", specification="Venue, 2 days",
                                unit="pax", quantity=Decimal("20"), unit_cost=Decimal("500.50"), amount=Decimal("10010"),
                                allocations=[{"period": "Y1Q1", "amount": "4010.00"}, {"period": "Y1Q4", "amount": "6000.00"}])
        LineItem.objects.create(budget=self.budget, category="co", description="Office Equipment", unit="unit",
                                quantity=Decimal("1"), unit_cost=Decimal("5490"), amount=Decimal("5490"),
                                allocations=[{"period": "Y2Q1", "amount": "5490.00"}])
        self.url = f"/api/budget/budgets/{self.budget.id}/export/"

    def download(self, user=None):
        return self.client_for(user or self.officer).get(self.url)

    def test_the_finance_downloads_the_saved_lib_in_the_template_layout(self):
        import openpyxl

        response = self.download()

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Disposition"], 'attachment; filename="LIB_NTP-9_v1.xlsx"')
        sheet = openpyxl.load_workbook(io.BytesIO(response.content))["LIB"]
        header = next(r for r in range(1, 20) if sheet.cell(row=r, column=1).value == "Particulars")
        self.assertEqual([c.value for c in sheet[header] if c.value],
                         ["Particulars", "Item", "Specification (optional)", "Unit", "Qty", "Unit Cost (₱)", "Total (₱)",
                          "Y1Q1", "Y1Q2", "Y1Q3", "Y1Q4", "Y2Q1"])
        rows = [[c.value for c in sheet[header + i]][:12] for i in (1, 2)]
        self.assertEqual(rows[0], ["MOOE", "Training Expenses", "Venue, 2 days", "pax", 20, 500.5, 10010, 4010, None, None, 6000, None])
        self.assertEqual(rows[1], ["CO", "Office Equipment", None, "unit", 1, 5490, 5490, None, None, None, None, 5490])
        self.assertIn("LEGEND", [sheet.cell(row=r, column=1).value for r in range(header + 1, sheet.max_row + 1)])

    def test_the_download_reads_back_as_the_same_lib(self):
        from research_projects.importer import read_lib

        rows, errors = read_lib(io.BytesIO(self.download().content), ["Y1Q1", "Y1Q2", "Y1Q3", "Y1Q4", "Y2Q1"])

        self.assertEqual(errors, [])
        self.assertEqual([(r["category"], r["description"], r["specification"], r["unit"], Decimal(r["amount"])) for r in rows],
                         [("mooe", "Training Expenses", "Venue, 2 days", "pax", Decimal("10010")),
                          ("co", "Office Equipment", "", "unit", Decimal("5490"))])
        self.assertEqual([(a["period"], Decimal(a["amount"])) for a in rows[0]["allocations"] if Decimal(a["amount"])],
                         [("Y1Q1", Decimal("4010")), ("Y1Q4", Decimal("6000"))])

    def test_who_may_download(self):
        self.assertEqual(self.download(self.leader).status_code, 200)
        self.assertEqual(self.download(self.make_user("crc_chair", campus="Siniloan")).status_code, 200)
        self.assertEqual(self.download(self.make_user("finance_budget", campus="Sta. Cruz")).status_code, 404)
        self.assertEqual(self.download(self.make_user("project_leader", campus="Siniloan")).status_code, 404)
        LineItemBudget.objects.filter(pk=self.budget.pk).update(status="submitted")
        self.assertEqual(self.download().status_code, 404)  # Finance sees it only from For Budget Review on
        self.assertEqual(self.client.get(self.url).status_code, 401)

    def test_the_empty_template_is_unchanged(self):
        import openpyxl

        from research_projects.importer import build_lib_template

        sheet = openpyxl.load_workbook(io.BytesIO(build_lib_template("2027-01-01", "2027-12-31", ["Q1", "Q2", "Q3", "Q4"])))["LIB"]
        header = next(r for r in range(1, 20) if sheet.cell(row=r, column=1).value == "Particulars")
        self.assertEqual(sheet.cell(row=header + 1, column=7).value, f'=IF(AND(E{header + 1}<>"",F{header + 1}<>""),E{header + 1}*F{header + 1},"")')
        self.assertIsNone(sheet.cell(row=header + 1, column=1).value)
        self.assertEqual(sheet.cell(row=header + 30 + 2, column=1).value, "LEGEND")


class FinanceEditTests(RMISTestCase):
    """Task 20 (2026-10-30): Update This LIB. Finance/Accounting/Budget changes existing line items of a LIB forwarded to
    them (For Budget Review or Ongoing Review - Budget), each change with a reason kept in an append-only history;
    no item is added, removed or reclassed, and the Grand Total must still equal the Total Project/Study Cost. The
    Project Leader and the campus CRC/CRD get an e-mail of the committed changes."""

    def setUp(self):
        super().setUp()
        self.leader = self.make_user("project_leader", campus="Siniloan", first_name="Ana", last_name="Reyes")
        self.officer = self.make_user("finance_budget", campus="Siniloan", first_name="Fe", last_name="Cruz")
        self.crc = self.make_user("crc_chair", campus="Siniloan")
        self.make_user("crc_chair", campus="Sta. Cruz")
        self.project = Project.objects.create(
            title="P", project_code="NTP-9", funding_type="core_funded", lead=self.leader, campus="Siniloan",
            total_cost="15500", start_date=date(2027, 1, 1), target_end_date=date(2028, 3, 31),
        )
        self.budget = LineItemBudget.objects.create(project=self.project, version_number=1, status="endorsed")
        self.a = LineItem.objects.create(budget=self.budget, category="mooe", description="Training Expenses", unit="pax",
                                         quantity=Decimal("20"), unit_cost=Decimal("500.50"), amount=Decimal("10010"),
                                         allocations=[{"period": "Y1Q1", "amount": "4010.00"}, {"period": "Y1Q4", "amount": "6000.00"}])
        self.b = LineItem.objects.create(budget=self.budget, category="co", description="Office Equipment", unit="unit",
                                         quantity=Decimal("1"), unit_cost=Decimal("5490"), amount=Decimal("5490"),
                                         allocations=[{"period": "Y2Q1", "amount": "5490.00"}])
        self.brevo_post.reset_mock()

    def update(self, items, user=None):
        with self.captureOnCommitCallbacks(execute=True):
            return self.client_for(user or self.officer).post(f"/api/budget/budgets/{self.budget.id}/finance-edit/", {"items": items}, format="json")

    def move_cost(self):
        """Training Expenses down by 1,000 and Office Equipment up by 1,000: the Grand Total stays 15,500."""
        return [
            {"id": self.a.id, "unit_cost": "450.50", "allocations": [{"period": "Y1Q1", "amount": "3010"}, {"period": "Y1Q4", "amount": "6000"}],
             "reason": "Venue rate per COA ceiling"},
            {"id": self.b.id, "unit_cost": "6490", "allocations": [{"period": "Y2Q1", "amount": "6490"}], "reason": "Updated quotation"},
        ]

    def sent(self):
        return [call.kwargs["json"] for call in self.brevo_post.call_args_list]

    def test_finance_updates_existing_items_with_reasons_and_the_history_is_kept(self):
        response = self.update(self.move_cost())

        self.assertEqual(response.status_code, 200, response.data)
        self.a.refresh_from_db()
        self.b.refresh_from_db()
        self.assertEqual((self.a.amount, self.b.amount), (Decimal("9010.00"), Decimal("6490.00")))
        edit = LineItemEdit.objects.get(line_item=self.a)
        self.assertEqual((edit.edited_by, edit.reason, edit.budget), (self.officer, "Venue rate per COA ceiling", self.budget))
        self.assertEqual(edit.changes["unit_cost"], ["500.50", "450.50"])
        self.assertEqual(edit.changes["amount"], ["10010.00", "9010.00"])
        self.assertIn("allocations", edit.changes)
        self.assertNotIn("quantity", edit.changes)
        item = next(i for i in response.data["line_items"] if i["id"] == self.a.id)
        self.assertEqual([(e["reason"], e["edited_by_name"]) for e in item["edits"]], [("Venue rate per COA ceiling", "Fe Cruz")])
        self.assertEqual(LineItem.objects.filter(budget=self.budget).count(), 2)

    def test_a_later_edit_adds_to_the_history(self):
        self.update(self.move_cost())
        self.update([{"id": self.a.id, "specification": "Venue, 2 days", "reason": "Clarified"}])

        self.assertEqual([e.reason for e in LineItemEdit.objects.filter(line_item=self.a)], ["Venue rate per COA ceiling", "Clarified"])
        self.assertEqual(LineItemEdit.objects.filter(line_item=self.a).last().changes, {"specification": ["", "Venue, 2 days"]})

    def test_the_leader_and_the_campus_crc_crd_get_the_committed_changes_by_e_mail(self):
        self.update(self.move_cost())

        mails = self.sent()
        self.assertEqual(sorted(m["to"][0]["email"] for m in mails), sorted([self.leader.email, self.crc.email]))
        self.assertEqual(mails[0]["subject"], "LIB Updated by Finance/Accounting/Budget — P")
        for text in ("Fe Cruz", "Training Expenses", "Unit Cost", "₱500.50 → ₱450.50", "Venue rate per COA ceiling", "Office Equipment"):
            self.assertIn(text, mails[0]["htmlContent"])

    def test_every_changed_item_needs_a_reason_and_nothing_is_saved_without(self):
        items = self.move_cost()
        items[1]["reason"] = "  "

        response = self.update(items)

        self.assertEqual(response.status_code, 400)
        self.assertIn("reason", response.data["detail"])
        self.a.refresh_from_db()
        self.assertEqual(self.a.unit_cost, Decimal("500.50"))  # the first item's save was rolled back
        self.assertFalse(LineItemEdit.objects.exists())
        self.assertEqual(self.sent(), [])

    def test_the_grand_total_must_still_equal_the_total_cost(self):
        response = self.update(self.move_cost()[:1])

        self.assertEqual(response.status_code, 400)
        self.assertIn("Total Project/Study Cost", response.data["detail"])
        self.a.refresh_from_db()
        self.assertEqual(self.a.amount, Decimal("10010.00"))
        self.assertFalse(LineItemEdit.objects.exists())
        self.assertEqual(self.sent(), [])

    def test_no_real_change_is_refused_and_sends_nothing(self):
        response = self.update([{"id": self.a.id, "quantity": "20.00", "unit_cost": "500.5", "reason": "x"}])

        self.assertEqual((response.status_code, response.data["detail"]), (400, "Nothing was changed."))
        self.assertFalse(LineItemEdit.objects.exists())
        self.assertEqual(self.sent(), [])

    def test_items_cant_be_added_removed_or_reclassed(self):
        other = LineItemBudget.objects.create(project=Project.objects.create(title="Q", project_code="Q-1", funding_type="core_funded",
                                                                             lead=self.leader, campus="Siniloan"), version_number=1)
        foreign = LineItem.objects.create(budget=other, category="mooe", description="Training Expenses", amount=Decimal("1"))
        for items in ([{"category": "mooe", "description": "Training Expenses", "reason": "new"}],
                      [{"id": foreign.id, "unit_cost": "1", "reason": "x"}],
                      [{"id": self.a.id, "category": "ps", "reason": "x"}],
                      [{"id": self.a.id, "amount": "1", "reason": "x"}]):
            self.assertEqual(self.update(items).status_code, 400, items)
        self.assertEqual(LineItem.objects.filter(budget=self.budget).count(), 2)
        client = self.client_for(self.officer)
        self.assertEqual(client.delete(f"/api/budget/line-items/{self.a.id}/").status_code, 403)
        self.assertEqual(client.post("/api/budget/line-items/", {"budget": self.budget.id, "category": "mooe",
                                                                  "description": "Training Expenses", "amount": "1"}, format="json").status_code, 403)
        self.assertFalse(LineItemEdit.objects.exists())

    def test_only_finance_before_certification(self):
        LineItemBudget.objects.filter(pk=self.budget.pk).update(status="certifying")
        self.assertEqual(self.update(self.move_cost()).status_code, 200)
        LineItemBudget.objects.filter(pk=self.budget.pk).update(status="certified")
        self.assertEqual(self.update([{"id": self.a.id, "specification": "x", "reason": "x"}]).status_code, 400)
        LineItemBudget.objects.filter(pk=self.budget.pk).update(status="endorsed")
        for user in (self.leader, self.crc):
            self.assertEqual(self.update([{"id": self.a.id, "specification": "x", "reason": "x"}], user).status_code, 403, user.role.code)
        self.assertEqual(self.update([{"id": self.a.id, "specification": "x", "reason": "x"}],
                                     self.make_user("finance_budget", campus="Sta. Cruz")).status_code, 404)
        LineItemBudget.objects.filter(pk=self.budget.pk).update(status="submitted")
        self.assertEqual(self.update([{"id": self.a.id, "specification": "x", "reason": "x"}]).status_code, 404)  # not forwarded yet
        self.assertEqual(LineItemEdit.objects.count(), 2)  # only the certifying edit
