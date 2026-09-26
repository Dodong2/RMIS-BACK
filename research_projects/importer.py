"""Excel registration of a project in the LSPU-RDO-SF-018 (Research Proposal Form) layout.

build_template() makes the blank workbook; import_workbook() reads a filled one and saves the project and its
tables through the same serializers the manual endpoints use, so both paths follow the same validation rules.
"""
import datetime
import io
import re
from decimal import Decimal

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from accounts.models import User
from accounts.permissions import LEADER_ROLES
from budget_lib.models import LineItem
from budget_lib.serializers import LineItemBudgetSerializer, LineItemSerializer
from outputs.models import SIX_PS
from outputs.serializers import ExpectedOutputSerializer

from .models import Program, Project, ProjectTeamMember
from .serializers import (
    ProjectSerializer, ProjectTeamMemberSerializer, StudySerializer, TargetBeneficiarySerializer,
    WorkPlanMilestoneSerializer, leader_owns,
)

NEW_OR_CONTINUING = (("new", "New Proposal"), ("continuing", "Continuing"))
WET_OR_DRY = (("wet", "Wet Research"), ("dry", "Dry Research"))

# (label in column A, serializer field, kind, choices, note in column C). Order follows the form.
PROJECT_FIELDS = [
    ("Title", "title", "text", None, "Required"),
    ("LSPU Faculty Research Number (Project Code)", "project_code", "text", None, "Required, must be unique"),
    ("Program Code (if under a Program)", "program", "program", None, "Optional; code of an existing RMIS Program"),
    ("Funding Type", "funding_type", "choice", Program.FUNDING_CHOICES, "Blank = Institutional (LSPU-Funded)"),
    ("Project Leader E-mail Address", "lead", "user", None, "Required; RMIS account with the Project Leader role"),
    ("Project Leader Gender", "lead_gender", "choice", Project.GENDER_CHOICES, "Male or Female"),
    ("Contact No./s.", "contact_number", "text", None, ""),
    ("Start Date", "start_date", "date", None, "Date cell or YYYY-MM-DD"),
    ("End Date", "target_end_date", "date", None, "Date cell or YYYY-MM-DD"),
    ("Total Project/Study Cost", "total_cost", "decimal", None, "Number"),
    ("Implementing Unit", "implementing_unit", "text", None, ""),
    ("Campus", "campus", "text", None, ""),
    ("College Unit", "college", "text", None, "e.g. CTE"),
    ("Cooperating Agency/ies", "cooperating_agencies", "text", None, ""),
    ("Sector", "sectors", "multi", Project.SECTOR_CHOICES, "Required; one or more, comma-separated"),
    ("Sector (if Others)", "sector_other", "text", None, "Required when Sector includes Others"),
    ("New or Continuing", "is_continuing", "choice", NEW_OR_CONTINUING, "Blank = New Proposal"),
    ("Continuing Year", "continuing_year", "int", None, "2, 3, ... (continuing proposals only)"),
    ("Basic or Applied Research", "research_type", "choice", Project.RESEARCH_TYPE_CHOICES, ""),
    ("Wet or Dry Research", "is_dry_research", "choice", WET_OR_DRY, "Blank = Dry"),
    ("Research Priority Area", "research_priority_area", "choice", Project.PRIORITY_AREA_CHOICES, ""),
    ("Research Typology", "research_typology", "multi", Project.TYPOLOGY_CHOICES, "One or more, comma-separated"),
    ("Sustainable Development Goals (SDGs)", "sdgs", "sdgs", None, "Required; goal numbers 1-17, comma-separated"),
    ("REI Thrust", "rei_thrust", "text", None, "Optional"),
    ("II. Background of the Study", "background", "text", None, ""),
    ("III. Objectives of the Study", "objectives", "text", None, ""),
    ("IV. Project Descriptions/Methodology", "methodology", "text", None, ""),
    ("VI. Socio-Economic Significance", "socio_economic_significance", "text", None, ""),
    ("VIII. Monitoring/Evaluation", "monitoring_evaluation", "text", None, ""),
    ("IX. List of References", "references", "text", None, "APA format"),
    ("Date Submitted (Project Leader)", "proposal_submitted_on", "date", None, "Annex A"),
    ("Endorsed By (Dean/Associate Dean)", "endorsed_by_dean", "text", None, "Annex A"),
    ("Date Endorsed by Dean", "endorsed_by_dean_on", "date", None, "Annex A"),
    ("Noted By (RDS Director/Chairperson)", "noted_by_rds_director", "text", None, "Annex A"),
    ("Date Noted by RDS Director", "noted_by_rds_director_on", "date", None, "Annex A"),
    ("Recommending Approval (Campus Director)", "recommended_by_campus_director", "text", None, "Annex A"),
    ("Date Recommended by Campus Director", "recommended_by_campus_director_on", "date", None, "Annex A"),
    ("Recommending Approval (VPRDE)", "recommended_by_vprde", "text", None, "Annex A"),
    ("Date Recommended by VPRDE", "recommended_by_vprde_on", "date", None, "Annex A"),
    ("Approved By (University President)", "approved_by_president", "text", None, "Annex A"),
    ("Date Approved by University President", "proposal_approved_on", "date", None, "Annex A"),
]

# Table sheets: sheet name -> [(column header, serializer field, kind, choices)]
TEAM_ROLES = ProjectTeamMember.ROLE_CHOICES
CATEGORIES = LineItem.CATEGORY_CHOICES + (("mooe", "MOOE"), ("co", "Equipment Outlay"), ("ps", "PS"))
# The form's own 6P wording (Section V), accepted alongside the RMIS labels
SIX_P_CHOICES = SIX_PS + (
    ("patents", "Patent"), ("places_partnerships", "Places/Partnerships"), ("policies", "Policy Recommendations"),
)
TABLE_SHEETS = {
    "Project Team": [
        ("Role", "member_role", "choice", TEAM_ROLES),
        ("Name", "name", "text", None),
        ("Gender", "gender", "choice", Project.GENDER_CHOICES),
        ("RMIS E-mail (optional)", "user", "user", None),
    ],
    "Study Components": [
        ("Study Title", "title", "text", None),
        ("Study Leader E-mail", "lead", "user", None),
    ],
    "Expected Outputs (6Ps)": [
        ("Item", "category", "choice", SIX_P_CHOICES),
        ("Particulars", "description", "text", None),
        ("Quantity", "target_count", "int", None),
    ],
    "Target Beneficiaries": [
        ("Target Beneficiaries", "group", "text", None),
        ("Description", "description", "text", None),
        ("Total", "total", "int", None),
    ],
    "Budget Requirements": [
        ("Fiscal Year", "fiscal_year", "int", None),
        ("Category", "category", "choice", CATEGORIES),
        ("Particulars", "description", "text", None),
        ("QTR1", "q1_amount", "decimal", None),
        ("QTR2", "q2_amount", "decimal", None),
        ("QTR3", "q3_amount", "decimal", None),
        ("QTR4", "q4_amount", "decimal", None),
        ("Total", "amount", "decimal", None),
        ("Funding Source", "funding_source", "text", None),
    ],
    "Work Plan": [
        ("Activity", "title", "text", None),
        ("Start Date", "start_date", "date", None),
        ("Target Date", "target_date", "date", None),
    ],
}
TABLE_SERIALIZERS = {
    "Project Team": ProjectTeamMemberSerializer,
    "Study Components": StudySerializer,
    "Expected Outputs (6Ps)": ExpectedOutputSerializer,
    "Target Beneficiaries": TargetBeneficiarySerializer,
    "Budget Requirements": LineItemSerializer,
    "Work Plan": WorkPlanMilestoneSerializer,
}


def _norm(value):
    return re.sub(r"\s+", " ", str(value)).strip().rstrip(".").lower()


class CellError(ValueError):
    pass


def _convert(raw, kind, choices):
    """Turn one cell into the value its serializer field expects. Blank cells return None (field left out)."""
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return None
    if kind == "text":
        return str(raw).strip()
    if kind == "date":
        if isinstance(raw, datetime.datetime):
            return raw.date()
        if isinstance(raw, datetime.date):
            return raw
        try:
            return datetime.date.fromisoformat(str(raw).strip())
        except ValueError:
            raise CellError(f"'{raw}' is not a date (use YYYY-MM-DD).")
    if kind in ("int", "decimal"):
        text = str(raw).replace(",", "").strip()
        try:
            number = Decimal(text)
        except ArithmeticError:
            raise CellError(f"'{raw}' is not a number.")
        if kind == "int":
            if number != number.to_integral_value():
                raise CellError(f"'{raw}' must be a whole number.")
            return int(number)
        return str(number)
    if kind == "user":
        user = User.objects.filter(email__iexact=str(raw).strip()).first()
        if not user:
            raise CellError(f"No RMIS account with e-mail '{raw}'.")
        return user.pk
    if kind == "program":
        program = Program.objects.filter(code__iexact=str(raw).strip()).first()
        if not program:
            raise CellError(f"No RMIS Program with code '{raw}'.")
        return program.pk
    if kind == "sdgs":
        try:
            return [int(part) for part in re.split(r"[,;]", str(raw)) if part.strip()]
        except ValueError:
            raise CellError(f"'{raw}' must be SDG numbers separated by commas.")
    lookup = {}
    for code, label in choices:
        lookup[_norm(code)] = code
        lookup[_norm(label)] = code
    if kind == "multi":
        parts = [p for p in re.split(r"[,;]", str(raw)) if p.strip()]
        unknown = [p.strip() for p in parts if _norm(p) not in lookup]
        if unknown:
            raise CellError(f"Unknown value(s): {', '.join(unknown)}.")
        return [lookup[_norm(p)] for p in parts]
    if _norm(raw) not in lookup:
        raise CellError(f"Unknown value '{raw}'.")
    return lookup[_norm(raw)]


def _allowed(kind, choices):
    return ", ".join(label for _, label in choices) if choices and kind in ("choice", "multi") else ""


def build_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Project"
    ws.append(["Field", "Value", "Notes / allowed values"])
    for label, _, kind, choices, note in PROJECT_FIELDS:
        allowed = _allowed(kind, choices)
        ws.append([label, None, f"{note}. {allowed}".strip(". ") if allowed else note])
    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 70
    for cell in ws["B"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for name, columns in TABLE_SHEETS.items():
        sheet = wb.create_sheet(name)
        sheet.append([header for header, *_ in columns])
        for i, (header, _, kind, choices) in enumerate(columns):
            sheet.column_dimensions[get_column_letter(i + 1)].width = max(18, len(header) + 4)
            allowed = _allowed(kind, choices)
            if allowed:
                sheet.cell(row=1, column=i + 1).comment = Comment(f"Allowed: {allowed}", "RMIS")
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _read_project_sheet(ws, errors):
    by_label = {_norm(label): (field, kind, choices) for label, field, kind, choices, _ in PROJECT_FIELDS}
    data = {}
    for row_number, row in enumerate(ws.iter_rows(min_row=2, max_col=2, values_only=True), start=2):
        if not row or row[0] is None or _norm(row[0]) not in by_label:
            continue
        field, kind, choices = by_label[_norm(row[0])]
        try:
            value = _convert(row[1] if len(row) > 1 else None, kind, choices)
        except CellError as exc:
            errors.append({"sheet": "Project", "row": row_number, "field": row[0], "message": str(exc)})
            continue
        if value is not None:
            data[field] = value
    if "is_continuing" in data:
        data["is_continuing"] = data["is_continuing"] == "continuing"
    if "is_dry_research" in data:
        data["is_dry_research"] = data["is_dry_research"] == "dry"
    data.setdefault("funding_type", "institutional")
    return data


def _read_table(ws, name, errors):
    """Rows of {field: value} keyed by the sheet's headers; blank rows are skipped. Returns [(row_number, data)]."""
    columns = TABLE_SHEETS[name]
    rows = ws.iter_rows(values_only=True)
    headers = [_norm(h) if h is not None else "" for h in next(rows, [])]
    index = {}
    for header, field, kind, choices in columns:
        if _norm(header) in headers:
            index[field] = (headers.index(_norm(header)), header, kind, choices)
    out = []
    for row_number, row in enumerate(rows, start=2):
        if not any(v is not None and str(v).strip() for v in row):
            continue
        data = {}
        for field, (col, header, kind, choices) in index.items():
            try:
                value = _convert(row[col] if col < len(row) else None, kind, choices)
            except CellError as exc:
                errors.append({"sheet": name, "row": row_number, "field": header, "message": str(exc)})
                continue
            if value is not None:
                data[field] = value
        out.append((row_number, data))
    return out


def _serializer_errors(sheet, row, serializer_errors, labels):
    items = []
    for field, messages in serializer_errors.items():
        messages = messages if isinstance(messages, list) else [messages]
        for message in messages:
            items.append({"sheet": sheet, "row": row, "field": labels.get(field, field), "message": str(message)})
    return items


def import_workbook(file_obj, request):
    """Create one project from a filled template. Returns (project, errors); on any error nothing should be kept,
    so the caller must run this inside transaction.atomic() and roll back when errors is non-empty."""
    try:
        wb = openpyxl.load_workbook(file_obj, data_only=True)
    except Exception as exc:  # not an xlsx, or corrupted
        return None, [{"sheet": None, "row": None, "field": "file", "message": f"Could not read workbook: {exc}"}]
    if "Project" not in wb.sheetnames:
        return None, [{"sheet": None, "row": None, "field": "file", "message": "Missing the 'Project' sheet; use the RMIS template."}]

    errors = []
    context = {"request": request, "defer_scope": True}
    project_data = _read_project_sheet(wb["Project"], errors)
    tables = {name: _read_table(wb[name], name, errors) for name in TABLE_SHEETS if name in wb.sheetnames}
    if errors:
        return None, errors

    serializer = ProjectSerializer(data=project_data, context=context)
    if not serializer.is_valid():
        labels = {field: label for label, field, *_ in PROJECT_FIELDS}
        return None, _serializer_errors("Project", None, serializer.errors, labels)
    project = serializer.save()

    budget = None
    for name, rows in tables.items():
        labels = {field: header for header, field, *_ in TABLE_SHEETS[name]}
        for row_number, data in rows:
            data["project"] = project.pk
            if name == "Budget Requirements":
                if "amount" not in data:
                    data["amount"] = str(sum(Decimal(data.get(f, 0)) for f in ("q1_amount", "q2_amount", "q3_amount", "q4_amount")))
                if budget is None:
                    budget = LineItemBudgetSerializer(data={"project": project.pk}, context=context)
                    budget.is_valid(raise_exception=True)
                    budget = budget.save()
                data["budget"] = budget.pk
            row_serializer = TABLE_SERIALIZERS[name](data=data, context=context)
            if row_serializer.is_valid():
                row_serializer.save()
            else:
                errors.extend(_serializer_errors(name, row_number, row_serializer.errors, labels))
    user = request.user
    if user.role and user.role.code in LEADER_ROLES and not leader_owns(user, project):
        errors.append({"sheet": "Project", "row": None, "field": "Project Leader E-mail Address", "message": (
            "Leaders can only upload a project they lead, one under a Program they lead (Program Code), "
            "or one with a Study Component they lead."
        )})
    return project, errors
