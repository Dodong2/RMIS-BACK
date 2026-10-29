import datetime
import io
from decimal import Decimal

import openpyxl

from accounts.testing import RMISTestCase
from accounts.models import User
from budget_lib.models import LibItem, LineItem, LineItemBudget
from outputs.models import ExpectedOutput
from research_projects.models import (
    CollegeUnit, CooperatingAgency, Project, ProjectEndorser, RdeThrust, Study, TargetBeneficiary, WorkPlanMilestone,
)
from research_projects.rde_agenda import seed_rde_agenda
from research_projects.serializers import REQUIRED_PROJECT_FIELDS


def project_payload(lead, **overrides):
    """Every Step 1 field filled (client 2026-10-29); the NTP number is the project code."""
    seed_rde_agenda(RdeThrust)  # a TransactionTestCase flush also drops the migration-seeded agenda
    data = {
        "title": "Bridging Academia and Communities", "ntp_number": "LSPU-1", "funding_type": "core_funded",
        "lead": lead.id, "sdgs": [4, 17], "sectors": ["education"],
        "research_type": "applied", "is_dry_research": True, "is_continuing": False,
        "start_date": "2027-01-01", "target_end_date": "2027-12-31", "total_cost": "100000.00",
        "rde_thrust": RdeThrust.objects.get(code="1").name, "rde_priority": RdeThrust.objects.get(code="1.3").name,
        "research_priority_area": "education_teacher_training", "research_typology": ["quantitative"],
        "objectives": "1. Map community needs", "background": "<p>Background</p>", "methodology": "<p>Method</p>",
        "socio_economic_significance": "<p>Significance</p>", "monitoring_evaluation": "<p>M&amp;E</p>",
        "references": "<p>References</p>", "description": "Summary", "expected_outcomes": "Outcomes",
        "expected_impacts": "Impacts",
    }
    data.update(overrides)
    return data


class StepOneRequiredTests(RMISTestCase):
    """Client 2026-10-29 (Task 4): every Step 1 field is required and the NTP number is the project code."""

    def setUp(self):
        super().setUp()
        self.admin, self.leader = self.make_user("system_admin"), self.make_user("project_leader")
        self.client_ = self.client_for(self.admin)

    def register(self, payload):
        return self.client_.post("/api/projects/", payload, format="json")

    def test_each_step_1_field_is_required(self):
        for field in REQUIRED_PROJECT_FIELDS:
            payload = project_payload(self.leader)
            del payload[field]
            response = self.register(payload)
            self.assertEqual(response.status_code, 400, field)
            self.assertIn(field, response.data, field)
        self.assertFalse(Project.objects.exists())

    def test_blank_values_and_an_emptied_editor_count_as_missing(self):
        response = self.register(project_payload(
            self.leader, ntp_number="  ", research_typology=[], objectives="<p></p>", background="<p></p>",
        ))

        self.assertEqual(response.status_code, 400)
        self.assertTrue({"ntp_number", "research_typology", "objectives"} <= set(response.data))
        self.assertNotIn("background", response.data)

    def test_priority_area_and_proposal_content_are_no_longer_required(self):
        """Client 2026-10-29 (Tasks 6/8): RDE Agenda replaces Research Priority Area; Proposal Content is dropped except
        the Objectives."""
        removed = ("research_priority_area", "background", "methodology", "socio_economic_significance",
                   "monitoring_evaluation", "references", "description", "expected_outcomes", "expected_impacts")
        payload = {k: v for k, v in project_payload(self.leader).items() if k not in removed}

        response = self.register(payload)

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(Project.objects.values_list("research_priority_area", "background", "objectives").get(),
                         ("", "", "1. Map community needs"))

    def test_an_edit_without_the_dropped_fields_keeps_their_values(self):
        project_id = self.register(project_payload(self.leader, background="<p>Old background</p>")).data["id"]

        response = self.client_.patch(f"/api/projects/{project_id}/", {"title": "New title"}, format="json")

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(Project.objects.values_list("title", "background", "research_priority_area").get(),
                         ("New title", "<p>Old background</p>", "education_teacher_training"))

    def test_wet_dry_and_new_continuing_must_be_picked(self):
        payload = project_payload(self.leader)
        del payload["is_dry_research"], payload["is_continuing"]

        response = self.register(payload)

        self.assertEqual(response.status_code, 400)
        self.assertEqual(set(response.data), {"is_dry_research", "is_continuing"})

    def test_not_applicable_is_a_valid_pick(self):
        response = self.register(project_payload(self.leader, is_dry_research=None, research_type="not_applicable"))

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual((response.data["is_dry_research"], response.data["research_type"]), (None, "not_applicable"))

    def test_the_ntp_number_is_the_project_code(self):
        response = self.register(project_payload(self.leader, ntp_number=" NTP-2027-001 ", project_code="IGNORED"))

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual((response.data["ntp_number"], response.data["project_code"]), ("NTP-2027-001", "NTP-2027-001"))

    def test_an_ntp_number_already_used_as_a_code_is_refused(self):
        Project.objects.create(title="Old", project_code="NTP-1", funding_type="core_funded", lead=self.leader)

        response = self.register(project_payload(self.leader, ntp_number="NTP-1"))

        self.assertEqual(response.status_code, 400)
        self.assertIn("ntp_number", response.data)

    def test_the_ntp_number_is_saved_uppercase_and_case_does_not_make_it_new(self):
        """Client 2026-10-29: NTP letters are uppercase (e.g. LSPU-RD-GRANT-SIN-2024-001), from any source."""
        response = self.register(project_payload(self.leader, ntp_number=" lspu-rd-grant-sin-2024-001 "))
        again = self.register(project_payload(self.leader, ntp_number="LSPU-RD-Grant-SIN-2024-001", title="Second"))

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual((response.data["ntp_number"], response.data["project_code"]),
                         ("LSPU-RD-GRANT-SIN-2024-001", "LSPU-RD-GRANT-SIN-2024-001"))
        self.assertEqual(again.status_code, 400)
        self.assertIn("ntp_number", again.data)

    def test_only_the_system_admin_changes_the_ntp_and_the_code_follows(self):
        project_id = self.register(project_payload(self.leader, ntp_number="NTP-1")).data["id"]

        by_leader = self.client_for(self.leader).patch(f"/api/projects/{project_id}/", {"ntp_number": "NTP-9"}, format="json")
        by_admin = self.client_.patch(f"/api/projects/{project_id}/", {"ntp_number": "NTP-2"}, format="json")

        self.assertEqual(by_leader.status_code, 400)
        self.assertEqual(by_admin.status_code, 200, by_admin.data)
        self.assertEqual(Project.objects.values_list("ntp_number", "project_code").get(), ("NTP-2", "NTP-2"))

    def test_an_edit_of_an_older_record_only_checks_the_fields_it_sends(self):
        project = Project.objects.create(title="Old", project_code="OLD-1", funding_type="core_funded", lead=self.leader,
                                         sdgs=[4], sectors=["education"])

        response = self.client_.patch(f"/api/projects/{project.pk}/", {"title": "Old, renamed"}, format="json")

        self.assertEqual(response.status_code, 200, response.data)


class NestedRegistrationTests(RMISTestCase):
    """Client 2026-10-29 (Task 4): Manual Entry sends the form's tables with the project, saved all-or-nothing."""

    def setUp(self):
        super().setUp()
        self.admin, self.leader = self.make_user("system_admin"), self.make_user("project_leader")

    def payload(self, **tables):
        data = project_payload(self.leader)
        data.update({
            "team": [{"member_role": "co_leader", "name": "Archieval M. Jain", "gender": "male", "user": None}],
            "studies": [{"title": "Study 1"}, {"title": "Study 2"}],
            "outputs": [{"category": "publications", "description": "Journal article", "target_count": 2}],
            "beneficiaries": [{"group": "Faculty", "description": "CTE faculty", "total": 10}],
            "milestones": [{"title": "Data gathering", "start_date": "2027-02-01", "target_date": "2027-04-30"}],
            "endorsers": [{"role_code": "", "user": None, "name": "Dean A", "designation": "Dean", "signed_on": None},
                          {"role_code": "", "user": None, "name": "Director B", "designation": "RDS Director", "signed_on": None}],
            "line_items": [{"category": "mooe", "description": "Traveling Expenses - Local", "fiscal_year": 2027, "unit": "lot",
                            "quantity": "1", "unit_cost": "100000", "allocations": [{"period": "Q1", "amount": "100000"}]}],
        })
        data.update(tables)
        return data

    def register(self, payload, **headers):
        return self.client_for(self.admin).post("/api/projects/", payload, format="json", **headers)

    def saved_rows(self):
        return (Project.objects.count(), Study.objects.count(), ExpectedOutput.objects.count(),
                TargetBeneficiary.objects.count(), WorkPlanMilestone.objects.count(), ProjectEndorser.objects.count(),
                LineItemBudget.objects.count(), LineItem.objects.count())

    def test_the_project_and_its_tables_are_saved_together(self):
        response = self.register(self.payload())

        self.assertEqual(response.status_code, 201, response.data)
        project = Project.objects.get()
        self.assertEqual(response.data["id"], project.pk)
        self.assertEqual(self.saved_rows(), (1, 2, 1, 1, 1, 2, 1, 1))
        self.assertEqual(project.team_members.get().name, "Archieval M. Jain")
        self.assertEqual(list(project.endorsers.order_by("id").values_list("name", flat=True)), ["Dean A", "Director B"])
        budget = LineItemBudget.objects.get()
        self.assertEqual((budget.version_number, budget.status), (1, "draft"))
        self.assertEqual(float(budget.line_items.get().amount), 100000)

    def test_one_bad_row_saves_nothing(self):
        response = self.register(self.payload(beneficiaries=[{"group": "", "description": "x", "total": 1}]))

        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.data["errors"]["beneficiaries"][0]["row"], 1)
        self.assertEqual(self.saved_rows(), (0,) * 8)

    def test_the_lib_keeps_the_lib_steps_rules(self):
        # A dated project's line items need their quarterly split, as on POST projects/<id>/lib/
        row = {"category": "mooe", "description": "Traveling Expenses - Local", "unit": "lot", "quantity": "1", "unit_cost": "100"}

        response = self.register(self.payload(line_items=[row]))

        self.assertEqual(response.status_code, 400)
        self.assertIn("line_items", response.data["errors"])
        self.assertEqual(self.saved_rows(), (0,) * 8)

    def test_a_table_that_is_not_a_list_is_refused(self):
        response = self.register(self.payload(studies=None, team="Archieval"))

        self.assertEqual(response.status_code, 400)
        self.assertEqual(set(response.data["errors"]), {"studies", "team"})
        self.assertFalse(Project.objects.exists())

    def test_empty_tables_add_no_rows(self):
        response = self.register(self.payload(team=[], outputs=[], milestones=[], endorsers=[], line_items=[]))

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(self.saved_rows(), (1, 2, 0, 1, 0, 0, 0, 0))

    def test_a_retried_registration_replays_instead_of_saving_twice(self):
        first = self.register(self.payload(), HTTP_IDEMPOTENCY_KEY="register-0001")
        again = self.register(self.payload(), HTTP_IDEMPOTENCY_KEY="register-0001")

        self.assertEqual((first.status_code, again.status_code), (201, 201))
        self.assertEqual(again.data["id"], first.data["id"])
        self.assertEqual(self.saved_rows(), (1, 2, 1, 1, 1, 2, 1, 1))


class ManualRegistrationTests(RMISTestCase):
    def test_registrar_can_register_a_project_with_several_sectors(self):
        admin, leader = self.make_user("system_admin"), self.make_user("project_leader")

        response = self.client_for(admin).post(
            "/api/projects/", project_payload(leader, sectors=["education", "community_development"]), format="json"
        )

        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data["sectors"], ["community_development", "education"])

    def test_others_sector_needs_a_description(self):
        admin, leader = self.make_user("system_admin"), self.make_user("project_leader")

        response = self.client_for(admin).post("/api/projects/", project_payload(leader, sectors=["others"]), format="json")

        self.assertEqual(response.status_code, 400)
        self.assertIn("sector_other", response.data)

    def test_continuing_proposal_needs_its_year(self):
        admin, leader = self.make_user("system_admin"), self.make_user("project_leader")

        response = self.client_for(admin).post("/api/projects/", project_payload(leader, is_continuing=True), format="json")

        self.assertEqual(response.status_code, 400)
        self.assertIn("continuing_year", response.data)

    def test_roles_outside_module_2_cannot_register(self):
        leader = self.make_user("project_leader")

        response = self.client_for(self.make_user("finance_budget")).post("/api/projects/", project_payload(leader), format="json")

        self.assertEqual(response.status_code, 403)

    def test_crc_chair_no_longer_registers_a_project(self):
        """Client 2026-10-24 (Carl, 1b): the CRC Chair reviews submitted projects; registering is the Project Leader's,
        RIUH's, DRD's or the System Admin's job, so a CRC Chair's draft can't hide from the CRC Chair."""
        leader = self.make_user("project_leader")
        crc = self.client_for(self.make_user("crc_chair"))

        self.assertEqual(crc.post("/api/projects/", project_payload(leader), format="json").status_code, 403)
        self.assertEqual(crc.get("/api/projects/import-template/").status_code, 403)
        self.assertFalse(Project.objects.exists())

    def test_project_leader_can_register_a_project_they_lead(self):
        """Client meeting 2026-10-01: the Project Leader registers their own approved project."""
        leader = self.make_user("project_leader")

        response = self.client_for(leader).post("/api/projects/", project_payload(leader), format="json")

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(Project.objects.get().lead, leader)

    def test_project_leader_cannot_register_a_project_for_another_leader(self):
        leader, other = self.make_user("project_leader"), self.make_user("project_leader")

        response = self.client_for(leader).post("/api/projects/", project_payload(other), format="json")

        self.assertEqual(response.status_code, 400)
        self.assertIn("lead", response.data)
        self.assertFalse(Project.objects.exists())

    def test_roles_without_projects_register_cannot_register(self):
        leader = self.make_user("project_leader")

        for code in ("crc_chair", "project_staff"):  # Program/Study Leaders merged into the Project Leader 2026-10-27
            response = self.client_for(self.make_user(code)).post("/api/projects/", project_payload(leader), format="json")
            self.assertEqual(response.status_code, 403)
        self.assertFalse(Project.objects.exists())

    def test_leader_can_edit_their_own_project_but_not_the_crc_owned_fields(self):
        leader = self.make_user("project_leader")
        project = Project.objects.create(
            title="P", project_code="P-1", ntp_number="P-1", funding_type="core_funded", lead=leader, sdgs=[4],
            sectors=["education"],
        )
        client = self.client_for(leader)

        edited = client.patch(f"/api/projects/{project.id}/", {"title": "Better title", "ntp_number": "P-1"}, format="json")
        renumbered = client.patch(f"/api/projects/{project.id}/", {"ntp_number": "P-2"}, format="json")
        recoded = client.patch(f"/api/projects/{project.id}/", {"project_code": "P-3"}, format="json")

        self.assertEqual(edited.status_code, 200, edited.data)
        self.assertEqual(renumbered.status_code, 400)
        self.assertIn("ntp_number", renumbered.data)
        self.assertEqual(recoded.status_code, 200)  # read-only since the NTP is the code (2026-10-29), so ignored
        project.refresh_from_db()
        self.assertEqual((project.title, project.ntp_number, project.project_code), ("Better title", "P-1", "P-1"))

    def test_project_leader_cannot_edit_another_leaders_project(self):
        leader, other = self.make_user("project_leader"), self.make_user("project_leader")
        project = Project.objects.create(title="P", project_code="P-1", funding_type="core_funded", lead=other)

        response = self.client_for(leader).patch(f"/api/projects/{project.id}/", {"title": "Hijacked"}, format="json")

        # 404 since strict read scope (client meeting 2026-10-01): another leader's project isn't visible at all.
        self.assertEqual(response.status_code, 404)
        project.refresh_from_db()
        self.assertEqual(project.title, "P")

    def test_code_available_sees_codes_outside_the_leaders_scope(self):
        """Client meeting 2026-10-01 (#3): a leader can't list other leaders' projects, so the wizard asks this."""
        leader, other = self.make_user("project_leader"), self.make_user("project_leader")
        Project.objects.create(title="P", project_code="LSPU-1", funding_type="core_funded", lead=other)
        client = self.client_for(leader)

        taken = client.get("/api/projects/code-available/", {"code": " lspu-1 "})
        free = client.get("/api/projects/code-available/", {"code": "LSPU-2"})
        staff = self.client_for(self.make_user("project_staff")).get("/api/projects/code-available/", {"code": "LSPU-2"})

        self.assertEqual((taken.status_code, taken.data["available"]), (200, False))
        self.assertEqual((free.status_code, free.data["available"]), (200, True))
        self.assertEqual(staff.status_code, 403)

    def test_study_components_need_only_a_title(self):
        """Client meeting 2026-10-01 (#7): SF-018 lists Study 1, Study 2 by title; a study leader is optional."""
        leader, study_leader = self.make_user("project_leader"), self.make_user("project_leader")
        project = Project.objects.create(title="P", project_code="P-1", funding_type="core_funded", lead=leader)
        client = self.client_for(leader)

        untitled_lead = client.post("/api/studies/", {"project": project.id, "title": "Study 1"}, format="json")
        with_lead = client.post(
            "/api/studies/", {"project": project.id, "title": "Study 2", "lead": study_leader.id}, format="json"
        )
        staff = self.make_user("project_staff")
        wrong_role = client.post("/api/studies/", {"project": project.id, "title": "Study 3", "lead": staff.id}, format="json")

        self.assertEqual(untitled_lead.status_code, 201, untitled_lead.data)
        self.assertIsNone(untitled_lead.data["lead"])
        self.assertEqual(with_lead.status_code, 201, with_lead.data)
        self.assertEqual(wrong_role.status_code, 400)

    def test_endorser_rows_pick_a_role_and_an_account_of_that_role(self):
        """Client meeting 2026-10-01 (#8): Annex A is role -> account rows; text-only rows cover non-accounts."""
        leader, vp = self.make_user("project_leader"), self.make_user("vprei", first_name="Robert", last_name="Agatep")
        project = Project.objects.create(title="P", project_code="P-1", funding_type="core_funded", lead=leader)
        client = self.client_for(leader)

        picked = client.post("/api/project-endorsers/", {
            "project": project.id, "role_code": "vprei", "user": vp.id, "name": "Robert Agatep",
            "designation": "VPRDE", "signed_on": "2025-02-21",
        }, format="json")
        typed = client.post("/api/project-endorsers/", {
            "project": project.id, "name": "Adriel G. Roman", "designation": "Dean/Associate Dean",
        }, format="json")
        mismatched = client.post("/api/project-endorsers/", {
            "project": project.id, "role_code": "crc_chair", "user": vp.id, "name": "Robert Agatep",
        }, format="json")
        outsider = self.client_for(self.make_user("project_leader")).post("/api/project-endorsers/", {
            "project": project.id, "name": "Someone",
        }, format="json")

        self.assertEqual(picked.status_code, 201, picked.data)
        self.assertEqual(typed.status_code, 201, typed.data)
        self.assertEqual(mismatched.status_code, 400)
        self.assertEqual(outsider.status_code, 400)
        self.assertEqual(ProjectEndorser.objects.filter(project=project).count(), 2)

    def test_registration_lib_step_creates_draft_lib_v1(self):
        """Client meeting 2026-10-01 (#11): the LIB is part of registration, for every registration role."""
        leader = self.make_user("project_leader")
        riuh_project = Project.objects.create(title="A", project_code="A-1", funding_type="core_funded", lead=leader)
        own_project = Project.objects.create(title="B", project_code="B-1", funding_type="core_funded", lead=leader)
        rows = [
            {"category": "mooe", "description": "Traveling Expenses - Local", "fiscal_year": 2026, "unit": "pax", "quantity": "4", "unit_cost": "2500"},
            {"category": "co", "description": "Information and Communication Technology Equipment", "fiscal_year": 2026, "unit": "unit", "quantity": "1", "unit_cost": "7500"},
        ]

        by_riuh = self.client_for(self.make_user("riuh")).post(f"/api/projects/{riuh_project.id}/lib/", {"line_items": rows}, format="json")
        by_leader = self.client_for(leader).post(f"/api/projects/{own_project.id}/lib/", {"line_items": rows}, format="json")
        again = self.client_for(leader).post(f"/api/projects/{own_project.id}/lib/", {"line_items": rows}, format="json")
        finance = self.client_for(self.make_user("finance_budget")).post(f"/api/projects/{own_project.id}/lib/", {"line_items": rows}, format="json")

        self.assertEqual(by_riuh.status_code, 201, by_riuh.data)
        self.assertEqual(by_leader.status_code, 201, by_leader.data)
        self.assertEqual((by_leader.data["version_number"], by_leader.data["status"]), (1, "draft"))
        self.assertEqual(float(by_leader.data["total_amount"]), 17500)
        self.assertEqual(again.status_code, 400)
        self.assertEqual(finance.status_code, 403)

    def test_registration_lib_step_is_all_or_nothing_and_scoped(self):
        leader, other = self.make_user("project_leader"), self.make_user("project_leader")
        project = Project.objects.create(title="P", project_code="P-1", funding_type="core_funded", lead=leader)
        rows = [{"category": "mooe", "description": "Traveling Expenses - Local", "unit": "lot", "quantity": "1", "unit_cost": "100"}, {"category": "nope", "description": "Bad"}]

        bad_row = self.client_for(leader).post(f"/api/projects/{project.id}/lib/", {"line_items": rows}, format="json")
        outsider = self.client_for(other).post(f"/api/projects/{project.id}/lib/", {"line_items": rows[:1]}, format="json")

        self.assertEqual(bad_row.status_code, 400)
        self.assertEqual(bad_row.data["errors"][0]["row"], 2)
        self.assertEqual(outsider.status_code, 400)
        self.assertFalse(LineItemBudget.objects.exists())

    def test_line_item_total_must_equal_qty_times_unit_cost(self):
        leader = self.make_user("project_leader")
        project = Project.objects.create(title="P", project_code="P-1", funding_type="core_funded", lead=leader)
        budget = self.client_for(leader).post("/api/budget/budgets/", {"project": project.id}, format="json").data
        row = {"budget": budget["id"], "category": "mooe", "description": "Representation Expenses", "unit": "pax", "quantity": "30", "unit_cost": "85.50"}

        wrong = self.client_for(leader).post("/api/budget/line-items/", {**row, "amount": "2500"}, format="json")
        right = self.client_for(leader).post("/api/budget/line-items/", {**row, "amount": "2565"}, format="json")

        self.assertEqual(wrong.status_code, 400)
        self.assertEqual(right.status_code, 201, right.data)

    def test_budget_reviewers_cannot_encode_line_items(self):
        """Client 2026-10-11: finance_budget and system_admin only certify/approve; the project leader encodes the LIB."""
        leader = self.make_user("project_leader")
        project = Project.objects.create(title="P", project_code="P-1", funding_type="core_funded", lead=leader)
        budget = self.client_for(leader).post("/api/budget/budgets/", {"project": project.id}, format="json").data
        row = {"budget": budget["id"], "category": "mooe", "description": "Traveling Expenses - Local", "unit": "lot", "quantity": "1", "unit_cost": "100", "amount": "100"}
        created = self.client_for(leader).post("/api/budget/line-items/", row, format="json")
        self.assertEqual(created.status_code, 201, created.data)
        item = created.data

        for role in ("finance_budget", "system_admin"):
            reviewer = self.client_for(self.make_user(role))
            self.assertEqual(reviewer.post("/api/budget/budgets/", {"project": project.id}, format="json").status_code, 403)
            self.assertEqual(reviewer.post("/api/budget/line-items/", row, format="json").status_code, 403)
            self.assertEqual(reviewer.patch(f"/api/budget/line-items/{item['id']}/", {"description": "X"}, format="json").status_code, 403)

    def test_line_item_keeps_its_justification(self):
        leader = self.make_user("project_leader")
        project = Project.objects.create(title="P", project_code="P-1", funding_type="core_funded", lead=leader)
        budget = self.client_for(leader).post("/api/budget/budgets/", {"project": project.id}, format="json").data

        response = self.client_for(leader).post("/api/budget/line-items/", {
            "budget": budget["id"], "category": "ps", "description": "Honoraria", "unit": "month", "quantity": "18",
            "unit_cost": "8000", "amount": "144000", "justification": "PI honorarium.",
        }, format="json")

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data["justification"], "PI honorarium.")

    def test_line_item_quarters_must_add_up_to_the_amount(self):
        leader = self.make_user("project_leader")
        project = Project.objects.create(title="P", project_code="P-1", funding_type="core_funded", lead=leader)
        budget = self.client_for(leader).post("/api/budget/budgets/", {"project": project.id}, format="json").data

        response = self.client_for(leader).post("/api/budget/line-items/", {
            "budget": budget["id"], "category": "mooe", "description": "Traveling Expenses - Local", "amount": "100",
            "q1_amount": "30", "q2_amount": "30",
        }, format="json")

        self.assertEqual(response.status_code, 400)


class ProjectCampusTests(RMISTestCase):
    """Client 2026-10-27: a project is on its Project Leader's campus, and the campus-scoped roles see only theirs."""

    def setUp(self):
        super().setUp()
        self.admin = self.make_user("system_admin")
        self.leader = self.make_user("project_leader", campus="Siniloan")

    def test_registration_takes_the_leaders_campus(self):
        blank = self.client_for(self.admin).post("/api/projects/", project_payload(self.leader), format="json")
        own = self.client_for(self.leader).post(
            "/api/projects/", project_payload(self.leader, ntp_number="LSPU-2", title="Second", campus="siniloan"),
            format="json",
        )
        other = self.client_for(self.admin).post(
            "/api/projects/", project_payload(self.leader, ntp_number="LSPU-3", title="Third", campus="Sta. Cruz"),
            format="json",
        )

        self.assertEqual((blank.status_code, blank.data["campus"]), (201, "Siniloan"))
        self.assertEqual((own.status_code, own.data["campus"]), (201, "Siniloan"))
        self.assertEqual(other.status_code, 400)
        self.assertEqual(str(other.data["campus"][0]), "The project's campus is its Project Leader's campus (Siniloan).")

    def test_a_leader_without_a_campus_keeps_the_free_choice(self):
        old = self.make_user("project_leader")

        response = self.client_for(self.admin).post("/api/projects/", project_payload(old, campus="Los Baños"), format="json")

        self.assertEqual((response.status_code, response.data["campus"]), (201, "Los Baños"))

    def test_campus_scoped_roles_see_only_their_campus(self):
        mine = Project.objects.create(title="S", project_code="S-1", funding_type="core_funded", lead=self.leader,
                                      campus="Siniloan")
        theirs = Project.objects.create(title="C", project_code="C-1", funding_type="core_funded",
                                        lead=self.make_user("project_leader"), campus="Sta. Cruz")
        for project in (mine, theirs):  # submitted, so the CRC/CRD and Finance see them (Drafts are hidden, D2)
            LineItemBudget.objects.create(project=project, version_number=1, status="submitted")
        from accounts.permissions import scoped_projects

        for code in ("crc_chair", "finance_budget", "riuh", "procurement_officer_lib"):
            with self.subTest(role=code):
                user = self.make_user(code, campus="Siniloan")
                self.assertEqual(list(scoped_projects(user).values_list("project_code", flat=True)), ["S-1"])
        self.assertIsNone(scoped_projects(self.make_user("vprei", campus="Siniloan")))  # university-wide


    def test_every_crc_and_finance_user_of_a_campus_sees_its_proposals_by_status_and_no_other_campus(self):
        """Task 17 (2026-10-30): all CRC/CRD and all Finance users of a campus share its proposals; Drafts stay hidden
        and Finance sees a proposal once the CRC/CRD forwards it (For Budget Review onward, like the LIB list)."""
        def project(code, lib_status, campus="Siniloan"):
            row = Project.objects.create(title=code, project_code=code, funding_type="core_funded", campus=campus,
                                         lead=self.leader if campus == "Siniloan" else self.make_user("project_leader"))
            if lib_status:
                LineItemBudget.objects.create(project=row, version_number=1, status=lib_status)
            return row

        draft, submitted, forwarded, active = project("D-1", "draft"), project("S-1", "submitted"), project("F-1", "endorsed"), project("A-1", "certified")
        elsewhere = project("X-1", "endorsed", campus="Sta. Cruz")
        listed = lambda user: sorted(p["project_code"] for p in self.client_for(user).get("/api/projects/").data)
        detail = lambda user, row: self.client_for(user).get(f"/api/projects/{row.pk}/").status_code

        for crc in (self.make_user("crc_chair", campus="Siniloan"), self.make_user("crc_chair", campus="Siniloan")):
            self.assertEqual(listed(crc), ["A-1", "F-1", "S-1"])
            self.assertEqual((detail(crc, draft), detail(crc, elsewhere)), (404, 404))
        for officer in (self.make_user("finance_budget", campus="Siniloan"), self.make_user("finance_budget", campus="Siniloan")):
            self.assertEqual(listed(officer), ["A-1", "F-1"])
            self.assertEqual((detail(officer, submitted), detail(officer, draft), detail(officer, elsewhere)), (404, 404, 404))
        self.assertEqual(listed(self.make_user("crc_chair", campus="Sta. Cruz")), ["X-1"])
        self.assertEqual(listed(self.leader), ["A-1", "D-1", "F-1", "S-1"])
        self.assertEqual(detail(self.leader, elsewhere), 404)

    def test_project_staff_see_only_their_projects_and_another_leaders_draft_stays_private(self):
        mine = Project.objects.create(title="M", project_code="M-1", funding_type="core_funded", lead=self.leader, campus="Siniloan")
        other_leader = self.make_user("project_leader", campus="Siniloan")
        theirs = Project.objects.create(title="T", project_code="T-1", funding_type="core_funded", lead=other_leader, campus="Siniloan")
        staff = self.make_user("project_staff", campus="Siniloan")
        from research_projects.models import ProjectTeamMember
        ProjectTeamMember.objects.create(project=mine, member_role="member", name="Staff", user=staff)

        self.assertEqual([p["project_code"] for p in self.client_for(staff).get("/api/projects/").data], ["M-1"])
        self.assertEqual(self.client_for(staff).get(f"/api/projects/{theirs.pk}/").status_code, 404)
        self.assertEqual(self.client_for(other_leader).get(f"/api/projects/{mine.pk}/").status_code, 404)

class ExcelImportTests(RMISTestCase):
    def setUp(self):
        super().setUp()
        self.admin = self.make_user("system_admin")
        self.leader = self.make_user("project_leader", email="leader@lspu.test")
        CollegeUnit.objects.create(name="CTE")
        seed_rde_agenda(RdeThrust)
        CooperatingAgency.objects.bulk_create([CooperatingAgency(name="DOST-PCAARRD"), CooperatingAgency(name="LGU of Siniloan, Laguna")])
        template = self.client_for(self.admin).get("/api/projects/import-template/")
        self.assertEqual(template.status_code, 200)
        self.workbook = openpyxl.load_workbook(io.BytesIO(template.content))

    def fill_project(self, **values):
        defaults = {
            "Title": "Bridging Academia and Communities",
            "NTP (Notice to Proceed No.)": "LSPU-X1",
            "Funding Type": "Institutional",
            "Project Leader E-mail Address": "Leader@LSPU.test",
            "Project Leader Gender": "Female",
            "Start Date": datetime.datetime(2026, 1, 1),
            "End Date": datetime.datetime(2026, 12, 31),
            "Total Project/Study Cost": 100000,
            "Sector": "Education, Community Development",
            "New or Continuing": "New Proposal",
            "Basic or Applied Research": "Applied Research",
            "Wet or Dry Research": "Dry Research",
            "Research Priority Area": "Education and Teacher's Training",
            "Research Typology": "Quantitative Research",
            "Sustainable Development Goals (SDGs)": "4, 11, 17",
            "RDE Thrust": RdeThrust.objects.get(code="1").name,
            "RDE Priority Area": RdeThrust.objects.get(code="1.3").name,
            "College Unit - Implementing Unit": "cte",
            # Every Step 1 field is required since 2026-10-29
            "II. Background of the Study": "Background",
            "III. Objectives of the Study": "Map community needs",
            "IV. Project Descriptions/Methodology": "Methodology",
            "VI. Socio-Economic Significance": "Significance",
            "VIII. Monitoring/Evaluation": "Monitoring",
            "IX. List of References": "References",
            "Executive Summary": "Summary",
            "Anticipated Outcomes": "Outcomes",
            "Long-term Potential Impacts": "Impacts",
        }
        defaults.update(values)
        for row in self.workbook["Project"].iter_rows(min_row=2):
            if row[0].value in defaults:
                row[1].value = defaults[row[0].value]

    def upload(self, user=None):
        buffer = io.BytesIO()
        self.workbook.save(buffer)
        buffer.seek(0)
        buffer.name = "registration.xlsx"
        return self.client_for(user or self.admin).post("/api/projects/import/", {"file": buffer}, format="multipart")

    def test_filled_template_registers_the_project_and_its_tables(self):
        self.fill_project()
        self.workbook["Project Team"].append(["Co-Project Leader", "Archieval M. Jain", "Male", None])
        self.workbook["Project Team"].append(["Team Member", "EIU Coordinators", None, None])
        self.workbook["Expected Outputs (6Ps)"].append(["Patent", "Novel technology", 1])
        self.workbook["Target Beneficiaries"].append(["Community beneficiaries", "Farmers", 2000])
        self.workbook["Budget Requirements"].append([2026, "MOOE", "Traveling Expenses - Local", None, "pax", 4, 2500, None])
        self.workbook["Work Plan"].append(["Logistic preparations", datetime.datetime(2026, 1, 1), datetime.datetime(2026, 3, 31)])

        response = self.upload()

        self.assertEqual(response.status_code, 201, response.data)
        project = Project.objects.get(project_code="LSPU-X1")
        self.assertEqual(project.lead, self.leader)
        self.assertEqual(project.sectors, ["community_development", "education"])
        self.assertEqual(project.team_members.count(), 2)
        self.assertEqual(project.expected_outputs.get().category, "patents")
        self.assertEqual(project.target_beneficiaries.get().total, 2000)
        self.assertEqual(LineItem.objects.get(budget__project=project).amount, 10000)
        self.assertEqual(project.milestones.count(), 1)

    def test_excel_registration_takes_the_leaders_campus(self):
        User.objects.filter(pk=self.leader.pk).update(campus="Siniloan")
        self.fill_project(Campus="Sta. Cruz")
        refused = self.upload()
        self.fill_project(Campus=None)

        response = self.upload()

        self.assertEqual(refused.status_code, 400)
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(Project.objects.get(project_code="LSPU-X1").campus, "Siniloan")

    def test_budget_rows_take_the_quarter_columns_of_each_year(self):
        """Client 2026-10-17 (yearly.png): a multi-year project's LIB rows carry Y1 Q1 .. Y3 Q4 amounts."""
        self.fill_project(**{"End Date": datetime.datetime(2027, 3, 31)})
        #                                    FY    Cat     Particulars      Unit  Qty Cost   Total Y1Q1  Q2    Q3    Q4    Y2Q1
        self.workbook["Budget Requirements"].append([2026, "MOOE", "Traveling Expenses - Local", None, "lot", 1, 10000, None, 5000, None, 2500, None, 2500])
        self.workbook["Budget Requirements"].append([2026, "MOOE", "Office Supplies Expenses", None, "lot", 1, 1000, None, 900])

        response = self.upload()

        self.assertEqual(response.status_code, 400)
        self.assertEqual((response.data["errors"][0]["row"], response.data["errors"][0]["field"]), (3, "Y1 Q1 - Y3 Q4"))
        self.workbook["Budget Requirements"].delete_rows(3)
        response = self.upload()
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual([(a["period"], a["amount"]) for a in LineItem.objects.get().allocations], [
            ("Y1Q1", "5000.00"), ("Y1Q2", "0.00"), ("Y1Q3", "2500.00"), ("Y1Q4", "0.00"), ("Y2Q1", "2500.00"),
        ])

    def test_one_year_project_reads_the_year_1_quarters_as_q1_to_q4(self):
        self.fill_project(**{"End Date": datetime.datetime(2026, 12, 31)})
        self.workbook["Budget Requirements"].append([2026, "MOOE", "Traveling Expenses - Local", None, "lot", 1, 10000, None, None, 6000, None, 4000])

        response = self.upload()

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual([(a["period"], a["amount"]) for a in LineItem.objects.get().allocations], [
            ("Q1", "0.00"), ("Q2", "6000.00"), ("Q3", "0.00"), ("Q4", "4000.00"),
        ])

    def test_college_unit_must_be_in_the_admin_list_and_fills_both_fields(self):
        self.fill_project()
        self.assertEqual(self.upload().status_code, 201)
        project = Project.objects.get()
        self.assertEqual((project.college, project.implementing_unit), ("CTE", "CTE"))

        self.fill_project(**{"NTP (Notice to Proceed No.)": "LSPU-X2", "College Unit - Implementing Unit": "CCS"})
        response = self.upload()
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.data["errors"][0]["field"], "College Unit - Implementing Unit")

    def test_rde_thrust_and_agencies_must_be_in_the_admin_lists(self):
        self.fill_project(**{
            "RDE Thrust": "7 - food, nutrition and health", "RDE Priority Area": "7.4 - Food Safety and Security",
            "Cooperating Agency/ies": "dost-pcaarrd; LGU of Siniloan, Laguna",
        })
        self.assertEqual(self.upload().status_code, 201)
        project = Project.objects.get()
        self.assertEqual(
            (project.rde_thrust, project.rde_priority), ("7 - Food, Nutrition and Health", "7.4 - Food Safety and Security"),
        )
        self.assertEqual(project.cooperating_agencies, "DOST-PCAARRD, LGU of Siniloan, Laguna")

        self.fill_project(**{"NTP (Notice to Proceed No.)": "LSPU-X2", "RDE Thrust": "Space", "Cooperating Agency/ies": "DOST-PCAARRD; CHED"})
        fields = sorted(e["field"] for e in self.upload().data["errors"])
        self.assertEqual(fields, ["Cooperating Agency/ies", "RDE Thrust"])

    def test_template_lists_the_rde_thrusts_only_on_the_thrust_row(self):
        notes = {row[0].value: row[2].value for row in self.workbook["Project"].iter_rows(min_row=2)}

        self.assertIn("8 - Philippine Global Competitiveness", notes["RDE Thrust"])
        self.assertNotIn("1.1 -", notes["RDE Thrust"])

    def test_template_lists_the_college_units(self):
        notes = {row[0].value: row[2].value for row in self.workbook["Project"].iter_rows(min_row=2)}
        self.assertIn("CTE", notes["College Unit - Implementing Unit"])

    def test_legend_sheet_explains_each_budget_category_with_lib_item_examples(self):
        # Client 2026-10-23: the Category column of Budget Requirements needs a legend; the codes themselves stay.
        rows = [[c for c in row if c is not None] for row in self.workbook["Legend"].iter_rows(values_only=True)]
        by_code = {row[0]: row for row in rows if row and row[0] in ("PS", "MOOE", "CO")}

        self.assertEqual(set(by_code), {"PS", "MOOE", "CO"})
        self.assertIn("Personal Services", by_code["PS"][1])
        self.assertIn("Equipment Outlay", by_code["CO"][2])
        self.assertTrue(by_code["CO"][2].startswith("CO, "))
        self.assertIn("Honoraria", by_code["PS"][3])
        self.assertIn("Traveling Expenses - Local", by_code["MOOE"][3])
        self.assertIn("Office Equipment", by_code["CO"][3])
        self.assertTrue(any("Particulars" in str(row) for row in rows))

    def test_new_system_admin_choices_show_in_the_next_template(self):
        """Client 2026-10-29 (Task 11): the template reads the System Admin lists on each download."""
        CollegeUnit.objects.create(name="CNEW - College of New Things")
        thrust = RdeThrust.objects.get(code="1")
        RdeThrust.objects.create(name="1.9 - Brand New Priority", code="1.9", title="Brand New Priority", parent=thrust)

        template = self.client_for(self.admin).get("/api/projects/import-template/")
        workbook = openpyxl.load_workbook(io.BytesIO(template.content))
        notes = {row[0]: row[2] for row in workbook["Project"].iter_rows(min_row=2, values_only=True)}
        legend = {row[0]: row[1] for row in workbook["Legend"].iter_rows(values_only=True) if row[0]}

        self.assertIn("CNEW - College of New Things", notes["College Unit - Implementing Unit"])
        self.assertIn("1.9 - Brand New Priority", legend[thrust.name])
        self.assertIn("1.3 - ", legend[thrust.name])

    def test_campus_row_lists_the_account_campuses_and_refuses_others(self):
        notes = {row[0].value: row[2].value for row in self.workbook["Project"].iter_rows(min_row=2)}
        for campus in ("San Pablo City", "Siniloan", "Los Baños", "Sta. Cruz"):
            self.assertIn(campus, notes["Campus"])

        self.fill_project(Campus="Calamba")
        response = self.upload()

        self.assertEqual(response.status_code, 400)
        self.assertTrue(any(e["field"] == "Campus" for e in response.data["errors"]), response.data)

    def test_campus_is_read_case_insensitively(self):
        leader = User.objects.get(email__iexact="leader@lspu.test")
        User.objects.filter(pk=leader.pk).update(campus="Sta. Cruz")
        self.fill_project(Campus="sta. cruz")

        response = self.upload()

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(Project.objects.get().campus, "Sta. Cruz")

    def test_legend_lists_only_active_lib_items(self):
        LibItem.objects.filter(name="Honoraria").update(is_active=False)

        template = self.client_for(self.admin).get("/api/projects/import-template/")
        legend = openpyxl.load_workbook(io.BytesIO(template.content))["Legend"]

        self.assertNotIn("Honoraria", " ".join(str(c) for row in legend.iter_rows(values_only=True) for c in row if c))

    def test_category_header_points_to_the_legend(self):
        self.assertIn("Legend", self.workbook["Budget Requirements"]["B1"].comment.text)

    def test_template_with_the_legend_still_imports_every_category_spelling(self):
        self.fill_project()
        for category, item in (("PS", "Honoraria"), ("Maintenance and Other Operating Expenses", "Training Expenses"),
                               ("Equipment Outlay", "Office Equipment"), ("CO", "Books")):
            self.workbook["Budget Requirements"].append([2026, category, item, None, "lot", 1, 1000, None])

        response = self.upload()

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(sorted(LineItem.objects.values_list("category", flat=True)), ["co", "co", "mooe", "ps"])

    def test_not_applicable_research_classification_imports(self):
        # Client 2026-10-23: Wet/Dry and Basic/Applied may both be "Not Applicable"; a blank Wet/Dry is still Dry.
        self.fill_project(**{"Wet or Dry Research": "Not Applicable", "Basic or Applied Research": "Not Applicable"})

        response = self.upload()

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual((response.data["is_dry_research"], response.data["research_type"]), (None, "not_applicable"))

    def test_blank_wet_or_dry_new_or_continuing_and_funding_type_must_be_filled(self):
        # Client 2026-10-29: no more "Blank = Dry", "Blank = New Proposal" or "Blank = Institutional"
        self.fill_project(**{"Wet or Dry Research": None, "New or Continuing": None})
        choices = self.upload()
        self.fill_project(**{"Funding Type": None})
        funding = self.upload()

        self.assertEqual(choices.status_code, 400)
        self.assertTrue({"Wet or Dry Research", "New or Continuing"} <= {e["field"] for e in choices.data["errors"]}, choices.data)
        self.assertEqual(funding.status_code, 400)
        self.assertIn("Funding Type", {e["field"] for e in funding.data["errors"]})
        self.assertFalse(Project.objects.exists())

    def test_an_old_template_with_only_the_project_code_is_refused(self):
        self.fill_project(**{"NTP (Notice to Proceed No.)": None})
        sheet = self.workbook["Project"]
        sheet.append(["LSPU Faculty Research Number (Project Code)", "LSPU-OLD"])

        response = self.upload()

        self.assertEqual(response.status_code, 400)
        messages = {e["field"]: e["message"] for e in response.data["errors"]}
        self.assertIn("no longer used", messages["LSPU Faculty Research Number (Project Code)"])
        self.assertFalse(Project.objects.exists())

    def test_a_project_code_row_with_a_different_ntp_is_refused_and_the_same_one_is_fine(self):
        self.fill_project()
        self.workbook["Project"].append(["LSPU Faculty Research Number (Project Code)", "OTHER-1"])
        self.assertEqual(self.upload().status_code, 400)

        self.workbook["Project"].cell(row=self.workbook["Project"].max_row, column=2).value = "lspu-x1"
        response = self.upload()

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(Project.objects.get().project_code, "LSPU-X1")

    def test_an_excel_ntp_is_saved_uppercase(self):
        self.fill_project(**{"NTP (Notice to Proceed No.)": "lspu-rd-grant-sin-2024-001"})

        response = self.upload()

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(Project.objects.values_list("ntp_number", "project_code").get(),
                         ("LSPU-RD-GRANT-SIN-2024-001", "LSPU-RD-GRANT-SIN-2024-001"))

    def test_the_template_has_the_ntp_row_and_not_the_dropped_rows(self):
        labels = [row[0].value for row in self.workbook["Project"].iter_rows(min_row=2)]

        self.assertIn("NTP (Notice to Proceed No.)", labels)
        self.assertIn("III. Objectives of the Study", labels)
        self.assertNotIn("LSPU Faculty Research Number (Project Code)", labels)
        for dropped in ("Research Priority Area", "II. Background of the Study", "IV. Project Descriptions/Methodology",
                        "VI. Socio-Economic Significance", "VIII. Monitoring/Evaluation", "IX. List of References",
                        "Executive Summary", "Anticipated Outcomes", "Long-term Potential Impacts",
                        # Task 13/14 (2026-10-29): Annex A and its dates left the registration
                        "Date Submitted (Project Leader)", "Endorsed By (Dean/Associate Dean)", "Date Endorsed by Dean",
                        "Noted By (RDS Director/Chairperson)", "Recommending Approval (Campus Director)",
                        "Recommending Approval (VPRDE)", "Approved By (University President)", "Date Approved by University President"):
            self.assertNotIn(dropped, labels)

    def test_an_older_template_still_saves_the_dropped_rows(self):
        self.fill_project()
        for label, value in (("Research Priority Area", "Health and Health Profession"),
                             ("II. Background of the Study", "Old background"), ("Executive Summary", "Old summary")):
            self.workbook["Project"].append([label, value])

        response = self.upload()

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(Project.objects.values_list("research_priority_area", "background", "description").get(),
                         ("health", "Old background", "Old summary"))

    def test_budget_rows_keep_their_specification_and_the_preview_shows_it(self):
        """Task 14 (2026-10-29): Budget Requirements has Specification (optional), as Manual Entry and the LIB template do."""
        header = [c.value for c in self.workbook["Budget Requirements"][1]]
        self.assertEqual(header[:5], ["Fiscal Year", "Category", "Particulars", "Specification (optional)", "Unit"])
        self.fill_project()
        self.workbook["Budget Requirements"].append([2026, "MOOE", "Office Supplies Expenses", "PAPER, 80 gsm, A4", "unit", 10, 225, None, 2250])
        buffer = io.BytesIO()
        self.workbook.save(buffer)
        buffer.seek(0)
        buffer.name = "registration.xlsx"

        response = self.client_for(self.admin).post("/api/projects/import/?dry_run=1", {"file": buffer}, format="multipart")

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["draft"]["line_items"][0]["specification"], "PAPER, 80 gsm, A4")
        self.assertEqual((response.data["preview"]["budget"][0]["description"], response.data["preview"]["budget"][0]["specification"]),
                         ("Office Supplies Expenses", "PAPER, 80 gsm, A4"))

    def test_an_older_template_without_the_specification_column_still_imports(self):
        self.fill_project()
        self.workbook["Budget Requirements"].delete_cols(4)
        self.workbook["Budget Requirements"].append([2026, "MOOE", "Traveling Expenses - Local", "pax", 4, 2500, None])

        response = self.upload()

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(LineItem.objects.values_list("description", "specification", "amount").get(),
                         ("Traveling Expenses - Local", "", Decimal("10000.00")))

    def work_plan(self):
        return list(WorkPlanMilestone.objects.order_by("id").values_list("title", "start_date", "target_date"))

    def test_work_plan_rows_keep_their_order_activities_and_dates(self):
        # Client 2026-10-23: the uploaded Work Plan must match the template, so nobody re-encodes it.
        self.fill_project()
        sheet = self.workbook["Work Plan"]
        self.assertEqual([c.value for c in sheet[1]], ["Activity", "Start Date", "End Date"])
        sheet.append(["Logistic preparations", datetime.datetime(2026, 1, 1), datetime.datetime(2026, 3, 31)])
        sheet.append(["Data gathering", "2026-04-01", "2026-06-30"])
        sheet.append(["Data analysis", "07/01/2026", "09/30/2026"])
        sheet.append(["Final report", "October 1, 2026", "Dec 15, 2026"])

        response = self.upload()

        self.assertEqual(response.status_code, 201, response.data)
        d = datetime.date
        self.assertEqual(self.work_plan(), [
            ("Logistic preparations", d(2026, 1, 1), d(2026, 3, 31)),
            ("Data gathering", d(2026, 4, 1), d(2026, 6, 30)),
            ("Data analysis", d(2026, 7, 1), d(2026, 9, 30)),
            ("Final report", d(2026, 10, 1), d(2026, 12, 15)),
        ])

    def test_an_old_template_with_target_date_still_imports(self):
        self.fill_project()
        sheet = self.workbook["Work Plan"]
        sheet["C1"] = "Target Date"
        sheet.append(["Logistic preparations", "2026-01-01", "2026-03-31"])

        self.assertEqual(self.upload().status_code, 201)
        self.assertEqual(self.work_plan()[0][2], datetime.date(2026, 3, 31))

    def test_work_plan_errors_name_the_row_and_save_nothing(self):
        self.fill_project()
        sheet = self.workbook["Work Plan"]
        sheet.append(["Logistic preparations", "2026-01-01", "2026-03-31"])
        sheet.append(["Data gathering", "2026-04-01", None])
        sheet.append(["Data analysis", "31/12/2026", "2026-12-31"])
        sheet.append(["Logistic preparations", "2026-01-01", "2026-03-31"])

        response = self.upload()

        self.assertEqual(response.status_code, 400)
        errors = {(e["row"], e["field"]): e["message"] for e in response.data["errors"] if e["sheet"] == "Work Plan"}
        self.assertEqual(errors[(3, "End Date")], "End Date is required.")
        self.assertIn("is not a date", errors[(4, "Start Date")])
        self.assertIn("row 2", errors[(5, "Activity")])
        self.assertFalse(Project.objects.exists())

    def test_an_activity_ending_before_it_starts_names_its_row(self):
        self.fill_project()
        self.workbook["Work Plan"].append(["Logistic preparations", "2026-01-01", "2026-03-31"])
        self.workbook["Work Plan"].append(["Reporting", "06/01/2026", "05/01/2026"])

        response = self.upload()

        self.assertEqual(response.status_code, 400)
        self.assertEqual([(e["sheet"], e["row"], e["field"]) for e in response.data["errors"]], [("Work Plan", 3, "Start Date")])
        self.assertFalse(Project.objects.exists())

    def test_the_preview_shows_the_work_plan_that_would_be_saved(self):
        self.fill_project()
        self.workbook["Work Plan"].append(["Data analysis", "07/01/2026", "09/30/2026"])
        buffer = io.BytesIO()
        self.workbook.save(buffer)
        buffer.seek(0)
        buffer.name = "registration.xlsx"

        response = self.client_for(self.admin).post("/api/projects/import/?dry_run=1", {"file": buffer}, format="multipart")

        self.assertEqual(response.data["preview"]["work_plan"], [{"title": "Data analysis", "start_date": "2026-07-01", "target_date": "2026-09-30"}])

    def test_an_older_templates_annex_a_rows_still_become_endorser_rows(self):
        # Task 13/14 (2026-10-29): Annex A is no longer in the template; an older file's rows are still read and saved
        self.fill_project()
        for label, value in (("Endorsed By (Dean/Associate Dean)", "Adriel G. Roman"), ("Recommending Approval (VPRDE)", "Robert C. Agatep"),
                             ("Date Approved by University President", datetime.datetime(2025, 12, 1))):
            self.workbook["Project"].append([label, value])

        response = self.upload()

        self.assertEqual(response.status_code, 201, response.data)
        rows = list(ProjectEndorser.objects.order_by("id").values_list("name", "designation", "role_code"))
        self.assertEqual(rows, [
            ("Adriel G. Roman", "Dean/Associate Dean", ""),
            ("Robert C. Agatep", "Vice President for Research, Development and Extension", "vprei"),
        ])
        self.assertEqual(Project.objects.get().proposal_approved_on, datetime.date(2025, 12, 1))

    def test_any_bad_row_rolls_back_the_whole_import_and_lists_the_error(self):
        self.fill_project()
        self.workbook["Expected Outputs (6Ps)"].append(["Not a 6P", "x", 1])

        response = self.upload()

        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.data["errors"][0]["sheet"], "Expected Outputs (6Ps)")
        self.assertFalse(Project.objects.exists())

    def test_dry_run_returns_the_preview_and_saves_nothing(self):
        """Client meeting 2026-10-01 (#9): the Excel mode View button previews the SF-018 without registering."""
        self.fill_project(**{"III. Objectives of the Study": "1. Train teachers\n2. Map communities"})
        self.workbook["Project Team"].append(["Co-Project Leader", "Archieval M. Jain", "Male", None])
        self.workbook["Budget Requirements"].append([2026, "MOOE", "Traveling Expenses - Local", None, "pax", 4, 2500, None])
        counts = lambda: (Project.objects.count(), ProjectEndorser.objects.count(), LineItem.objects.count())
        before = counts()

        buffer = io.BytesIO()
        self.workbook.save(buffer)
        buffer.seek(0)
        buffer.name = "registration.xlsx"
        response = self.client_for(self.leader).post("/api/projects/import/?dry_run=1", {"file": buffer}, format="multipart")

        self.assertEqual(response.status_code, 200, response.data)
        preview = response.data["preview"]
        self.assertEqual(preview["title"], "Bridging Academia and Communities")
        self.assertEqual(preview["lead_email"], "leader@lspu.test")
        self.assertEqual(preview["co_leaders"], [{"name": "Archieval M. Jain", "gender": "male"}])
        self.assertEqual(preview["objectives"], ["Train teachers", "Map communities"])
        self.assertEqual((preview["rde_thrust"], preview["rde_priority"]),
                         (RdeThrust.objects.get(code="1").name, RdeThrust.objects.get(code="1.3").name))
        self.assertEqual(preview["budget"][0], {
            "category": "mooe", "description": "Traveling Expenses - Local", "specification": "", "unit": "pax", "quantity": 4, "unit_cost": 2500, "total": 10000,
            "allocations": [],
        })
        self.assertEqual(counts(), before)

    def test_dry_run_returns_an_editable_draft_and_saves_nothing(self):
        """Client 2026-10-23: an uploaded workbook can be reviewed and edited in Manual Entry before + Register Project."""
        self.fill_project()
        self.workbook["Project"].append(["Endorsed By (Dean/Associate Dean)", "Adriel G. Roman"])
        self.workbook["Project Team"].append(["Co-Project Leader", "Archieval M. Jain", "Male", None])
        self.workbook["Study Components"].append(["Study on soils", None])
        self.workbook["Expected Outputs (6Ps)"].append(["Patent", "Novel technology", 1])
        self.workbook["Target Beneficiaries"].append(["Farmers", "Rice farmers", 200])
        self.workbook["Budget Requirements"].append([2026, "MOOE", "Traveling Expenses - Local", None, "pax", 4, 2500, None])
        self.workbook["Work Plan"].append(["Logistic preparations", "2026-01-01", "2026-03-31"])
        counts = lambda: (Project.objects.count(), ProjectEndorser.objects.count(), LineItem.objects.count(), WorkPlanMilestone.objects.count())
        before = counts()
        buffer = io.BytesIO()
        self.workbook.save(buffer)
        buffer.seek(0)
        buffer.name = "registration.xlsx"

        response = self.client_for(self.leader).post("/api/projects/import/?dry_run=1", {"file": buffer}, format="multipart")

        self.assertEqual(response.status_code, 200, response.data)
        draft = response.data["draft"]
        self.assertEqual((draft["project"]["project_code"], draft["project"]["lead"]), ("LSPU-X1", self.leader.pk))
        self.assertEqual(draft["project"]["sectors"], ["community_development", "education"])
        self.assertEqual([(t["member_role"], t["name"], t["gender"]) for t in draft["team"]], [("co_leader", "Archieval M. Jain", "male")])
        self.assertEqual([s["title"] for s in draft["studies"]], ["Study on soils"])
        self.assertEqual([(o["category"], o["description"], o["target_count"]) for o in draft["outputs"]], [("patents", "Novel technology", 1)])
        self.assertEqual([(b["group"], b["total"]) for b in draft["beneficiaries"]], [("Farmers", 200)])
        self.assertEqual([(m["title"], m["start_date"], m["target_date"]) for m in draft["milestones"]], [("Logistic preparations", "2026-01-01", "2026-03-31")])
        self.assertEqual([(e["name"], e["designation"]) for e in draft["endorsers"]], [("Adriel G. Roman", "Dean/Associate Dean")])
        item = draft["line_items"][0]
        self.assertEqual((item["category"], item["description"], item["unit"], item["quantity"], item["unit_cost"], item["amount"]),
                         ("mooe", "Traveling Expenses - Local", "pax", "4.00", "2500.00", "10000.00"))
        self.assertIn("preview", response.data)
        self.assertEqual(counts(), before)

    def test_dry_run_with_errors_lists_them_like_a_real_upload(self):
        self.fill_project()
        self.workbook["Expected Outputs (6Ps)"].append(["Not a 6P", "x", 1])
        buffer = io.BytesIO()
        self.workbook.save(buffer)
        buffer.seek(0)
        buffer.name = "registration.xlsx"

        response = self.client_for(self.admin).post("/api/projects/import/?dry_run=1", {"file": buffer}, format="multipart")

        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.data["errors"][0]["sheet"], "Expected Outputs (6Ps)")
        self.assertFalse(Project.objects.exists())

    def test_a_non_workbook_file_is_rejected_cleanly(self):
        buffer = io.BytesIO(b"not an excel file")
        buffer.name = "notes.xlsx"

        response = self.client_for(self.admin).post("/api/projects/import/", {"file": buffer}, format="multipart")

        self.assertEqual(response.status_code, 400)

    def test_project_leader_can_import_a_project_they_lead(self):
        """Client meeting 2026-10-01: the Excel import follows the same rule as manual entry."""
        self.fill_project()

        response = self.upload(self.leader)

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(Project.objects.get().lead, self.leader)

    def test_project_leader_cannot_import_another_leaders_project(self):
        self.fill_project()
        other = self.make_user("project_leader")

        response = self.upload(other)

        self.assertEqual(response.status_code, 400)
        self.assertFalse(Project.objects.exists())

    def test_roles_without_projects_register_cannot_use_the_excel_import(self):
        self.fill_project()

        for code in ("crc_chair", "project_staff"):
            self.assertEqual(self.upload(self.make_user(code)).status_code, 403)
        self.assertFalse(Project.objects.exists())


    def test_a_workbook_with_an_existing_title_is_refused(self):
        """Client 2026-10-24: an existing title can't be registered again, from Excel either."""
        Project.objects.create(title="Bridging Academia and Communities", project_code="OLD-1", funding_type="core_funded",
                               lead=self.make_user("project_leader"), sectors=["education"], sdgs=[4])
        self.fill_project(Title="  bridging academia   and COMMUNITIES ")

        response = self.upload()

        self.assertEqual(response.status_code, 400)
        self.assertEqual([e["field"] for e in response.data["errors"]], ["Title"])
        self.assertFalse(Project.objects.filter(project_code="LSPU-X1").exists())


class StaffIdExcelTests(RMISTestCase):
    """Client 2026-10-28: the Staff ID picks the Project Leader and Project Team accounts; the e-mail is the fallback
    for older templates. Blank Gender/Contact No./Name cells are filled from the account's profile."""

    fill_project = ExcelImportTests.fill_project
    upload = ExcelImportTests.upload

    def setUp(self):
        super().setUp()
        self.admin = self.make_user("system_admin")
        self.leader = self.make_user("project_leader", email="leader@lspu.test")
        CollegeUnit.objects.create(name="CTE")
        seed_rde_agenda(RdeThrust)
        self.workbook = openpyxl.load_workbook(io.BytesIO(self.client_for(self.admin).get("/api/projects/import-template/").content))
        User.objects.filter(pk=self.leader.pk).update(
            staff_id="012345", gender="female", mobile_number="09171234567", first_name="Lea", last_name="Dimaano",
        )
        self.staff = self.make_user("project_staff", staff_id="222222", gender="male", first_name="Ben", last_name="Reyes")

    def by_staff_id(self, value, **values):
        self.fill_project(**{"Project Leader Staff ID": value, "Project Leader E-mail Address": None,
                             "Project Leader Gender": None, **values})

    def dry_run(self):
        buffer = io.BytesIO()
        self.workbook.save(buffer)
        buffer.seek(0)
        buffer.name = "registration.xlsx"
        return self.client_for(self.admin).post("/api/projects/import/?dry_run=1", {"file": buffer}, format="multipart")

    def lead_error(self):
        response = self.upload()
        self.assertEqual(response.status_code, 400)
        return [e["message"] for e in response.data["errors"] if e["field"] == "Project Leader Staff ID"]

    def test_staff_ids_pick_the_accounts_and_fill_their_profiles(self):
        self.by_staff_id("012345")
        self.workbook["Project Team"].append(["Team Member", None, None, None, "222222"])

        response = self.upload()

        self.assertEqual(response.status_code, 201, response.data)
        project = Project.objects.get()
        self.assertEqual((project.lead, project.lead_gender, project.contact_number), (self.leader, "female", "09171234567"))
        member = project.team_members.get()
        self.assertEqual((member.user, member.name, member.gender), (self.staff, "Ben Reyes", "male"))

    def test_number_cells_and_spaces_are_read_as_typed(self):
        User.objects.filter(pk=self.leader.pk).update(staff_id="123456")
        for value in [123456, 123456.0, " 123456 "]:
            self.by_staff_id(value)
            response = self.dry_run()
            self.assertEqual(response.status_code, 200, (value, response.data))
            self.assertEqual(response.data["preview"]["lead_email"], "leader@lspu.test")

    def test_filled_cells_stay_the_projects_own_values(self):
        self.by_staff_id("012345", **{"Project Leader Gender": "Male", "Contact No./s.": "0998 164 9081"})
        self.workbook["Project Team"].append(["Team Member", "B. Reyes", "Female", None, "222222"])

        self.assertEqual(self.upload().status_code, 201)
        project = Project.objects.get()
        self.assertEqual((project.lead_gender, project.contact_number), ("male", "0998 164 9081"))
        self.assertEqual(project.team_members.values_list("name", "gender").get(), ("B. Reyes", "female"))

    def test_bad_or_unknown_staff_ids_are_errors_without_guessing(self):
        self.by_staff_id(12345)  # 012345 typed in a Number cell: the zero is gone
        self.assertIn("format the cell as Text", self.lead_error()[0])
        self.by_staff_id("12345a")
        self.assertIn("Staff ID must be 6 digits", self.lead_error()[0])
        self.by_staff_id("999999")
        self.assertIn("No RMIS account with Staff ID '999999'", self.lead_error()[0])
        other = self.make_user("project_leader", email="other@lspu.test")
        self.by_staff_id("012345", **{"Project Leader E-mail Address": other.email})
        self.assertEqual(self.lead_error(), ["The Staff ID and the e-mail belong to two different RMIS accounts."])
        self.assertFalse(Project.objects.exists())

    def test_team_staff_id_and_e_mail_must_be_the_same_account(self):
        self.fill_project()
        self.workbook["Project Team"].append(["Team Member", "Ben", None, self.leader.email, "222222"])

        response = self.upload()

        self.assertEqual(response.status_code, 400)
        self.assertEqual([(e["sheet"], e["row"], e["field"]) for e in response.data["errors"]],
                         [("Project Team", 2, "Staff ID (optional)")])

    def test_same_staff_id_and_e_mail_is_fine(self):
        self.by_staff_id("012345", **{"Project Leader E-mail Address": "leader@lspu.test"})
        self.assertEqual(self.upload().status_code, 201)

    def test_older_template_without_staff_id_columns_still_uses_the_e_mail(self):
        sheet = self.workbook["Project"]
        sheet.delete_rows(next(r.row for r in sheet["A"] if r.value == "Project Leader Staff ID"))
        self.workbook["Project Team"].delete_cols(5)
        self.fill_project(**{"Project Leader Gender": None})
        self.workbook["Project Team"].append(["Team Member", None, None, self.staff.email])

        response = self.upload()

        self.assertEqual(response.status_code, 201, response.data)
        project = Project.objects.get()
        self.assertEqual((project.lead, project.lead_gender), (self.leader, "female"))
        self.assertEqual(project.team_members.values_list("user", "name").get(), (self.staff.pk, "Ben Reyes"))

    def test_template_cells_for_staff_ids_are_text(self):
        sheet = self.workbook["Project"]
        row = next(r.row for r in sheet["A"] if r.value == "Project Leader Staff ID")
        self.assertEqual(sheet.cell(row=row, column=2).number_format, "@")
        self.assertEqual(self.workbook["Project Team"]["E1"].value, "Staff ID (optional)")

    def test_preview_shows_the_account_the_staff_id_picked(self):
        self.by_staff_id("012345")
        self.workbook["Project Team"].append(["Co-Project Leader", None, None, None, "222222"])

        response = self.dry_run()

        self.assertEqual(response.status_code, 200, response.data)
        preview = response.data["preview"]
        self.assertEqual((preview["lead_email"], preview["lead_gender"], preview["contact_number"]),
                         ("leader@lspu.test", "female", "09171234567"))
        self.assertEqual(preview["co_leaders"], [{"name": "Ben Reyes", "gender": "male"}])
        self.assertFalse(Project.objects.exists())
        self.assertEqual(self.upload().status_code, 201)
        self.assertEqual(Project.objects.get().team_members.get().user, self.staff)


class WorkPlanExcelTests(RMISTestCase):
    """Client 2026-10-23: Manual Entry Step 5 fills XI. Work Plan from the template's Work Plan sheet; nothing is saved."""

    URL = "/api/projects/import-work-plan/"

    def setUp(self):
        super().setUp()
        self.admin = self.make_user("system_admin")
        self.workbook = openpyxl.Workbook()
        self.sheet = self.workbook.active
        self.sheet.title = "Work Plan"
        self.sheet.append(["Activity", "Start Date", "End Date"])

    def post(self, user=None, workbook=None):
        buffer = io.BytesIO()
        (workbook or self.workbook).save(buffer)
        buffer.seek(0)
        buffer.name = "work_plan.xlsx"
        return self.client_for(user or self.admin).post(self.URL, {"file": buffer}, format="multipart")

    def test_rows_come_back_in_order_with_their_dates_and_nothing_is_saved(self):
        self.sheet.append(["Logistic preparations", datetime.datetime(2026, 1, 1), datetime.datetime(2026, 3, 31)])
        self.sheet.append(["Data analysis", "07/01/2026", "September 30, 2026"])
        self.sheet.append(["Reporting", None, "2026-12-15"])

        response = self.post()

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["rows"], [
            {"title": "Logistic preparations", "start_date": "2026-01-01", "target_date": "2026-03-31"},
            {"title": "Data analysis", "start_date": "2026-07-01", "target_date": "2026-09-30"},
            {"title": "Reporting", "start_date": "", "target_date": "2026-12-15"},
        ])
        self.assertFalse(WorkPlanMilestone.objects.exists())

    def test_the_work_plan_template_has_only_the_work_plan_sheet_and_imports(self):
        """Client 2026-10-29 (Task 10): a Work Plan-only template for Manual Entry, same columns as the full one."""
        response = self.client_for(self.admin).get("/api/projects/work-plan-template/")
        template = openpyxl.load_workbook(io.BytesIO(response.content))

        self.assertEqual(response.status_code, 200)
        self.assertIn("rmis_work_plan_template.xlsx", response["Content-Disposition"])
        self.assertEqual(template.sheetnames, ["Work Plan"])
        self.assertEqual([c.value for c in template["Work Plan"][1]], ["Activity", "Start Date", "End Date"])

        template["Work Plan"].append(["Logistic preparations", "2026-01-01", "2026-03-31"])
        imported = self.post(workbook=template)

        self.assertEqual(imported.status_code, 200, imported.data)
        self.assertEqual(imported.data["rows"], [{"title": "Logistic preparations", "start_date": "2026-01-01", "target_date": "2026-03-31"}])
        self.assertFalse(WorkPlanMilestone.objects.exists())

    def test_the_work_plan_template_is_for_registrants_only(self):
        staff = self.client_for(self.make_user("project_staff")).get("/api/projects/work-plan-template/")
        leader = self.client_for(self.make_user("project_leader")).get("/api/projects/work-plan-template/")

        self.assertEqual((staff.status_code, leader.status_code), (403, 200))

    def test_the_full_registration_template_works_too(self):
        template = openpyxl.load_workbook(io.BytesIO(self.client_for(self.admin).get("/api/projects/import-template/").content))
        template["Work Plan"].append(["Logistic preparations", "2026-01-01", "2026-03-31"])

        response = self.post(workbook=template)

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual([r["title"] for r in response.data["rows"]], ["Logistic preparations"])

    def test_errors_name_the_row(self):
        self.sheet.append(["Logistic preparations", "2026-01-01", "2026-03-31"])
        self.sheet.append(["Data gathering", "2026-04-01", None])
        self.sheet.append(["Data analysis", "31/12/2026", "2026-12-31"])
        self.sheet.append(["Logistic preparations", "2026-01-01", "2026-03-31"])
        self.sheet.append(["Reporting", "2026-06-01", "2026-05-01"])
        self.sheet.append([None, "2026-06-01", "2026-06-30"])

        response = self.post()

        self.assertEqual(response.status_code, 400)
        fields = sorted((e["row"], e["field"]) for e in response.data["errors"])
        self.assertEqual(fields, [(3, "End Date"), (4, "Start Date"), (5, "Activity"), (6, "Start Date"), (7, "Activity")])

    def test_a_workbook_without_the_work_plan_sheet_or_not_a_workbook_is_rejected(self):
        self.sheet.title = "Sheet1"
        self.assertEqual(self.post().status_code, 400)

        buffer = io.BytesIO(b"not an excel file")
        buffer.name = "notes.xlsx"
        self.assertEqual(self.client_for(self.admin).post(self.URL, {"file": buffer}, format="multipart").status_code, 400)

    def test_only_registrants_can_use_it(self):
        self.sheet.append(["Logistic preparations", "2026-01-01", "2026-03-31"])

        self.assertEqual(self.post(self.make_user("project_leader")).status_code, 200)
        for code in ("project_staff", "crc_chair", "finance_budget"):
            self.assertEqual(self.post(self.make_user(code)).status_code, 403, code)



class LibExcelTests(RMISTestCase):
    """Task 12 (2026-10-29): Step 4 Budget (LIB) has its own Excel template and import, separate from Upload via Excel.
    The periods follow the Step 1 Start and End Dates exactly; the import is all-or-nothing and saves nothing."""

    TEMPLATE_URL = "/api/projects/lib-template/"
    PARSE_URL = "/api/projects/parse-lib/"
    FIXED = ["Particulars", "Item", "Specification (optional)", "Unit", "Qty", "Unit Cost (₱)", "Total (₱)"]
    ONE_YEAR = {"start_date": "2027-01-01", "target_end_date": "2027-12-31"}

    def setUp(self):
        super().setUp()
        self.admin = self.make_user("system_admin")

    def template(self, dates=None, user=None):
        return self.client_for(user or self.admin).get(self.TEMPLATE_URL, dates or self.ONE_YEAR)

    def workbook(self, dates=None):
        return openpyxl.load_workbook(io.BytesIO(self.template(dates).content))

    def header_row(self, sheet):
        return next(r for r in range(1, 20) if sheet.cell(row=r, column=1).value == "Particulars")

    def fill(self, sheet, rows):
        first = self.header_row(sheet) + 1
        for offset, values in enumerate(rows):
            for column, value in enumerate(values, start=1):
                sheet.cell(row=first + offset, column=column, value=value)
        return first

    def parse(self, workbook, dates=None, user=None):
        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        buffer.name = "lib.xlsx"
        return self.client_for(user or self.admin).post(self.PARSE_URL, {"file": buffer, **(dates or self.ONE_YEAR)}, format="multipart")

    def test_the_template_is_one_lib_sheet_with_the_projects_periods_and_a_legend_below(self):
        response = self.template()
        sheet = openpyxl.load_workbook(io.BytesIO(response.content))["LIB"]
        header = self.header_row(sheet)
        column_a = [sheet.cell(row=r, column=1).value for r in range(header + 1, sheet.max_row + 1)]
        legend_text = " ".join(str(c.value) for row in sheet.iter_rows(min_row=header + 1) for c in row if c.value)

        self.assertEqual(response.status_code, 200)
        self.assertIn("rmis_lib_template.xlsx", response["Content-Disposition"])
        self.assertEqual(openpyxl.load_workbook(io.BytesIO(response.content)).sheetnames, ["LIB"])
        self.assertEqual([c.value for c in sheet[header] if c.value], self.FIXED + ["Q1", "Q2", "Q3", "Q4"])
        self.assertIn("LEGEND", column_a)
        self.assertIn("2027-01-01", " ".join(str(c.value) for row in sheet.iter_rows(max_row=header) for c in row if c.value))
        self.assertIn("Office Supplies Expenses", legend_text)

    def test_the_template_periods_follow_the_project_duration(self):
        cases = [
            ({"start_date": "2027-01-01", "target_end_date": "2027-08-31"}, ["Q1", "Q2", "Q3"]),
            ({"start_date": "2027-01-01", "target_end_date": "2028-03-31"}, ["Y1Q1", "Y1Q2", "Y1Q3", "Y1Q4", "Y2Q1"]),
            ({"start_date": "2027-01-01", "target_end_date": "2028-12-31"}, [f"Y{y}Q{q}" for y in (1, 2) for q in (1, 2, 3, 4)]),
        ]
        for dates, periods in cases:
            sheet = self.workbook(dates)["LIB"]
            self.assertEqual([c.value for c in sheet[self.header_row(sheet)] if c.value][len(self.FIXED):], periods, dates)

    def test_the_template_needs_valid_project_dates(self):
        for dates in ({}, {"start_date": "2027-01-01"}, {"start_date": "2027-13-01", "target_end_date": "2027-12-31"},
                      {"start_date": "2027-12-31", "target_end_date": "2027-01-01"}):
            response = self.client_for(self.admin).get(self.TEMPLATE_URL, dates)
            self.assertEqual(response.status_code, 400, dates)
            self.assertTrue(response.data["errors"], dates)

    def test_a_filled_template_maps_each_column_to_the_lib_row_and_saves_nothing(self):
        workbook = self.workbook()
        self.fill(workbook["LIB"], [
            ["MOOE", "office supplies expenses", "PAPER, 80 gsm, A4", "unit", 10, 225, None, 2250, None, None, None],
            ["PS", "Honoraria", None, "month", 2, "5,000", None, 5000, 5000, None, None],
        ])
        LibItem.objects.get_or_create(category="ps", name="Honoraria")

        response = self.parse(workbook)

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["rows"], [
            {"category": "mooe", "description": "Office Supplies Expenses", "specification": "PAPER, 80 gsm, A4", "unit": "unit",
             "quantity": "10", "unit_cost": "225", "amount": "2250.00",
             "allocations": [{"period": "Q1", "amount": "2250.00"}, {"period": "Q2", "amount": "0.00"},
                             {"period": "Q3", "amount": "0.00"}, {"period": "Q4", "amount": "0.00"}]},
            {"category": "ps", "description": "Honoraria", "specification": "", "unit": "month",
             "quantity": "2", "unit_cost": "5000", "amount": "10000.00",
             "allocations": [{"period": "Q1", "amount": "5000.00"}, {"period": "Q2", "amount": "5000.00"},
                             {"period": "Q3", "amount": "0.00"}, {"period": "Q4", "amount": "0.00"}]},
        ])
        self.assertFalse(LineItem.objects.exists())
        self.assertFalse(LineItemBudget.objects.exists())

    def test_a_multi_year_lib_maps_each_amount_to_its_year_and_quarter(self):
        dates = {"start_date": "2027-01-01", "target_end_date": "2028-03-31"}
        workbook = self.workbook(dates)
        self.fill(workbook["LIB"], [["CO", "Office Equipment", "Printer with Scanner (3 in 1)", "unit", 2, 15000, None, 0, 30000, 0, 0, 0]])

        response = self.parse(workbook, dates)

        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["rows"][0]["allocations"], [
            {"period": "Y1Q1", "amount": "0.00"}, {"period": "Y1Q2", "amount": "30000.00"}, {"period": "Y1Q3", "amount": "0.00"},
            {"period": "Y1Q4", "amount": "0.00"}, {"period": "Y2Q1", "amount": "0.00"},
        ])

    def test_a_workbook_for_another_duration_lists_the_missing_and_extra_periods(self):
        workbook = self.workbook()  # Q1-Q4
        self.fill(workbook["LIB"], [["MOOE", "Office Supplies Expenses", None, "unit", 1, 100, None, 100]])

        response = self.parse(workbook, {"start_date": "2027-01-01", "target_end_date": "2028-03-31"})

        self.assertEqual(response.status_code, 400)
        messages = " ".join(e["message"] for e in response.data["errors"])
        for period in ("Y1Q1", "Y1Q4", "Y2Q1", "Q1", "Q4"):
            self.assertIn(period, messages)
        self.assertNotIn("rows", response.data)

    def test_row_errors_name_the_row_and_nothing_comes_back(self):
        workbook = self.workbook()
        first = self.fill(workbook["LIB"], [
            ["MOOE", "Office Supplies Expenses", None, "unit", 1, 100, None, 100],
            ["XYZ", "Office Supplies Expenses", None, "unit", 1, 100, None, 100],
            ["MOOE", "Common Office Supplies", None, "unit", 1, 100, None, 100],
            ["MOOE", "Training Expenses", None, "unit", 1, 100, None, 60, 30],
            ["MOOE", "Training Expenses", None, "unit", "two", 100, None, 200],
            ["MOOE", "Training Expenses", None, None, 1, 100, None, 100],
            ["CO", "Office Equipment", None, "unit", 1, -5, None, -5],
        ])

        response = self.parse(workbook)

        self.assertEqual(response.status_code, 400)
        self.assertEqual(sorted({e["row"] for e in response.data["errors"]}), [first + n for n in range(1, 7)])
        self.assertTrue(all(e["sheet"] == "LIB" for e in response.data["errors"]))

    def test_an_empty_lib_a_missing_sheet_or_a_non_workbook_is_rejected(self):
        self.assertEqual(self.parse(self.workbook()).status_code, 400)

        other = openpyxl.Workbook()
        other.active.title = "Budget Requirements"
        self.assertEqual(self.parse(other).status_code, 400)

        buffer = io.BytesIO(b"not an excel file")
        buffer.name = "lib.xlsx"
        response = self.client_for(self.admin).post(self.PARSE_URL, {"file": buffer, **self.ONE_YEAR}, format="multipart")
        self.assertEqual(response.status_code, 400)

    def test_the_import_needs_the_project_dates(self):
        workbook = self.workbook()
        self.fill(workbook["LIB"], [["MOOE", "Office Supplies Expenses", None, "unit", 1, 100, None, 25, 25, 25, 25]])

        response = self.parse(workbook, {"start_date": "", "target_end_date": "2027-12-31"})

        self.assertEqual(response.status_code, 400)
        self.assertTrue(response.data["errors"])

    def test_a_failed_import_keeps_no_file_on_the_server(self):
        from document_management.models import Document

        self.parse(self.workbook())

        self.assertFalse(Document.objects.exists())

    def test_only_registrants_can_use_it(self):
        workbook = self.workbook()
        self.fill(workbook["LIB"], [["MOOE", "Office Supplies Expenses", None, "unit", 1, 100, None, 100]])
        leader = self.make_user("project_leader")

        self.assertEqual((self.template(user=leader).status_code, self.parse(workbook, user=leader).status_code), (200, 200))
        for code in ("project_staff", "finance_budget"):
            user = self.make_user(code)
            self.assertEqual((self.template(user=user).status_code, self.parse(workbook, user=user).status_code), (403, 403), code)

class ProjectReadScopeTests(RMISTestCase):
    """Strict RBAC (client meeting 2026-10-01): leaders and staff read only projects they lead or belong to."""

    def setUp(self):
        super().setUp()
        self.leader, self.other = self.make_user("project_leader"), self.make_user("project_leader")
        self.mine = Project.objects.create(title="Mine", project_code="M-1", funding_type="core_funded", lead=self.leader)
        self.theirs = Project.objects.create(title="Theirs", project_code="T-1", funding_type="core_funded", lead=self.other)

    def codes(self, user, path="/api/projects/"):
        response = self.client_for(user).get(path)
        self.assertEqual(response.status_code, 200, response.data)
        return sorted(row.get("project_code") or row.get("title") for row in response.data)

    def test_leader_lists_and_opens_only_their_projects(self):
        self.assertEqual(self.codes(self.leader), ["M-1"])
        self.assertEqual(self.client_for(self.leader).get(f"/api/projects/{self.theirs.id}/").status_code, 404)
        self.assertEqual(self.client_for(self.leader).get(f"/api/projects/{self.mine.id}/").status_code, 200)

    def test_team_members_and_assigned_staff_see_the_project(self):
        member, staff = self.make_user("project_staff"), self.make_user("project_staff")
        self.theirs.team_members.create(name="Member", user=member)
        from personnel.models import ProjectAssignment
        ProjectAssignment.objects.create(user=staff, project=self.theirs, start_date="2026-01-01")

        self.assertEqual(self.codes(member), ["T-1"])
        self.assertEqual(self.codes(staff), ["T-1"])
        self.assertEqual(self.codes(self.make_user("project_staff")), [])

    def test_project_leader_sees_only_projects_they_lead(self):
        """Client 2026-10-18: a Project Leader sees only their own projects, not ones where they are listed as a
        Co-Project Leader, team member or assigned elsewhere."""
        self.theirs.team_members.create(member_role="co_leader", name="Co", user=self.leader)
        self.theirs.studies.create(title="Their study", lead=self.leader)

        self.assertEqual(self.codes(self.leader), ["M-1"])
        self.assertEqual(self.client_for(self.leader).get(f"/api/projects/{self.theirs.id}/").status_code, 404)

    def test_university_wide_roles_still_see_everything(self):
        self.assertEqual(self.codes(self.make_user("vprei")), ["M-1", "T-1"])
        self.assertEqual(self.codes(self.make_user("system_admin")), ["M-1", "T-1"])

    def test_studies_and_milestones_follow_the_project_scope(self):
        self.theirs.studies.create(title="Their study")
        self.mine.studies.create(title="My study")

        self.assertEqual(self.codes(self.leader, "/api/studies/"), ["My study"])


class CollegeUnitTests(RMISTestCase):
    """System Admin manages the "College Unit - Implementing Unit" choices; everyone else only reads them."""

    def test_system_admin_adds_a_unit_as_abbreviation_and_college(self):
        client = self.client_for(self.make_user("system_admin"))

        created = client.post("/api/college-units/", {"code": " CA ", "title": " College of Agriculture "}, format="json")
        self.assertEqual(created.status_code, 201)
        self.assertEqual((created.data["code"], created.data["title"]), ("CA", "College of Agriculture"))
        self.assertEqual(created.data["name"], "CA - College of Agriculture")
        renamed = client.patch(f"/api/college-units/{created.data['id']}/", {"title": "College of Agri"}, format="json")
        self.assertEqual(renamed.data["name"], "CA - College of Agri")
        self.assertEqual(client.delete(f"/api/college-units/{created.data['id']}/").status_code, 204)

    def test_code_and_title_are_required(self):
        client = self.client_for(self.make_user("system_admin"))

        self.assertEqual(client.post("/api/college-units/", {"code": "CA"}, format="json").status_code, 400)
        self.assertEqual(client.post("/api/college-units/", {"code": " ", "title": "College of Agriculture"}, format="json").status_code, 400)

    def test_duplicate_codes_are_rejected_regardless_of_case(self):
        client = self.client_for(self.make_user("system_admin"))
        client.post("/api/college-units/", {"code": "CCS", "title": "College of Computer Studies"}, format="json")

        self.assertEqual(client.post("/api/college-units/", {"code": "ccs", "title": "Other"}, format="json").status_code, 400)

    def test_other_roles_can_read_but_not_edit(self):
        self.client_for(self.make_user("system_admin")).post("/api/college-units/", {"code": "CTE", "title": "College of Teacher Education"}, format="json")
        riuh = self.client_for(self.make_user("riuh"))

        self.assertEqual([u["name"] for u in riuh.get("/api/college-units/").data], ["CTE - College of Teacher Education"])
        self.assertEqual(riuh.post("/api/college-units/", {"code": "CAS", "title": "Arts"}, format="json").status_code, 403)

    def test_rde_thrust_name_is_code_then_title(self):
        admin, riuh = self.client_for(self.make_user("system_admin")), self.client_for(self.make_user("riuh"))

        created = admin.post("/api/rde-thrusts/", {"code": "9", "title": "Space Science"}, format="json")
        self.assertEqual(created.data["name"], "9 - Space Science")
        self.assertEqual(admin.post("/api/rde-thrusts/", {"code": "9", "title": "X"}, format="json").status_code, 400)
        self.assertEqual(riuh.delete(f"/api/rde-thrusts/{created.data['id']}/").status_code, 403)

    def test_cooperating_agencies_keep_a_single_name(self):
        admin, riuh = self.client_for(self.make_user("system_admin")), self.client_for(self.make_user("riuh"))
        created = admin.post("/api/cooperating-agencies/", {"name": "X"}, format="json")
        self.assertEqual(created.status_code, 201)
        self.assertEqual(admin.post("/api/cooperating-agencies/", {"name": "x"}, format="json").status_code, 400)
        self.assertEqual([c["name"] for c in riuh.get("/api/cooperating-agencies/").data], ["X"])
        self.assertEqual(riuh.delete(f"/api/cooperating-agencies/{created.data['id']}/").status_code, 403)
        self.assertEqual(admin.delete(f"/api/cooperating-agencies/{created.data['id']}/").status_code, 204)

    def test_endorsers_carry_a_designation(self):
        admin, riuh = self.client_for(self.make_user("system_admin")), self.client_for(self.make_user("riuh"))
        created = admin.post("/api/endorsers/", {"name": " Adriel G. Roman ", "designation": "Dean/Associate Dean"}, format="json")
        self.assertEqual((created.data["name"], created.data["designation"]), ("Adriel G. Roman", "Dean/Associate Dean"))
        self.assertEqual(admin.post("/api/endorsers/", {"name": "adriel g. roman"}, format="json").status_code, 400)
        self.assertEqual(admin.patch(f"/api/endorsers/{created.data['id']}/", {"designation": "Dean"}, format="json").data["designation"], "Dean")
        self.assertEqual(riuh.get("/api/endorsers/").data[0]["designation"], "Dean")
        self.assertEqual(riuh.post("/api/endorsers/", {"name": "X"}, format="json").status_code, 403)
        self.assertEqual(admin.delete(f"/api/endorsers/{created.data['id']}/").status_code, 204)


class EditRegisteredProjectTests(RMISTestCase):
    """Client request 2026-10-09: the Project Leader (and System Admin) edits a registered project through the
    SF-018 form on the project page; these are the calls its Save Changes makes."""

    def setUp(self):
        super().setUp()
        from outputs.models import ExpectedOutput
        from research_projects.models import ProjectTeamMember, Study, TargetBeneficiary, WorkPlanMilestone

        self.leader = self.make_user("project_leader")
        self.project = Project.objects.create(
            title="Old", project_code="LSPU-9", funding_type="core_funded", lead=self.leader, sectors=["education"], sdgs=[4],
        )
        self.member = ProjectTeamMember.objects.create(project=self.project, member_role="member", name="A")
        self.study = Study.objects.create(project=self.project, title="S1")
        self.output = ExpectedOutput.objects.create(project=self.project, category="publications", description="D")
        self.beneficiary = TargetBeneficiary.objects.create(project=self.project, group="Farmers", total=5)
        self.milestone = WorkPlanMilestone.objects.create(project=self.project, title="M", target_date=datetime.date(2026, 5, 1))
        self.endorser = ProjectEndorser.objects.create(project=self.project, name="X", designation="Dean")

    def test_the_leader_saves_every_section_but_not_the_code(self):
        c = self.client_for(self.leader)
        ok = lambda r: self.assertIn(r.status_code, (200, 201, 204), getattr(r, "data", r))
        ok(c.patch(f"/api/projects/{self.project.pk}/", {"title": "New", "sectors": ["education", "others"], "sector_other": "X",
                                                         "objectives": "1. A\n2. B", "start_date": "2027-02-01"}, format="json"))
        ok(c.patch(f"/api/project-team/{self.member.pk}/", {"member_role": "co_leader", "name": "B", "gender": "female"}, format="json"))
        ok(c.post("/api/project-team/", {"project": self.project.pk, "member_role": "member", "name": "C", "gender": ""}, format="json"))
        ok(c.patch(f"/api/studies/{self.study.pk}/", {"title": "S1b"}, format="json"))
        ok(c.post("/api/studies/", {"project": self.project.pk, "title": "S2"}, format="json"))
        ok(c.patch(f"/api/outputs/expected-outputs/{self.output.pk}/", {"category": "patents", "description": "E", "target_count": 2}, format="json"))
        ok(c.patch(f"/api/target-beneficiaries/{self.beneficiary.pk}/", {"group": "Fisherfolk", "description": "", "total": 9}, format="json"))
        ok(c.patch(f"/api/milestones/{self.milestone.pk}/", {"title": "M2", "start_date": None, "target_date": "2026-06-01"}, format="json"))
        ok(c.patch(f"/api/project-endorsers/{self.endorser.pk}/", {"name": "Y", "designation": "Dean", "signed_on": None}, format="json"))
        ok(c.delete(f"/api/target-beneficiaries/{self.beneficiary.pk}/"))
        ok(c.delete(f"/api/outputs/expected-outputs/{self.output.pk}/"))
        ok(c.delete(f"/api/milestones/{self.milestone.pk}/"))
        ok(c.delete(f"/api/project-endorsers/{self.endorser.pk}/"))
        ok(c.delete(f"/api/project-team/{self.member.pk}/"))

        self.project.refresh_from_db()
        self.assertEqual((self.project.title, self.project.sector_other), ("New", "X"))
        self.assertEqual(c.patch(f"/api/projects/{self.project.pk}/", {"ntp_number": "LSPU-10"}, format="json").status_code, 400)
        self.assertEqual(c.patch(f"/api/projects/{self.project.pk}/", {"start_date": None}, format="json").status_code, 400)

    def test_only_the_system_admin_and_the_projects_leader_edit_it(self):
        from budget_lib.models import BudgetReview

        # Returned for revision: visible to the CRC Chair (D2) and open for edits (D6)
        budget = LineItemBudget.objects.create(project=self.project, version_number=1, status="returned")
        BudgetReview.objects.create(budget=budget, action="held_crc", acted_by=self.make_user("crc_chair"))
        for code in ("crc_chair", "riuh"):
            c = self.client_for(self.make_user(code))
            self.assertEqual(c.patch(f"/api/projects/{self.project.pk}/", {"title": "Changed"}, format="json").status_code, 403, code)
            self.assertEqual(c.patch(f"/api/project-team/{self.member.pk}/", {"name": "Z"}, format="json").status_code, 403, code)
            self.assertEqual(c.patch(f"/api/project-endorsers/{self.endorser.pk}/", {"name": "Z"}, format="json").status_code, 403, code)
            self.assertEqual(c.patch(f"/api/studies/{self.study.pk}/", {"title": "Z"}, format="json").status_code, 403, code)
            self.assertEqual(c.delete(f"/api/target-beneficiaries/{self.beneficiary.pk}/").status_code, 403, code)
        crc = self.client_for(self.make_user("crc_chair"))
        closed = crc.patch(f"/api/projects/{self.project.pk}/", {"status": "completed", "status_remarks": "Done"}, format="json")
        self.assertEqual(closed.status_code, 200, closed.data)
        admin = self.client_for(self.make_user("system_admin"))
        recoded = admin.patch(f"/api/projects/{self.project.pk}/", {"title": "Fixed", "project_code": "LSPU-10"}, format="json")
        self.assertEqual(recoded.status_code, 200, recoded.data)


class LibAllocationTests(RMISTestCase):
    """Client meeting 2026-10-14 + yearly.png (2026-10-17): each LIB item is split per quarter of each project year
    (Y1Q1..YnQ4) on a multi-year project, per quarter (Q1..Q4) on a one-year project; only quarters the project covers."""

    def make_project(self, start, end, **fields):
        leader = self.make_user("project_leader")
        project = Project.objects.create(
            title="P", project_code=f"P-{Project.objects.count()}", funding_type="core_funded", lead=leader,
            start_date=start, target_end_date=end, **fields,
        )
        return leader, project

    def row(self, allocations, unit_cost="10000"):
        return {"category": "mooe", "description": "Traveling Expenses - Local", "unit": "lot", "quantity": "1", "unit_cost": unit_cost, "allocations": allocations}

    def test_periods_follow_the_project_duration(self):
        from budget_lib.models import allocation_periods
        d = datetime.date
        cases = [
            ((d(2026, 1, 1), d(2026, 12, 31)), ["Q1", "Q2", "Q3", "Q4"]),
            ((d(2026, 1, 15), d(2026, 7, 14)), ["Q1", "Q2"]),
            ((d(2026, 1, 1), d(2026, 8, 31)), ["Q1", "Q2", "Q3"]),
            ((d(2026, 1, 1), d(2027, 3, 31)), ["Y1Q1", "Y1Q2", "Y1Q3", "Y1Q4", "Y2Q1"]),
            ((d(2026, 1, 1), d(2028, 12, 31)), [f"Y{y}Q{q}" for y in (1, 2, 3) for q in (1, 2, 3, 4)]),
            ((d(2026, 1, 1), None), None),
        ]
        for (start, end), expected in cases:
            self.assertEqual(allocation_periods(Project(start_date=start, target_end_date=end)), expected, (start, end))

    def test_multi_year_project_saves_the_split_per_quarter_of_each_year(self):
        leader, project = self.make_project(datetime.date(2026, 1, 1), datetime.date(2027, 3, 31))
        split = [{"period": "Y1Q1", "amount": "5000"}, {"period": "Y1Q3", "amount": "2500"}, {"period": "Y2Q1", "amount": 2500}]

        response = self.client_for(leader).post(f"/api/projects/{project.id}/lib/", {"line_items": [self.row(split)]}, format="json")

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(LineItem.objects.get().allocations, [
            {"period": "Y1Q1", "amount": "5000.00"}, {"period": "Y1Q2", "amount": "0.00"}, {"period": "Y1Q3", "amount": "2500.00"},
            {"period": "Y1Q4", "amount": "0.00"}, {"period": "Y2Q1", "amount": "2500.00"},
        ])

    def test_quarters_after_the_end_date_are_refused(self):
        leader, project = self.make_project(datetime.date(2026, 1, 1), datetime.date(2027, 3, 31))
        late = [{"period": "Y1Q1", "amount": "5000"}, {"period": "Y2Q2", "amount": "5000"}]

        response = self.client_for(leader).post(f"/api/projects/{project.id}/lib/", {"line_items": [self.row(late)]}, format="json")

        self.assertEqual(response.status_code, 400)
        self.assertIn("allocations", response.data["errors"][0]["errors"])

    def test_one_year_project_uses_quarters_and_missing_quarters_are_zero(self):
        leader, project = self.make_project(datetime.date(2026, 1, 1), datetime.date(2026, 12, 31))
        split = [{"period": "Q1", "amount": "6000"}, {"period": "Q3", "amount": "4000"}]

        response = self.client_for(leader).post(f"/api/projects/{project.id}/lib/", {"line_items": [self.row(split)]}, format="json")

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual([a["amount"] for a in LineItem.objects.get().allocations], ["6000.00", "0.00", "4000.00", "0.00"])

    def test_split_must_add_up_and_use_the_projects_periods(self):
        leader, project = self.make_project(datetime.date(2026, 1, 1), datetime.date(2028, 12, 31))
        url = f"/api/projects/{project.id}/lib/"
        client = self.client_for(leader)

        short = client.post(url, {"line_items": [self.row([{"period": "Y1Q1", "amount": "5000"}])]}, format="json")
        per_year = client.post(url, {"line_items": [self.row([{"period": "Y1", "amount": "10000"}])]}, format="json")
        quarter = client.post(url, {"line_items": [self.row([{"period": "Q1", "amount": "10000"}])]}, format="json")
        missing = client.post(url, {"line_items": [self.row([])]}, format="json")
        negative = client.post(url, {"line_items": [self.row([{"period": "Y1Q1", "amount": "12000"}, {"period": "Y2Q1", "amount": "-2000"}])]}, format="json")

        for response in (short, per_year, quarter, missing, negative):
            self.assertEqual(response.status_code, 400)
            self.assertIn("allocations", response.data["errors"][0]["errors"])
        self.assertFalse(LineItemBudget.objects.exists())

    def test_projects_without_dates_and_budget_management_items_stay_optional(self):
        leader, undated = self.make_project(None, None)
        _, dated = self.make_project(datetime.date(2026, 1, 1), datetime.date(2028, 12, 31))
        plain = {"category": "mooe", "description": "Traveling Expenses - Local", "unit": "lot", "quantity": "1", "unit_cost": "100"}

        registered = self.client_for(leader).post(f"/api/projects/{undated.id}/lib/", {"line_items": [plain]}, format="json")
        budget = self.client_for(dated.lead).post("/api/budget/budgets/", {"project": dated.id}, format="json").data
        added = self.client_for(dated.lead).post("/api/budget/line-items/", {**plain, "budget": budget["id"], "amount": "100"}, format="json")

        self.assertEqual(registered.status_code, 201, registered.data)
        self.assertEqual(added.status_code, 201, added.data)
        self.assertEqual(added.data["allocations"], [])

    def test_changing_the_amount_needs_a_matching_split(self):
        leader, project = self.make_project(datetime.date(2026, 1, 1), datetime.date(2027, 12, 31))
        split = [{"period": "Y1Q1", "amount": "5000"}, {"period": "Y2Q1", "amount": "5000"}]
        self.client_for(leader).post(f"/api/projects/{project.id}/lib/", {"line_items": [self.row(split)]}, format="json")
        item = LineItem.objects.get()

        stale = self.client_for(leader).patch(f"/api/budget/line-items/{item.id}/", {"quantity": "2", "amount": "20000"}, format="json")

        self.assertEqual(stale.status_code, 400)
        self.assertIn("allocations", stale.data)

    def test_dry_cap_counts_each_years_allocation(self):
        leader, project = self.make_project(datetime.date(2026, 1, 1), datetime.date(2027, 12, 31), is_dry_research=True)
        Project.objects.filter(pk=project.pk).update(funding_type="institutional")
        even = [{"period": "Y1Q1", "amount": "40000"}, {"period": "Y1Q2", "amount": "40000"}, {"period": "Y2Q1", "amount": "80000"}]
        # no single quarter is over P100k, but Year 1's four quarters together are
        lopsided = [{"period": "Y1Q1", "amount": "60000"}, {"period": "Y1Q4", "amount": "60000"}, {"period": "Y2Q1", "amount": "40000"}]

        ok = self.client_for(leader).post(f"/api/projects/{project.id}/lib/", {"line_items": [self.row(even, "160000")]}, format="json")
        LineItemBudget.objects.all().delete()
        over = self.client_for(leader).post(f"/api/projects/{project.id}/lib/", {"line_items": [self.row(lopsided, "160000")]}, format="json")

        self.assertFalse(ok.data["exceeds_dry_cap"])
        self.assertTrue(over.data["exceeds_dry_cap"])

    def test_not_applicable_wet_or_dry_has_no_dry_cap(self):
        leader, project = self.make_project(datetime.date(2026, 1, 1), datetime.date(2027, 12, 31), is_dry_research=None)
        Project.objects.filter(pk=project.pk).update(funding_type="institutional")
        over = [{"period": "Y1Q1", "amount": "60000"}, {"period": "Y1Q4", "amount": "60000"}, {"period": "Y2Q1", "amount": "40000"}]

        response = self.client_for(leader).post(f"/api/projects/{project.id}/lib/", {"line_items": [self.row(over, "160000")]}, format="json")

        self.assertFalse(response.data["exceeds_dry_cap"])


class NotApplicableClassificationTests(RMISTestCase):
    """Client 2026-10-23: "Not Applicable" for Wet or Dry Research and for Basic or Applied Research."""

    def setUp(self):
        super().setUp()
        self.leader = self.make_user("project_leader")
        self.project = Project.objects.create(
            title="P", project_code="NA-1", funding_type="core_funded", lead=self.leader, sectors=["education"], sdgs=[4],
        )
        self.client_ = self.client_for(self.leader)
        self.url = f"/api/projects/{self.project.pk}/"

    def test_not_applicable_is_saved_and_read_back(self):
        saved = self.client_.patch(self.url, {"is_dry_research": None, "research_type": "not_applicable"}, format="json")

        read = self.client_.get(self.url).data
        self.assertEqual(saved.status_code, 200, saved.data)
        self.assertEqual((read["is_dry_research"], read["research_type"]), (None, "not_applicable"))

    def test_the_existing_values_still_save(self):
        for dry, kind in ((False, "basic"), (True, "applied")):
            response = self.client_.patch(self.url, {"is_dry_research": dry, "research_type": kind}, format="json")
            self.assertEqual(response.status_code, 200, response.data)
            self.assertEqual((response.data["is_dry_research"], response.data["research_type"]), (dry, kind))
        # Required since 2026-10-29: it can't be blanked any more
        self.assertEqual(self.client_.patch(self.url, {"research_type": ""}, format="json").status_code, 400)

    def test_a_new_project_is_still_dry_by_default(self):
        self.assertIs(Project.objects.get(pk=self.project.pk).is_dry_research, True)

    def test_unknown_research_type_is_rejected(self):
        self.assertEqual(self.client_.patch(self.url, {"research_type": "mixed"}, format="json").status_code, 400)

    def test_sf018_leaves_every_box_unchecked_when_not_applicable(self):
        from reports.forms import proposal_form, render_form_csv

        Project.objects.filter(pk=self.project.pk).update(is_dry_research=None, research_type="not_applicable")

        text = render_form_csv(proposal_form(self.project.pk)).decode()
        for label in ("Basic Research", "Applied Research", "Wet Research (i.e., with laboratory)", "Dry Research"):
            self.assertIn(f"[ ] {label}", text)
        self.assertNotIn("Not Applicable", text)

    def test_sf018_still_ticks_the_chosen_boxes(self):
        from reports.forms import proposal_form, render_form_csv

        Project.objects.filter(pk=self.project.pk).update(is_dry_research=False, research_type="applied")

        text = render_form_csv(proposal_form(self.project.pk)).decode()
        self.assertIn("[/] Applied Research", text)
        self.assertIn("[/] Wet Research (i.e., with laboratory)", text)
        self.assertIn("[ ] Dry Research", text)


class RichTextProposalTests(RMISTestCase):
    """Client 2026-10-23: Sections II, IV, VI, VIII and IX are rich text (TipTap HTML). It is sanitized on save, older
    plain text stays as it is, and the downloaded SF-018 / Appendix F show it as readable text."""

    def setUp(self):
        super().setUp()
        self.leader = self.make_user("project_leader")
        self.project = Project.objects.create(
            title="P", project_code="RT-1", funding_type="core_funded", lead=self.leader, sectors=["education"], sdgs=[4],
        )
        self.url = f"/api/projects/{self.project.pk}/"

    def save(self, **fields):
        response = self.client_for(self.leader).patch(self.url, fields, format="json")
        self.assertEqual(response.status_code, 200, response.data)
        return response.data

    def test_unsafe_html_is_removed_and_formatting_kept(self):
        data = self.save(background='<p onclick="steal()">Hi <strong>there</strong> <em>and</em> <u>you</u></p>'
                                    '<script>alert(1)</script><p><a href="javascript:alert(1)">x</a></p>')

        self.assertEqual(data["background"], "<p>Hi <strong>there</strong> <em>and</em> <u>you</u></p><p><a rel=\"noopener noreferrer\">x</a></p>")

    def test_lists_headings_and_tables_are_kept(self):
        html = ("<h2>Design</h2><ul><li><p>One</p></li></ul><ol><li><p>First</p></li></ol>"
                "<table><tbody><tr><th><p>Item</p></th><td colspan=\"2\"><p>Cost</p></td></tr></tbody></table>")

        self.assertEqual(self.save(methodology=html)["methodology"], html)

    def test_every_rich_section_is_sanitized(self):
        fields = ("background", "methodology", "socio_economic_significance", "monitoring_evaluation", "references")
        data = self.save(**{f: "<p>ok</p><img src=x onerror=alert(1)>" for f in fields})

        self.assertEqual([data[f] for f in fields], ["<p>ok</p>"] * 5)

    def test_plain_text_stays_as_it_is(self):
        text = "Rice & corn yields\n2. Farmers < 50 years old"

        self.assertEqual(self.save(background=text)["background"], text)

    def test_sf018_shows_rich_text_as_readable_lines(self):
        from reports.forms import proposal_form, render_form_csv

        Project.objects.filter(pk=self.project.pk).update(
            methodology="<h2>Design</h2><p>We <strong>survey</strong> farms &amp; fisheries.</p>"
                        "<ul><li><p>One</p></li><li><p>Two</p></li></ul><ol><li><p>First</p></li><li><p>Second</p></li></ol>"
                        "<table><tbody><tr><th><p>Item</p></th><th><p>Cost</p></th></tr><tr><td><p>Seeds</p></td><td><p>500</p></td></tr></tbody></table>",
            references="Plain reference, 2020.\nAnother, 2021.",
        )

        text = render_form_csv(proposal_form(self.project.pk)).decode()
        for line in ("Design", "We survey farms & fisheries.", "• One", "• Two", "1. First", "2. Second", "Item | Cost", "Seeds | 500"):
            self.assertIn(line, text)
        self.assertIn("Plain reference, 2020.\nAnother, 2021.", text)
        self.assertNotIn("<", text.split("IV. PROJECT DESCRIPTIONS/METHODOLOGY")[1].split("VI.")[0])

    def test_appendix_f_shows_rich_text_as_readable_lines(self):
        from monitoring.models import TerminalReport
        from reports.forms import appendix_f_form, render_form_csv

        Project.objects.filter(pk=self.project.pk).update(background="<p>Why <em>this</em> matters</p><ul><li><p>Food</p></li></ul>")
        TerminalReport.objects.create(project=self.project, submitted_by=self.leader)

        text = render_form_csv(appendix_f_form(self.project.pk)).decode()
        self.assertIn("Why this matters\n• Food", text)
        self.assertNotIn("<p>", text)

    def test_marked_lines_keep_bold_italic_underline(self):
        """Carl 2026-10-25 (Task 9): the SF-018 PDF/Word keep bold, italic and underline; links and strikethrough stay text."""
        from reports.richtext import rich_text_marks

        self.assertEqual(rich_text_marks("<p>We <strong>survey</strong> <em>farms</em> &amp; <u>fish</u></p><p><b>Both <i>marks</i></b></p>"),
                         "We <b>survey</b> <i>farms</i> &amp; <u>fish</u>\n<b>Both <i>marks</i></b>")
        self.assertEqual(rich_text_marks("<p><strong>One<br>Two</strong></p>"), "<b>One</b>\n<b>Two</b>")
        self.assertEqual(rich_text_marks("<ul><li><p><em>Item</em></p></li></ul>"), "• <i>Item</i>")
        self.assertEqual(rich_text_marks('<p><s>old</s> <a href="https://x.ph">site</a> &lt;tag&gt;</p>'), "old site &lt;tag&gt;")
        self.assertEqual(rich_text_marks("Rice & corn < 50\nline 2"), "Rice &amp; corn &lt; 50\nline 2")

    def test_sf018_word_keeps_bold_italic_underline(self):
        import io

        from docx import Document

        from reports.forms import proposal_form, render_form_docx

        Project.objects.filter(pk=self.project.pk).update(methodology="<p>We <strong>survey</strong> <em>farms</em> &amp; <u>fish</u></p>")

        doc = Document(io.BytesIO(render_form_docx(proposal_form(self.project.pk))))
        runs = {run.text: run for table in doc.tables for row in table.rows for cell in row.cells
                for paragraph in cell.paragraphs for run in paragraph.runs}
        self.assertTrue(runs["survey"].bold)
        self.assertTrue(runs["farms"].italic)
        self.assertTrue(runs["fish"].underline)
        self.assertFalse(runs["We "].bold)
        self.assertIn(" & ", runs)

    def test_sf018_pdf_keeps_marks_and_still_breaks_across_pages(self):
        from reports.forms import proposal_form, render_form_pdf

        plain = render_form_pdf(proposal_form(self.project.pk))
        Project.objects.filter(pk=self.project.pk).update(
            background="".join(f"<p>Line {i} <strong>bold</strong> <em>italic</em> &amp; <u>under</u></p>" for i in range(150)))

        marked = render_form_pdf(proposal_form(self.project.pk))
        self.assertFalse(b"/Times-Italic" in plain)
        self.assertTrue(b"/Times-Italic" in marked)  # the font dictionaries aren't compressed

    def test_appendix_f_word_keeps_bold_italic_underline(self):
        """Carl 2026-10-25 (Task 9b): Appendix F's sections taken from the SF-018 keep their formatting too."""
        import io

        from docx import Document

        from monitoring.models import TerminalReport
        from reports.forms import appendix_f_form, render_form_docx

        Project.objects.filter(pk=self.project.pk).update(methodology="<p>We <strong>survey</strong> <em>farms</em> &amp; <u>fish</u></p><p>Next line</p>")
        TerminalReport.objects.create(project=self.project, submitted_by=self.leader)

        doc = Document(io.BytesIO(render_form_docx(appendix_f_form(self.project.pk))))
        runs = {run.text: run for paragraph in doc.paragraphs for run in paragraph.runs}
        self.assertTrue(runs["survey"].bold)
        self.assertTrue(runs["farms"].italic)
        self.assertTrue(runs["fish"].underline)
        self.assertFalse(runs["We "].bold)
        self.assertIn("Next line", runs)

    def test_appendix_f_pdf_keeps_marks(self):
        from monitoring.models import TerminalReport
        from reports.forms import appendix_f_form, render_form_pdf

        TerminalReport.objects.create(project=self.project, submitted_by=self.leader)
        plain = render_form_pdf(appendix_f_form(self.project.pk))
        Project.objects.filter(pk=self.project.pk).update(
            background="".join(f"<p>Line {i} <strong>bold</strong> <em>italic</em> &amp; <u>under</u></p>" for i in range(150)))

        marked = render_form_pdf(appendix_f_form(self.project.pk))
        self.assertFalse(b"/Times-Italic" in plain)
        self.assertTrue(b"/Times-Italic" in marked)


class DuplicateTitleTests(RMISTestCase):
    """Client 2026-10-24: a project whose title already exists can't be registered; spacing and capitals don't count."""

    def setUp(self):
        super().setUp()
        self.admin, self.leader = self.make_user("system_admin"), self.make_user("project_leader")
        self.existing = Project.objects.create(
            title="Bridging Academia and Communities", project_code="OLD-1", funding_type="core_funded", lead=self.leader,
            sectors=["education"], sdgs=[4], status="archived",
        )

    def test_the_same_title_is_refused_even_with_other_spacing_or_capitals(self):
        for title in ("Bridging Academia and Communities", "  bridging academia   AND communities "):
            response = self.client_for(self.admin).post("/api/projects/", project_payload(self.leader, title=title, ntp_number="NEW-1"), format="json")

            self.assertEqual(response.status_code, 400, title)
            self.assertEqual(str(response.data["title"][0]), "A project titled \"Bridging Academia and Communities\" already exists.")
        self.assertFalse(Project.objects.filter(project_code="NEW-1").exists())

    def test_a_different_title_registers(self):
        response = self.client_for(self.admin).post(
            "/api/projects/", project_payload(self.leader, title="Bridging Academia and Communities 2", ntp_number="NEW-1"), format="json",
        )

        self.assertEqual(response.status_code, 201, response.data)

    def test_editing_a_project_keeps_its_own_title(self):
        response = self.client_for(self.admin).patch(
            f"/api/projects/{self.existing.pk}/", {"title": "Bridging Academia and Communities", "description": "x"}, format="json",
        )

        self.assertEqual(response.status_code, 200, response.data)

    def test_editing_a_project_cannot_take_another_projects_title(self):
        other = Project.objects.create(title="Other", project_code="OLD-2", funding_type="core_funded", lead=self.leader, sectors=["education"], sdgs=[4])

        response = self.client_for(self.admin).patch(f"/api/projects/{other.pk}/", {"title": "BRIDGING academia and communities"}, format="json")

        self.assertEqual(response.status_code, 400)
        self.assertIn("title", response.data)


class ProjectReviewWorkflowTests(RMISTestCase):
    """Client 2026-10-24 (Figma ProjectReviewWorkflow, PL-after-submit.png): a project is a Draft until its Project
    Leader submits it (with its LIB, D1/D3) and the Notice of Approval, NTP and MOA (Task 8); once submitted nobody
    edits its registration until the CRC Chair or Budget Officer returns it (D6, Carl: every role)."""

    def setUp(self):
        super().setUp()
        from outputs.models import ExpectedOutput
        from research_projects.models import ProjectTeamMember, Study

        self.admin, self.leader, self.crc = self.make_user("system_admin"), self.make_user("project_leader"), self.make_user("crc_chair")
        self.project = Project.objects.create(
            title="Bridging", project_code="LSPU-7", funding_type="core_funded", lead=self.leader, sectors=["education"], sdgs=[4],
            total_cost=50000,
        )
        self.budget = LineItemBudget.objects.create(project=self.project, version_number=1)
        LineItem.objects.create(budget=self.budget, category="mooe", description="Travel Expenses", amount=50000)
        self.member = ProjectTeamMember.objects.create(project=self.project, member_role="member", name="A")
        self.study = Study.objects.create(project=self.project, title="S1")
        self.output = ExpectedOutput.objects.create(project=self.project, category="publications", description="D")
        self.milestone = WorkPlanMilestone.objects.create(project=self.project, title="M", target_date=datetime.date(2026, 5, 1))

    def attach(self, *types):
        from document_management.models import Document

        for t in types:
            Document.objects.create(project=self.project, document_type=t, version_number=1, storage_path=f"x/{t}", file_name=f"{t}.pdf",
                                    file_size=1, uploaded_by=self.leader)

    def review_status(self):
        return self.client_for(self.leader).get(f"/api/projects/{self.project.pk}/").data["review_status"]

    def lib_status(self, status, hold=None):
        from budget_lib.models import BudgetReview

        LineItemBudget.objects.filter(pk=self.budget.pk).update(status=status)
        if hold:
            BudgetReview.objects.create(budget=self.budget, action=hold, acted_by=self.crc)

    def test_the_review_status_follows_the_lib(self):
        for status, hold, expected in (
            ("draft", None, "draft"), ("submitted", None, "submitted"), ("crc_processing", None, "pending_research_review"),
            ("for_checking", None, "pending_research_review"), ("returned", "held_crc", "returned_research"),
            # Task 17 (2026-10-30): forwarded to the Budget Officer (Ready for Certification) is its own status
            ("endorsed", None, "for_budget_review"), ("certifying", None, "pending_budget_review"),
            ("returned", "held_budget", "returned_budget"), ("certified", None, "reviews_completed"),
        ):
            self.lib_status(status, hold)
            self.assertEqual(self.review_status(), expected, status)
        response = self.client_for(self.leader).get(f"/api/projects/{self.project.pk}/")
        self.assertEqual(response.data["review_status_label"], "Active")

    def test_the_review_status_labels_are_the_clients(self):
        """Task 17 (2026-10-30): the labels every dashboard shows, the same for the Project Leader and the CRC/CRD."""
        from research_projects.workflow import REVIEW_STATUS_LABELS

        self.assertEqual(REVIEW_STATUS_LABELS, {
            "draft": "Draft", "submitted": "Submitted", "pending_research_review": "Ongoing Review - CRC/CRD",
            "returned_research": "Returned for Revision - CRC/CRD", "for_budget_review": "For Budget Review",
            "pending_budget_review": "Ongoing Review - Budget", "returned_budget": "Returned Review - Budget",
            "reviews_completed": "Active",
        })
        self.lib_status("submitted")  # a Draft is hidden from the CRC/CRD
        crc_view = self.client_for(self.crc).get(f"/api/projects/{self.project.pk}/").data
        self.assertEqual(crc_view["review_status_label"], self.client_for(self.leader).get(f"/api/projects/{self.project.pk}/").data["review_status_label"])

    def test_a_project_without_a_lib_is_a_draft(self):
        LineItemBudget.objects.all().delete()

        self.assertEqual(self.review_status(), "draft")

    def test_submitting_needs_the_approve_proposal_and_ntp_but_not_the_moa(self):
        # Task 13 (2026-10-29): the Notice of Approval slot is labelled Approve Proposal; the MOA is optional now
        submit = lambda: self.client_for(self.leader).post(f"/api/budget/budgets/{self.budget.pk}/submit/")

        nothing = submit()

        self.assertEqual(nothing.status_code, 400)
        self.assertIn("Approve Proposal", nothing.data["detail"])
        self.assertIn("Notice to Proceed", nothing.data["detail"])
        self.assertNotIn("MOA", nothing.data["detail"])
        self.attach("ntp")
        only_ntp = submit()
        self.assertEqual(only_ntp.status_code, 400)
        self.assertNotIn("Notice to Proceed", only_ntp.data["detail"])
        self.attach("notice_of_approval")
        Study.objects.create(project=self.project, title="S2")
        TargetBeneficiary.objects.create(project=self.project, group="Faculty", description="CTE", total=10)
        self.assertEqual(submit().status_code, 200)
        self.assertEqual(self.review_status(), "submitted")

    def test_submitting_needs_2_studies_a_6p_and_a_beneficiary(self):
        # Client 2026-10-29 (Task 4): the form's minimums, checked on submit so Edit can still save row by row
        self.attach("notice_of_approval", "ntp", "moa")
        self.output.delete()
        submit = lambda: self.client_for(self.leader).post(f"/api/budget/budgets/{self.budget.pk}/submit/")

        refused = submit()

        self.assertEqual(refused.status_code, 400)
        for text in ("2 Study Component Titles", "6P expected output", "Target Beneficiary"):
            self.assertIn(text, refused.data["detail"])
        self.assertEqual(self.review_status(), "draft")
        Study.objects.create(project=self.project, title="S2")
        ExpectedOutput.objects.create(project=self.project, category="publications", description="D")
        TargetBeneficiary.objects.create(project=self.project, group="Faculty", description="CTE", total=10)
        self.assertEqual(submit().status_code, 200)

    def test_nobody_edits_a_submitted_registration(self):
        self.lib_status("submitted")
        for user in (self.leader, self.admin):
            c = self.client_for(user)
            responses = {
                "project": c.patch(f"/api/projects/{self.project.pk}/", {"title": "New"}, format="json"),
                "team add": c.post("/api/project-team/", {"project": self.project.pk, "member_role": "member", "name": "B"}, format="json"),
                "team edit": c.patch(f"/api/project-team/{self.member.pk}/", {"name": "B"}, format="json"),
                "team remove": c.delete(f"/api/project-team/{self.member.pk}/"),
                "study": c.patch(f"/api/studies/{self.study.pk}/", {"title": "S2"}, format="json"),
                "study add": c.post("/api/studies/", {"project": self.project.pk, "title": "S3"}, format="json"),
                "6P": c.patch(f"/api/outputs/expected-outputs/{self.output.pk}/", {"target_count": 3}, format="json"),
                "6P add": c.post("/api/outputs/expected-outputs/", {"project": self.project.pk, "category": "patents", "description": "E"}, format="json"),
                "6P remove": c.delete(f"/api/outputs/expected-outputs/{self.output.pk}/"),
                "work plan": c.patch(f"/api/milestones/{self.milestone.pk}/", {"title": "M2"}, format="json"),
                "work plan add": c.post("/api/milestones/", {"project": self.project.pk, "title": "M3", "target_date": "2026-06-01"}, format="json"),
                "work plan remove": c.delete(f"/api/milestones/{self.milestone.pk}/"),
            }
            for name, response in responses.items():
                self.assertEqual(response.status_code, 400, f"{user.role.code} {name}: {getattr(response, 'data', '')}")
        self.project.refresh_from_db()
        self.assertEqual(self.project.title, "Bridging")

    def test_closure_status_and_6p_actual_counts_still_change_after_submitting(self):
        self.lib_status("certified")
        c = self.client_for(self.admin)

        self.assertEqual(c.patch(f"/api/projects/{self.project.pk}/", {"status": "completed"}, format="json").status_code, 200)
        self.assertEqual(c.patch(f"/api/outputs/expected-outputs/{self.output.pk}/", {"manual_actual_count": 2}, format="json").status_code, 200)
        self.assertEqual(c.patch(f"/api/studies/{self.study.pk}/", {"status": "completed"}, format="json").status_code, 200)

    def test_a_returned_project_can_be_edited_again(self):
        for hold in ("held_crc", "held_budget"):
            self.lib_status("returned", hold)

            response = self.client_for(self.leader).patch(f"/api/projects/{self.project.pk}/", {"title": f"Fixed {hold}"}, format="json")

            self.assertEqual(response.status_code, 200, response.data)
